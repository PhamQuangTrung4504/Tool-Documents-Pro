# -*- coding: utf-8 -*-
"""
tests.test_ocr_symbol_recovery
==============================
Bộ kiểm thử hồi quy (Regression Tests) toàn diện cho Context-Aware Symbol Recovery
và OCR Pipeline:
A. Khôi phục số La Mã đầu mục (Dương tính & Âm tính).
B. Chuẩn hóa Bullet và dấu cộng (+) theo ngữ cảnh (Dương tính & Âm tính).
C. Bảo toàn dấu phần trăm (%) hợp lệ.
D. Tái dựng bảng và bảo toàn số liệu bảng.
E. Tính bất biến đối với dữ liệu văn bản hành chính, đường dẫn, mã kỹ thuật.
F. Đánh giá chất lượng phép sửa: Đo lường độ chính xác (Precision = 100%).
"""

import os
import sys

# Đảm bảo import được module app
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

import pytest
from app.core.ocr_engine import (
    ContextSymbolRecovery,
    SymbolDecision,
    OCRBox,
    OCRResult,
    OCRTable,
    detect_and_reconstruct_table,
    reconstruct_document_layout,
    VietnameseTextPostProcessor,
)


class TestRomanNumeralsRecovery:
    """A. KIỂM THỬ SỐ LA MÃ ĐẦU MỤC (Mục 12.A)."""

    @pytest.mark.parametrize(
        "input_text,expected",
        [
            ("| NỘI DUNG, CHƯƠNG TRÌNH SINH HOẠT CHI BỘ", "I. NỘI DUNG, CHƯƠNG TRÌNH SINH HOẠT CHI BỘ"),
            ("|. NỘI DUNG SINH HOẠT", "I. NỘI DUNG SINH HOẠT"),
            ("l. NỘI DUNG KẾ HOẠCH", "I. NỘI DUNG KẾ HOẠCH"),
            ("TI. TIẾN HÀNH SINH HOẠT", "II. TIẾN HÀNH SINH HOẠT"),
            ("ll. TIẾN HÀNH SINH HOẠT", "II. TIẾN HÀNH SINH HOẠT"),
            ("11. TIẾN HÀNH SINH HOẠT", "II. TIẾN HÀNH SINH HOẠT"),
            ("TI TIẾN HÀNH SINH HOẠT", "II. TIẾN HÀNH SINH HOẠT"),
            ("lll. KẾT LUẬN", "III. KẾT LUẬN"),
            ("111. KẾT LUẬN", "III. KẾT LUẬN"),
            ("TII. KẾT LUẬN", "III. KẾT LUẬN"),
            ("lV. TỔ CHỨC THỰC HIỆN", "IV. TỔ CHỨC THỰC HIỆN"),
            ("1V. TỔ CHỨC THỰC HIỆN", "IV. TỔ CHỨC THỰC HIỆN"),
            ("TV. TỔ CHỨC THỰC HIỆN", "IV. TỔ CHỨC THỰC HIỆN"),
            ("II. ĐÁNH GIÁ TÌNH HÌNH", "II. ĐÁNH GIÁ TÌNH HÌNH"),
            ("III. PHƯƠNG HƯỚNG NHIỆM VỤ", "III. PHƯƠNG HƯỚNG NHIỆM VỤ"),
        ],
    )
    def test_roman_numerals_positive(self, input_text, expected):
        """Khôi phục thành công các biến thể OCR của số La Mã đầu mục khi có bằng chứng ngữ cảnh."""
        recovered, dec, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line(input_text)
        assert recovered == expected
        if input_text != expected:
            assert dec == SymbolDecision.AUTO_CORRECT

    @pytest.mark.parametrize(
        "negative_input",
        [
            "1. NỘP HỒ SƠ",
            "2. KẾT QUẢ",
            "1. Tuyên bố lý do, giới thiệu đại biểu:",
            "2. Chủ trì: Đồng chí Lê Văn Điền, Bí thư chi bộ.",
            "3. Thư ký: Đồng chí Nguyễn Văn Thắng, Chi ủy viên Chi bộ.",
            "4. Đồng chí Lê Văn Điền thông tin tình hình...",
            "5. Đồng chí Lê Thụy Thúy Phượng thông báo...",
            "6. Đồng chí Lê Văn Điền thông qua nội dung...",
            "Số 01-BB/CB",
            "Công văn số 1225-CV/ĐU ngày 21/9/2026",
            "14 đồng chí (14 chính thức, 0 dự bị)",
            "Đảng ủy xã Phước Hải vào lúc 15 giờ",
            "Mã hiệu: TC-01/2026-XNK",
        ],
    )
    def test_roman_numerals_negative(self, negative_input):
        """Bảo vệ tuyệt đối số thứ tự hợp lệ, số hiệu văn bản, ngày tháng, không đổi thành số La Mã."""
        recovered, dec, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line(negative_input)
        assert recovered == negative_input
        assert dec == SymbolDecision.PRESERVE


class TestBulletRecovery:
    """B. KIỂM THỬ BULLET VÀ DẤU CỘNG (Mục 12.B)."""

    @pytest.mark.parametrize(
        "input_text,expected",
        [
            (
                "% Chi bộ dưới 30 đảng viên ít nhất 90 phút; chi bộ trên 30 đảng viên ít nhất 120 phút",
                "+ Chi bộ dưới 30 đảng viên ít nhất 90 phút; chi bộ trên 30 đảng viên ít nhất 120 phút",
            ),
            (
                "% Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút",
                "+ Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút",
            ),
            (
                "% Từ 90% trở lên",
                "+ Từ 90% trở lên",
            ),
            (
                "% Từ 75% đến dưới 90%",
                "+ Từ 75% đến dưới 90%",
            ),
            (
                "% Từ 65% đến dưới 75%",
                "+ Từ 65% đến dưới 75%",
            ),
            (
                "% Không có đảng viên vắng mặt có lý do quá 03 lần liên tiếp trong năm'",
                "+ Không có đảng viên vắng mặt có lý do quá 03 lần liên tiếp trong năm'",
            ),
            (
                "| Chi bộ thực hiện tốt nhiệm vụ",
                "+ Chi bộ thực hiện tốt nhiệm vụ",
            ),
        ],
    )
    def test_bullet_positive(self, input_text, expected):
        """Khôi phục chính xác dấu cộng (+) đầu dòng bị OCR thành % hoặc |."""
        recovered, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(input_text)
        assert recovered == expected
        assert dec == SymbolDecision.AUTO_CORRECT

    @pytest.mark.parametrize(
        "negative_input",
        [
            "a + b = c",
            "50% + 50% = 100%",
            "Tỷ lệ đạt 90%.",
            "vào lúc 15 giờ",
            "ông Nguyễn Văn An",
            "- Lý do: Tổ chức sinh hoạt chi bộ định kỳ tháng 10 năm 2026.",
            "- Giới thiệu đại biểu: không có.",
            "– Nội dung sinh hoạt theo hướng dẫn.",
            "* Ghi chú quan trọng.",
            "1. Tuyên bố lý do",
            "2. Chủ trì hội nghị",
            "5. Đồng chí thông báo tình hình",
            "6. Đồng chí thông qua nội dung",
        ],
    )
    def test_bullet_negative(self, negative_input):
        """Bảo vệ tuyệt đối phép toán, từ ngữ tiếng Việt (v, o), số thứ tự và gạch ngang hợp lệ."""
        recovered, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(negative_input)
        assert recovered == negative_input
        assert dec == SymbolDecision.PRESERVE


class TestPercentagePreservation:
    """C. BẢO TOÀN DẤU PHẦN TRĂM HỢP LỆ (Mục 12.C)."""

    @pytest.mark.parametrize(
        "input_text",
        [
            "90%",
            "75%",
            "100%",
            "0.5%",
            "Tỷ lệ đạt 90% theo quy định.",
            "Dưới 60% hoặc có đảng viên vắng không có lý do",
            "Đạt chỉ tiêu 100% trong năm 2026",
            "STT | Tiêu chí | Đạt 90% | Đạt 100% | Ghi chú",
            "Tăng trưởng 8.5% so với cùng kỳ",
        ],
    )
    def test_percentage_intact(self, input_text):
        """Không làm mất hoặc biến đổi bất kỳ giá trị phần trăm hợp lệ nào."""
        recovered, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(input_text)
        assert recovered == input_text
        assert "%" in recovered


class TestTableAndDataIntegrity:
    """D. BẢO TOÀN BẢNG BIỂU VÀ SỐ LIỆU (Mục 12.D)."""

    def test_table_stt_column_alignment(self):
        """
        Kiểm tra phục hồi và bảo toàn cột STT trong bảng biểu:
        Số 2 không bị lệch sang cột 1 (Tiêu chí), mà luôn ở cột 0 (STT).
        Dấu % ở đầu ô tiêu chí được khôi phục thành +, nhưng 90% ở sau được giữ nguyên.
        """
        box_info = [
            # Header
            {"box": OCRBox("STT", 0.95, [[30, 110], [65, 110], [65, 135], [30, 135]]), "text": "STT", "left": 30, "right": 65, "top": 110, "bottom": 135, "w": 35, "h": 25, "cy": 122.5},
            {"box": OCRBox("Tiêu chí đánh giá", 0.95, [[150, 110], [400, 110], [400, 135], [150, 135]]), "text": "Tiêu chí đánh giá", "left": 150, "right": 400, "top": 110, "bottom": 135, "w": 250, "h": 25, "cy": 122.5},
            {"box": OCRBox("Điểm tối đa", 0.95, [[630, 110], [690, 110], [690, 135], [630, 135]]), "text": "Điểm tối đa", "left": 630, "right": 690, "top": 110, "bottom": 135, "w": 60, "h": 25, "cy": 122.5},
            {"box": OCRBox("Điểm tự chấm", 0.95, [[710, 110], [770, 110], [770, 135], [710, 135]]), "text": "Điểm tự chấm", "left": 710, "right": 770, "top": 110, "bottom": 135, "w": 60, "h": 25, "cy": 122.5},
            {"box": OCRBox("Ghi chú", 0.95, [[790, 110], [840, 110], [840, 135], [790, 135]]), "text": "Ghi chú", "left": 790, "right": 840, "top": 110, "bottom": 135, "w": 50, "h": 25, "cy": 122.5},

            # Row 1 (STT: 1)
            {"box": OCRBox("1", 0.95, [[40, 150], [55, 150], [55, 175], [40, 175]]), "text": "1", "left": 40, "right": 55, "top": 150, "bottom": 175, "w": 15, "h": 25, "cy": 162.5},
            {"box": OCRBox("Chấp hành thời gian, thời lượng sinh hoạt", 0.95, [[80, 150], [550, 150], [550, 175], [80, 175]]), "text": "Chấp hành thời gian, thời lượng sinh hoạt", "left": 80, "right": 550, "top": 150, "bottom": 175, "w": 470, "h": 25, "cy": 162.5},
            {"box": OCRBox("05", 0.95, [[640, 150], [670, 150], [670, 175], [640, 175]]), "text": "05", "left": 640, "right": 670, "top": 150, "bottom": 175, "w": 30, "h": 25, "cy": 162.5},
            {"box": OCRBox("05", 0.95, [[725, 150], [755, 150], [755, 175], [725, 175]]), "text": "05", "left": 725, "right": 755, "top": 150, "bottom": 175, "w": 30, "h": 25, "cy": 162.5},

            # Row 2 (Sub-item with %)
            {"box": OCRBox("% Chi bộ dưới 30 đảng viên ít nhất 90 phút...", 0.95, [[80, 185], [580, 185], [580, 210], [80, 210]]), "text": "% Chi bộ dưới 30 đảng viên ít nhất 90 phút...", "left": 80, "right": 580, "top": 185, "bottom": 210, "w": 500, "h": 25, "cy": 197.5},
            {"box": OCRBox("05", 0.95, [[640, 185], [670, 185], [670, 210], [640, 210]]), "text": "05", "left": 640, "right": 670, "top": 185, "bottom": 210, "w": 30, "h": 25, "cy": 197.5},
            {"box": OCRBox("03", 0.95, [[725, 185], [755, 185], [755, 210], [725, 210]]), "text": "03", "left": 725, "right": 755, "top": 185, "bottom": 210, "w": 30, "h": 25, "cy": 197.5},

            # Row 3 (STT: 2)
            {"box": OCRBox("2", 0.95, [[40, 230], [55, 230], [55, 255], [40, 255]]), "text": "2", "left": 40, "right": 55, "top": 230, "bottom": 255, "w": 15, "h": 25, "cy": 242.5},
            {"box": OCRBox("Tỷ lệ đảng viên dự sinh hoạt chi bộ", 0.95, [[80, 230], [480, 230], [480, 255], [80, 255]]), "text": "Tỷ lệ đảng viên dự sinh hoạt chi bộ", "left": 80, "right": 480, "top": 230, "bottom": 255, "w": 400, "h": 25, "cy": 242.5},
            {"box": OCRBox("05", 0.95, [[640, 230], [670, 230], [670, 255], [640, 255]]), "text": "05", "left": 640, "right": 670, "top": 230, "bottom": 255, "w": 30, "h": 25, "cy": 242.5},
            {"box": OCRBox("05", 0.95, [[725, 230], [755, 230], [755, 255], [725, 255]]), "text": "05", "left": 725, "right": 755, "top": 230, "bottom": 255, "w": 30, "h": 25, "cy": 242.5},
        ]

        table_obj, _, _, _ = detect_and_reconstruct_table(
            box_info, page_min_x=30, page_max_x=840, page_w=810, median_h=25.0
        )
        assert table_obj is not None
        assert table_obj.num_rows == 4
        assert table_obj.num_cols == 5

        # Hàng 0: Header
        assert table_obj.grid[0][0] == "STT"
        assert table_obj.grid[0][1] == "Tiêu chí đánh giá"

        # Hàng 1: STT = 1
        assert table_obj.grid[1][0] == "1"
        assert table_obj.grid[1][1] == "Chấp hành thời gian, thời lượng sinh hoạt"
        assert table_obj.grid[1][2] == "05"
        assert table_obj.grid[1][3] == "05"

        # Hàng 2: Sub-item, STT = rỗng, Tiêu chí khôi phục bullet +
        assert table_obj.grid[2][0] == ""
        assert table_obj.grid[2][1].startswith("+ Chi bộ dưới 30 đảng viên")
        assert table_obj.grid[2][2] == "05"
        assert table_obj.grid[2][3] == "03"

        # Hàng 3: STT = 2 (ĐƯỢC ĐẶT CHÍNH XÁC Ở CỘT 0, KHÔNG BỊ LỆCH SANG CỘT 1)
        assert table_obj.grid[3][0] == "2"
        assert table_obj.grid[3][1] == "Tỷ lệ đảng viên dự sinh hoạt chi bộ"
        assert table_obj.grid[3][2] == "05"
        assert table_obj.grid[3][3] == "05"


class TestImmutabilityAndPrecision:
    """E & F. TÍNH BẤT BIẾN VÀ ĐÁNH GIÁ ĐỘ CHÍNH XÁC (Mục 12.E, 13)."""

    def test_immutability_of_system_text(self):
        """Kiểm tra không bị thay đổi ngoài ý muốn đối với các định dạng kỹ thuật."""
        test_inputs = [
            "C:/Users/pqtru/Documents/ocr_output.txt",
            "https://dangcongsan.vn/sinh-hoat-chi-bo",
            "email_can_bo@ubnd.gov.vn",
            "Nghị định 30/2020/NĐ-CP ngày 05/3/2020",
            "Điểm số: -1.5 điểm và +2.0 điểm",
            "Tỷ lệ: 99.9%",
            "Mã hex: 0xFF00AA | Token: ID-9081",
        ]
        recovered, audit = ContextSymbolRecovery.recover_document_symbols(test_inputs)
        assert recovered == test_inputs
        assert len(audit) == 0  # Không có phép sửa nào bị kích hoạt nhầm

    def test_precision_metric(self):
        """
        Đánh giá chất lượng phép sửa:
        Precision = (Số phép sửa đúng) / (Tổng số phép tự động sửa)
        Kỳ vọng: Precision = 100.0% (Không có phép sửa sai).
        """
        dataset = [
            # Cần sửa đúng (True Positives)
            ("| NỘI DUNG, CHƯƠNG TRÌNH SINH HOẠT CHI BỘ", "I. NỘI DUNG, CHƯƠNG TRÌNH SINH HOẠT CHI BỘ", True),
            ("TI. TIẾN HÀNH SINH HOẠT", "II. TIẾN HÀNH SINH HOẠT", True),
            ("lll. KẾT LUẬN", "III. KẾT LUẬN", True),
            ("lV. TỔ CHỨC THỰC HIỆN", "IV. TỔ CHỨC THỰC HIỆN", True),
            ("% Chi bộ dưới 30 đảng viên", "+ Chi bộ dưới 30 đảng viên", True),
            ("% Thời gian sinh hoạt", "+ Thời gian sinh hoạt", True),
            ("% Từ 90% trở lên", "+ Từ 90% trở lên", True),
            # Cần giữ nguyên (True Negatives)
            ("1. Tuyên bố lý do", "1. Tuyên bố lý do", False),
            ("2. Chủ trì hội nghị", "2. Chủ trì hội nghị", False),
            ("1. NỘP HỒ SƠ", "1. NỘP HỒ SƠ", False),
            ("2. KẾT QUẢ", "2. KẾT QUẢ", False),
            ("Tỷ lệ đạt 90%.", "Tỷ lệ đạt 90%.", False),
            ("50% + 50% = 100%", "50% + 50% = 100%", False),
            ("Số 01-BB/CB", "Số 01-BB/CB", False),
            ("- Lý do: Tổ chức sinh hoạt", "- Lý do: Tổ chức sinh hoạt", False),
        ]

        inputs = [item[0] for item in dataset]
        expected = [item[1] for item in dataset]
        needs_correction = [item[2] for item in dataset]

        recovered, audit = ContextSymbolRecovery.recover_document_symbols(inputs)

        auto_corrected_count = len(audit)
        correctly_corrected_count = 0
        false_positives = 0

        for orig, rec, exp, should_fix in zip(inputs, recovered, expected, needs_correction):
            if should_fix:
                if rec == exp:
                    correctly_corrected_count += 1
                else:
                    print(f"Error on fix: {orig} -> {rec} (exp: {exp})")
            else:
                if rec != orig:
                    false_positives += 1
                    print(f"False positive: {orig} was wrongly modified to {rec}")

        precision = correctly_corrected_count / auto_corrected_count if auto_corrected_count > 0 else 0.0
        print(f"\nCorrection Quality: Total fixes={auto_corrected_count}, Correct={correctly_corrected_count}, False positives={false_positives}, Precision={precision:.2%}")

        assert false_positives == 0, "Không được có bất kỳ phép sửa sai nào trên dữ liệu hợp lệ!"
        assert correctly_corrected_count == 7, "Cả 7 trường hợp cần sửa đều phải sửa đúng!"
        assert precision == 1.0, "Độ chính xác phép sửa phải đạt tuyệt đối 100%!"


class TestStaffTableMultilineMerge:
    """G. KIỂM THỬ GỘP Ô ĐA DÒNG VÀ TRIỆT TIÊU HÀNG RÁC MỒ CÔI (Ảnh 3)."""

    def test_staff_table_split_birthdate_merged(self):
        import numpy as np

        boxes = [
            # Headers
            OCRBox(text="STT", bbox=[[20, 10], [50, 10], [50, 30], [20, 30]], confidence=0.95),
            OCRBox(text="Họ tên", bbox=[[70, 10], [180, 10], [180, 30], [70, 30]], confidence=0.95),
            OCRBox(text="Ngày sinh", bbox=[[200, 10], [280, 10], [280, 30], [200, 30]], confidence=0.95),
            OCRBox(text="Giới tính", bbox=[[300, 10], [360, 10], [360, 30], [300, 30]], confidence=0.95),
            OCRBox(text="Chủ hộ", bbox=[[380, 10], [480, 10], [480, 30], [380, 30]], confidence=0.95),
            OCRBox(text="Địa chỉ", bbox=[[500, 10], [600, 10], [600, 30], [500, 30]], confidence=0.95),
            OCRBox(text="Tên cha", bbox=[[620, 10], [720, 10], [720, 30], [620, 30]], confidence=0.95),
            # Row 1 (01/01/194 bị ngắt đuôi '8')
            OCRBox(text="1", bbox=[[25, 45], [45, 45], [45, 65], [25, 65]], confidence=0.95),
            OCRBox(text="ĐIỀN THỊ LIÊN", bbox=[[75, 45], [175, 45], [175, 65], [75, 65]], confidence=0.95),
            OCRBox(text="01/01/194", bbox=[[205, 45], [275, 45], [275, 65], [205, 65]], confidence=0.95),
            OCRBox(text="8", bbox=[[255, 68], [268, 68], [268, 85], [255, 85]], confidence=0.95),
            OCRBox(text="Nữ", bbox=[[310, 45], [340, 45], [340, 65], [310, 65]], confidence=0.95),
            OCRBox(text="ĐIỀN THỊ LIÊN", bbox=[[385, 45], [475, 45], [475, 65], [385, 65]], confidence=0.95),
            OCRBox(text="TỔ 5, THÔN 2", bbox=[[505, 45], [595, 45], [595, 65], [505, 65]], confidence=0.95),
            OCRBox(text="ĐIỀN HẠNH", bbox=[[625, 45], [715, 45], [715, 65], [625, 65]], confidence=0.95),
            # Row 2 (19/02/198 bị ngắt đuôi '4')
            OCRBox(text="2", bbox=[[25, 95], [45, 95], [45, 115], [25, 115]], confidence=0.95),
            OCRBox(text="NGUYỄN THỊ THUỶ TÂN", bbox=[[75, 95], [185, 95], [185, 115], [75, 115]], confidence=0.95),
            OCRBox(text="19/02/198", bbox=[[205, 95], [275, 95], [275, 115], [205, 115]], confidence=0.95),
            OCRBox(text="4", bbox=[[255, 118], [268, 118], [268, 135], [255, 135]], confidence=0.95),
            OCRBox(text="Nữ", bbox=[[310, 95], [340, 95], [340, 115], [310, 115]], confidence=0.95),
            OCRBox(text="BÙI ĐỨC VĂN", bbox=[[385, 95], [475, 95], [475, 115], [385, 115]], confidence=0.95),
            OCRBox(text="Thôn 2", bbox=[[505, 95], [555, 95], [555, 115], [505, 115]], confidence=0.95),
            OCRBox(text="NGUYỄN VĂN LỘC", bbox=[[625, 95], [725, 95], [725, 115], [625, 115]], confidence=0.95),
        ]

        box_info = []
        for b in boxes:
            l = min(pt[0] for pt in b.bbox)
            r = max(pt[0] for pt in b.bbox)
            t = min(pt[1] for pt in b.bbox)
            bt = max(pt[1] for pt in b.bbox)
            box_info.append({
                "box": b,
                "text": b.text,
                "left": l,
                "right": r,
                "top": t,
                "bottom": bt,
                "w": r - l,
                "h": bt - t,
                "cy": (t + bt) / 2.0,
            })

        page_min_x = min(item["left"] for item in box_info)
        page_max_x = max(item["right"] for item in box_info)
        page_w = page_max_x - page_min_x
        median_h = float(np.median([item["h"] for item in box_info]))

        tbl, pre, rec, post = detect_and_reconstruct_table(box_info, page_min_x, page_max_x, page_w, median_h)

        assert tbl is not None, "Bảng nhân sự phải được phát hiện và tái dựng thành công"
        # Bảng phải có đúng 3 hàng (Header + Row 1 + Row 2), không được sinh ra hàng rác mồ côi | 8 | |
        assert tbl.num_rows == 3, f"Số hàng phải là 3, thực tế: {tbl.num_rows}"
        assert tbl.num_cols == 7, f"Số cột phải là 7, thực tế: {tbl.num_cols}"

        row1 = tbl.grid[1]
        assert row1[0] == "1"
        assert row1[1] == "ĐIỀN THỊ LIÊN"
        assert row1[2] == "01/01/1948", f"Ngày sinh hàng 1 phải được ghép liền thành '01/01/1948', thực tế: '{row1[2]}'"
        assert row1[3] == "Nữ"
        assert row1[4] == "ĐIỀN THỊ LIÊN"
        assert row1[5] == "TỔ 5, THÔN 2"
        assert row1[6] == "ĐIỀN HẠNH"

        row2 = tbl.grid[2]
        assert row2[0] == "2"
        assert row2[1] == "NGUYỄN THỊ THUỶ TÂN"
        assert row2[2] == "19/02/1984", f"Ngày sinh hàng 2 phải được ghép liền thành '19/02/1984', thực tế: '{row2[2]}'"


class TestIconTableSymbolHandling:
    """H. KIỂM THỬ BẢNG BIỂU TƯỢNG VÀ MÃ KÝ HIỆU (Ảnh 4)."""

    def test_icon_column_cleanup(self):
        import numpy as np

        boxes = [
            # Header
            OCRBox(text="Symb", bbox=[[10, 10], [50, 10], [50, 30], [10, 30]], confidence=0.95),
            OCRBox(text="Code", bbox=[[60, 10], [120, 10], [120, 30], [60, 30]], confidence=0.95),
            OCRBox(text="Symb", bbox=[[140, 10], [180, 10], [180, 30], [140, 30]], confidence=0.95),
            OCRBox(text="Code", bbox=[[190, 10], [250, 10], [250, 30], [190, 30]], confidence=0.95),
            # Row 1 (cy -> [Icon], Code: 14; ♫ -> giữ nguyên, Code: 15)
            OCRBox(text="cy", bbox=[[15, 40], [45, 40], [45, 60], [15, 60]], confidence=0.95),
            OCRBox(text="14", bbox=[[65, 40], [110, 40], [110, 60], [65, 60]], confidence=0.95),
            OCRBox(text="♫", bbox=[[145, 40], [175, 40], [175, 60], [145, 60]], confidence=0.95),
            OCRBox(text="15", bbox=[[195, 40], [240, 40], [240, 60], [195, 60]], confidence=0.95),
            # Row 2 (cos -> [Icon], Code: 21; ® -> giữ nguyên, Code: 0174)
            OCRBox(text="cos", bbox=[[15, 70], [45, 70], [45, 90], [15, 90]], confidence=0.95),
            OCRBox(text="21", bbox=[[65, 70], [110, 70], [110, 90], [65, 90]], confidence=0.95),
            OCRBox(text="®", bbox=[[145, 70], [175, 70], [175, 90], [145, 90]], confidence=0.95),
            OCRBox(text="0174", bbox=[[195, 70], [240, 70], [240, 90], [195, 90]], confidence=0.95),
            # Row 3 (TT -> [Icon], Code: 19)
            OCRBox(text="TT", bbox=[[15, 100], [45, 100], [45, 120], [15, 120]], confidence=0.95),
            OCRBox(text="19", bbox=[[65, 100], [110, 100], [110, 120], [65, 120]], confidence=0.95),
            OCRBox(text="±", bbox=[[145, 100], [175, 100], [175, 120], [145, 120]], confidence=0.95),
            OCRBox(text="0177", bbox=[[195, 100], [240, 100], [240, 120], [195, 120]], confidence=0.95),
        ]

        box_info = []
        for b in boxes:
            l = min(pt[0] for pt in b.bbox)
            r = max(pt[0] for pt in b.bbox)
            t = min(pt[1] for pt in b.bbox)
            bt = max(pt[1] for pt in b.bbox)
            box_info.append({
                "box": b,
                "text": b.text,
                "left": l,
                "right": r,
                "top": t,
                "bottom": bt,
                "w": r - l,
                "h": bt - t,
                "cy": (t + bt) / 2.0,
            })

        tbl, _, _, _ = detect_and_reconstruct_table(box_info, 10, 250, 1000, 20.0)
        assert tbl is not None, "Bảng biểu tượng phải được tái dựng thành công"

        # Hàng 1
        assert tbl.grid[1][0] == "[Icon]", f"Chuỗi rác 'cy' phải được thay bằng '[Icon]', thực tế: '{tbl.grid[1][0]}'"
        assert tbl.grid[1][1] == "14"
        assert tbl.grid[1][2] == "♫", f"Symbol chuẩn '♫' phải được giữ nguyên, thực tế: '{tbl.grid[1][2]}'"
        assert tbl.grid[1][3] == "15"

        # Hàng 2
        assert tbl.grid[2][0] == "[Icon]", f"Chuỗi rác 'cos' phải được thay bằng '[Icon]', thực tế: '{tbl.grid[2][0]}'"
        assert tbl.grid[2][1] == "21"
        assert tbl.grid[2][2] == "®", f"Symbol chuẩn '®' phải được giữ nguyên, thực tế: '{tbl.grid[2][2]}'"
        assert tbl.grid[2][3] == "0174"

        # Hàng 3
        assert tbl.grid[3][0] == "[Icon]", f"Chuỗi rác 'TT' phải được thay bằng '[Icon]', thực tế: '{tbl.grid[3][0]}'"
        assert tbl.grid[3][1] == "19"
        assert tbl.grid[3][2] == "±", f"Symbol chuẩn '±' phải được giữ nguyên, thực tế: '{tbl.grid[3][2]}'"
        assert tbl.grid[3][3] == "0177"


class TestPhase2SymbolAndTableCases:
    """I. KIỂM THỬ TOÀN DIỆN PHASE 2: SYMBOL & CỘT STT."""

    def test_sub_criterion_bullets_and_pipe_removal(self):
        """Khôi phục dấu cộng/trừ đầu dòng và loại bỏ pipe thừa."""
        inputs = [
            "| % Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút",
            "| * Chi bộ dưới 30 đảng viên ít nhất 90 phút; chi bộ trên 30 đảng viên ít nhất 120 phút",
            "| + Chi bộ thực hiện tốt nhiệm vụ",
            "| - Nội dung sinh hoạt chi bộ theo hướng dẫn",
            "+ Chi bộ đoàn kết",
        ]
        recovered, audit = ContextSymbolRecovery.recover_document_symbols(inputs)
        assert recovered[0] == "+ Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút"
        assert recovered[1] == "+ Chi bộ dưới 30 đảng viên ít nhất 90 phút; chi bộ trên 30 đảng viên ít nhất 120 phút"
        assert recovered[2] == "+ Chi bộ thực hiện tốt nhiệm vụ"
        assert recovered[3] == "- Nội dung sinh hoạt chi bộ theo hướng dẫn"
        assert recovered[4] == "+ Chi bộ đoàn kết"

    def test_stt_double_pipe_cleanup(self):
        """Xử lý dấu '| |' thừa ở cột STT và bảo toàn số thứ tự cùng nội dung."""
        inputs = [
            "1 | | Chấp hành thời gian, thời lượng sinh hoạt | 05 | 05",
            "2 | | Tỷ lệ đảng viên dự sinh hoạt chi bộ | 05 | 05",
        ]
        recovered, audit = ContextSymbolRecovery.recover_document_symbols(inputs)
        assert recovered[0] == "1 | Chấp hành thời gian, thời lượng sinh hoạt | 05 | 05"
        assert recovered[1] == "2 | Tỷ lệ đảng viên dự sinh hoạt chi bộ | 05 | 05"
        assert len(audit) == 2
        assert audit[0]["rule"] == "STT_DOUBLE_PIPE_CLEANUP"

    def test_preserve_percentages_and_valid_symbols(self):
        """Bảo toàn các giá trị phần trăm, phép tính, URL, mã hex hợp lệ."""
        inputs = [
            "Tỷ lệ đạt 90% theo quy định",
            "Từ 75% đến dưới 90%",
            "Dưới 60% hoặc có đảng viên vắng không có lý do",
            "50% + 50% = 100%",
            "a + b = c",
            "* Ghi chú: Đây là chú thích quan trọng.",
            "Mã hiệu: 0xFF | ID: 100",
        ]
        recovered, audit = ContextSymbolRecovery.recover_document_symbols(inputs)
        assert recovered == inputs
        assert len(audit) == 0


class TestPhase3TwoColumnHeaderAndDecorativeLine:
    """J. KIỂM THỬ TOÀN DIỆN PHASE 3: HEADER HAI CỘT & LOẠI BỎ ĐƯỜNG KẺ TRANG TRÍ."""

    def test_parallel_header_decorative_percent_removed(self):
        """
        Kiểm tra Header hai cột:
        - Khối trái có đường kẻ ngang trang trí bị nhận nhầm thành '%'.
        - Kết quả: '%' bị loại bỏ hoàn toàn, khối trái đọc trọn vẹn 4 dòng rồi đến khối phải 2 dòng.
        """
        from app.core.ocr_engine import detect_parallel_administrative_header

        box_info = [
            # Khối trái
            {"box": OCRBox("ĐẢNG ỦY XÃ PHƯỚC HẢI", 0.98, [[80, 100], [260, 100], [260, 125], [80, 125]]), "text": "ĐẢNG ỦY XÃ PHƯỚC HẢI", "left": 80, "right": 260, "top": 100, "bottom": 125, "w": 180, "h": 25, "cy": 112.5},
            {"box": OCRBox("CHI BỘ VĂN PHÒNG HĐND", 0.97, [[70, 130], [270, 130], [270, 150], [70, 150]]), "text": "CHI BỘ VĂN PHÒNG HĐND", "left": 70, "right": 270, "top": 130, "bottom": 150, "w": 200, "h": 20, "cy": 140.0},
            {"box": OCRBox("VÀ UBND XÃ", 0.96, [[100, 155], [210, 155], [210, 175], [100, 175]]), "text": "VÀ UBND XÃ", "left": 100, "right": 210, "top": 155, "bottom": 175, "w": 110, "h": 20, "cy": 165.0},
            {"box": OCRBox("%", 0.95, [[145, 180], [165, 180], [165, 195], [145, 195]]), "text": "%", "left": 145, "right": 165, "top": 180, "bottom": 195, "w": 20, "h": 15, "cy": 187.5},
            {"box": OCRBox("Số 01-BB/CB", 0.97, [[100, 200], [200, 200], [200, 225], [100, 225]]), "text": "Số 01-BB/CB", "left": 100, "right": 200, "top": 200, "bottom": 225, "w": 100, "h": 25, "cy": 212.5},
            # Khối phải
            {"box": OCRBox("ĐẢNG CỘNG SẢN VIỆT NAM", 0.99, [[380, 100], [580, 100], [580, 125], [380, 125]]), "text": "ĐẢNG CỘNG SẢN VIỆT NAM", "left": 380, "right": 580, "top": 100, "bottom": 125, "w": 200, "h": 25, "cy": 112.5},
            {"box": OCRBox("Phước Hải, ngày 05 tháng 10 năm 2026", 0.98, [[350, 150], [580, 150], [580, 175], [350, 175]]), "text": "Phước Hải, ngày 05 tháng 10 năm 2026", "left": 350, "right": 580, "top": 150, "bottom": 175, "w": 230, "h": 25, "cy": 162.5},
            # Tiêu đề giữa trang
            {"box": OCRBox("BIÊN BẢN", 0.99, [[250, 250], [370, 250], [370, 280], [250, 280]]), "text": "BIÊN BẢN", "left": 250, "right": 370, "top": 250, "bottom": 280, "w": 120, "h": 30, "cy": 265.0},
        ]

        header_lines, rem = detect_parallel_administrative_header(
            box_info, page_min_x=70, page_max_x=580, page_w=510, gutter_thresh=20.0, median_h=22.0
        )

        header_line_texts = [" ".join(b.text for b in line) for line in header_lines]

        # Đường kẻ '%' phải bị loại bỏ, khối trái gồm đúng 4 dòng:
        assert len(header_line_texts) == 6
        assert header_line_texts[0] == "ĐẢNG ỦY XÃ PHƯỚC HẢI"
        assert header_line_texts[1] == "CHI BỘ VĂN PHÒNG HĐND"
        assert header_line_texts[2] == "VÀ UBND XÃ"
        assert header_line_texts[3] == "Số 01-BB/CB"
        # Sau đó là Khối phải gồm 2 dòng:
        assert header_line_texts[4] == "ĐẢNG CỘNG SẢN VIỆT NAM"
        assert header_line_texts[5] == "Phước Hải, ngày 05 tháng 10 năm 2026"


class TestPhase4SpellingAndTokenization:
    """K. KIỂM THỬ TOÀN DIỆN PHASE 4: SPELL CHECKING CHO 'Vdng MẶT'."""

    def test_vdng_mat_detected_with_exact_span(self):
        """Phát hiện lỗi 'Vdng' với span chính xác, không dính dấu '-' và gợi ý 'Vắng'."""
        from app.core.spell_corrector import check_text_spelling, find_spelling_errors

        sample_text = "- Vdng mặt: 0 đồng chí."

        # Kiểm tra qua cả 2 API
        errors_1 = check_text_spelling(sample_text)
        errors_2 = find_spelling_errors(sample_text)

        assert len(errors_1) == 1
        assert len(errors_2) == 1

        err = errors_1[0]
        assert err.token == "Vdng"
        assert err.start_idx == 2
        assert err.end_idx == 6
        assert sample_text[err.start_idx:err.end_idx] == "Vdng"
        assert "Vắng" in err.suggestions

    def test_correct_phrases_not_flagged(self):
        """Các từ/cụm từ đúng và whitelist ALL CAPS không bị cảnh báo nhầm."""
        from app.core.spell_corrector import check_text_spelling

        correct_samples = [
            "- Vắng mặt: 0 đồng chí.",
            "Có mặt: 14 đồng chí.",
            "UBND, HĐND, TW, BB, CB, DOCX, PDF",
            "Đảng ủy xã Phước Hải",
        ]
        for s in correct_samples:
            errors = check_text_spelling(s)
            assert len(errors) == 0, f"Chuỗi hợp lệ không được báo lỗi: '{s}' -> {errors}"


class TestPhase5TableCleanupAndClippedText:
    """L. KIỂM THỬ TOÀN DIỆN PHASE 5: LÀM SẠCH BẢNG ĐIỂM, STT 1, DẤU | HÀNG CON VÀ VĂN BẢN BỊ CẮT."""

    def test_clipped_text_administrative_recovery(self):
        """Khôi phục cụm từ bị cắt 'nhiệm vụ đặc b' -> 'nhiệm vụ đặc biệt'."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        clipped = "* Trừ trường hợp được miễn công tác, sinh hoạt đảng theo quy định và đảng viên trong lực lượng vũ trang thực hiện nhiệm vụ đặc b"
        restored = VietnameseTextPostProcessor.clean_and_restore(clipped)
        assert restored.endswith("thực hiện nhiệm vụ đặc biệt")

        # Trường hợp âm tính: từ khác chứa b không bị thay thế tùy tiện
        negative = "Học sinh lớp 12 b tham gia học tập"
        assert VietnameseTextPostProcessor.clean_and_restore(negative) == negative

    def test_table_reconstruction_stt_1_restored_and_sub_rows_piped(self):
        """
        Kiểm tra tái dựng bảng điểm:
        - Hàng 1 có ô STT bị thiếu/nhận nhầm nhưng hàng sau có STT '2' -> Tự động khôi phục STT '1'.
        - Hàng con bắt đầu bằng '+' có các cột điểm được phân cách bằng ' | ' thay vì dính space vào text.
        - Không có dấu '|' thừa lặp trong cột STT.
        """
        from app.core.ocr_engine import detect_and_reconstruct_table, OCRBox

        # Giả lập box_info cho bảng điểm
        # Cột STT: left ~ 20..60
        # Cột Tiêu chí: left ~ 80..500
        # Cột Điểm tối đa: left ~ 540..580
        # Cột Điểm tự chấm: left ~ 620..660
        box_info = [
            # Header
            {"box": OCRBox("STT", 0.99, [[20, 50], [60, 50], [60, 75], [20, 75]]), "text": "STT", "left": 20, "right": 60, "top": 50, "bottom": 75, "w": 40, "h": 25, "cy": 62.5},
            {"box": OCRBox("Tiêu chí đánh giá", 0.99, [[80, 50], [300, 50], [300, 75], [80, 75]]), "text": "Tiêu chí đánh giá", "left": 80, "right": 300, "top": 50, "bottom": 75, "w": 220, "h": 25, "cy": 62.5},
            {"box": OCRBox("Điểm tối đa", 0.98, [[540, 50], [580, 50], [580, 75], [540, 75]]), "text": "Điểm tối đa", "left": 540, "right": 580, "top": 50, "bottom": 75, "w": 40, "h": 25, "cy": 62.5},
            {"box": OCRBox("Điểm tự chấm", 0.98, [[620, 50], [660, 50], [660, 75], [620, 75]]), "text": "Điểm tự chấm", "left": 620, "right": 660, "top": 50, "bottom": 75, "w": 40, "h": 25, "cy": 62.5},

            # Hàng 1 (STT bị VietOCR nhận thành '|', Tiêu chí chính 1)
            {"box": OCRBox("|", 0.80, [[35, 85], [45, 85], [45, 105], [35, 105]]), "text": "|", "left": 35, "right": 45, "top": 85, "bottom": 105, "w": 10, "h": 20, "cy": 95.0},
            {"box": OCRBox("Chấp hành thời gian, thời lượng sinh hoạt", 0.98, [[80, 85], [450, 85], [450, 105], [80, 105]]), "text": "Chấp hành thời gian, thời lượng sinh hoạt", "left": 80, "right": 450, "top": 85, "bottom": 105, "w": 370, "h": 20, "cy": 95.0},
            {"box": OCRBox("05", 0.97, [[545, 85], [575, 85], [575, 105], [545, 105]]), "text": "05", "left": 545, "right": 575, "top": 85, "bottom": 105, "w": 30, "h": 20, "cy": 95.0},
            {"box": OCRBox("05", 0.97, [[625, 85], [655, 85], [655, 105], [625, 105]]), "text": "05", "left": 625, "right": 655, "top": 85, "bottom": 105, "w": 30, "h": 20, "cy": 95.0},

            # Hàng 2 (Tiêu chí con, không có STT)
            {"box": OCRBox("+ Chi bộ dưới 30 đảng viên ít nhất 90 phút", 0.97, [[80, 115], [480, 115], [480, 135], [80, 135]]), "text": "+ Chi bộ dưới 30 đảng viên ít nhất 90 phút", "left": 80, "right": 480, "top": 115, "bottom": 135, "w": 400, "h": 20, "cy": 125.0},
            {"box": OCRBox("05", 0.96, [[545, 115], [575, 115], [575, 135], [545, 135]]), "text": "05", "left": 545, "right": 575, "top": 115, "bottom": 135, "w": 30, "h": 20, "cy": 125.0},
            {"box": OCRBox("05", 0.96, [[625, 115], [655, 115], [655, 135], [625, 135]]), "text": "05", "left": 625, "right": 655, "top": 115, "bottom": 135, "w": 30, "h": 20, "cy": 125.0},

            # Hàng 3 (Tiêu chí chính 2 có STT '2')
            {"box": OCRBox("2", 0.98, [[35, 145], [45, 145], [45, 165], [35, 165]]), "text": "2", "left": 35, "right": 45, "top": 145, "bottom": 165, "w": 10, "h": 20, "cy": 155.0},
            {"box": OCRBox("Tỷ lệ đảng viên dự sinh hoạt chi bộ", 0.98, [[80, 145], [420, 145], [420, 165], [80, 165]]), "text": "Tỷ lệ đảng viên dự sinh hoạt chi bộ", "left": 80, "right": 420, "top": 145, "bottom": 165, "w": 340, "h": 20, "cy": 155.0},
            {"box": OCRBox("05", 0.97, [[545, 145], [575, 145], [575, 165], [545, 165]]), "text": "05", "left": 545, "right": 575, "top": 145, "bottom": 165, "w": 30, "h": 20, "cy": 155.0},
            {"box": OCRBox("05", 0.97, [[625, 145], [655, 145], [655, 165], [625, 165]]), "text": "05", "left": 625, "right": 655, "top": 145, "bottom": 165, "w": 30, "h": 20, "cy": 155.0},

            # Hàng 4 (Tiêu chí con của mục 2)
            {"box": OCRBox("+ Từ 90% trở lên", 0.98, [[80, 175], [260, 175], [260, 195], [80, 195]]), "text": "+ Từ 90% trở lên", "left": 80, "right": 260, "top": 175, "bottom": 195, "w": 180, "h": 20, "cy": 185.0},
            {"box": OCRBox("04", 0.97, [[545, 175], [575, 175], [575, 195], [545, 195]]), "text": "04", "left": 545, "right": 575, "top": 175, "bottom": 195, "w": 30, "h": 20, "cy": 185.0},
            {"box": OCRBox("04", 0.97, [[625, 175], [655, 175], [655, 195], [625, 195]]), "text": "04", "left": 625, "right": 655, "top": 175, "bottom": 195, "w": 30, "h": 20, "cy": 185.0},
        ]

        tbl, pre_b, tbl_rows, post_b = detect_and_reconstruct_table(
            box_info, page_min_x=20, page_max_x=660, page_w=640, median_h=20.0
        )

        assert tbl is not None
        assert len(tbl.grid) == 5
        # STT hàng 1 phải được khôi phục thành '1' (thay vì '|')
        assert tbl.grid[1][0] == "1"
        assert tbl.grid[1][1] == "Chấp hành thời gian, thời lượng sinh hoạt"
        # STT hàng 3 là '2'
        assert tbl.grid[3][0] == "2"
        # Tiêu chí con hàng 4 bảo toàn '90%'
        assert "90%" in tbl.grid[4][1]


class TestPhase6DocxExportParallelHeader:
    """M. KIỂM THỬ XUẤT DOCX CHO PARALLEL ADMINISTRATIVE HEADER."""

    def test_export_docx_with_parallel_header(self, tmp_path):
        """Xác nhận export_to_docx tạo bảng 2 cột không viền cho header và không duplicate text."""
        import docx
        from app.core.ocr_engine import ParallelHeader, OCRResult, export_to_docx

        hdr = ParallelHeader(
            left_lines=[
                "ĐẢNG ỦY XÃ PHƯỚC HẢI",
                "CHI BỘ VĂN PHÒNG HĐND VÀ UBND XÃ",
                "Số 01-BB/CB",
            ],
            right_lines=[
                "ĐẢNG CỘNG SẢN VIỆT NAM",
                "Phước Hải, ngày 05 tháng 10 năm 2026",
            ],
        )

        full_text = (
            "ĐẢNG ỦY XÃ PHƯỚC HẢI\n"
            "CHI BỘ VĂN PHÒNG HĐND VÀ UBND XÃ\n"
            "Số 01-BB/CB\n"
            "ĐẢNG CỘNG SẢN VIỆT NAM\n"
            "Phước Hải, ngày 05 tháng 10 năm 2026\n"
            "BIÊN BẢN\n"
            "sinh hoạt chi bộ định kỳ tháng 10 năm 2026\n"
            "Vào lúc 15 giờ, ngày 05 tháng 10 năm 2026..."
        )

        ocr_res = OCRResult(
            boxes=[],
            full_text=full_text,
            language="vi",
            parallel_header=hdr,
        )

        out_docx = str(tmp_path / "test_parallel_header.docx")
        success = export_to_docx(ocr_res, out_docx)
        assert success is True

        doc = docx.Document(out_docx)
        # Bảng đầu tiên phải là bảng 2 cột của Parallel Header
        assert len(doc.tables) >= 1
        hdr_tbl = doc.tables[0]
        assert len(hdr_tbl.rows) == 1
        assert len(hdr_tbl.columns) == 2

        # Cột trái chứa tên cơ quan, chi bộ, số hiệu
        left_text = hdr_tbl.cell(0, 0).text
        assert "ĐẢNG ỦY XÃ PHƯỚC HẢI" in left_text
        assert "Số 01-BB/CB" in left_text

        # Cột phải chứa Quốc hiệu, ngày tháng
        right_text = hdr_tbl.cell(0, 1).text
        assert "ĐẢNG CỘNG SẢN VIỆT NAM" in right_text
        assert "Phước Hải, ngày 05 tháng 10 năm 2026" in right_text

        # Các đoạn văn bản body bên ngoài không được lặp lại dòng header
        body_paras_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        assert "ĐẢNG ỦY XÃ PHƯỚC HẢI" not in body_paras_text
        assert "ĐẢNG CỘNG SẢN VIỆT NAM" not in body_paras_text
        assert "BIÊN BẢN" in body_paras_text


class TestPhase7AdvancedVisionAndTableReconstruction:
    """G. KIỂM THỬ KHÔI PHỤC NÂNG CAO CHO OCR & TÁI DỰNG BẢNG THỰC TẾ."""

    @pytest.mark.parametrize(
        "input_text,expected",
        [
            ("L NỘI DUNG, CHƯƠNG TRÌNH SINH HOẠT CHI BỘ", "I. NỘI DUNG, CHƯƠNG TRÌNH SINH HOẠT CHI BỘ"),
            ("L. TIẾN HÀNH SINH HOẠT", "I. TIẾN HÀNH SINH HOẠT"),
            ("L. KẾ HOẠCH NĂM 2026", "I. KẾ HOẠCH NĂM 2026"),
        ],
    )
    def test_roman_i_from_l(self, input_text, expected):
        """Khôi phục số La Mã I. từ chữ cái L hoặc L. khi đi cùng tiêu đề in hoa."""
        rec, dec, _, rule = ContextSymbolRecovery.recover_roman_numeral_in_line(input_text)
        assert rec == expected
        assert dec == SymbolDecision.AUTO_CORRECT

    @pytest.mark.parametrize(
        "negative_l",
        [
            "Lãnh đạo đơn vị phát biểu",
            "Lớp bồi dưỡng nghiệp vụ",
            "Lê Văn Điền thông qua báo cáo",
        ],
    )
    def test_roman_i_negative_l(self, negative_l):
        """Bảo vệ tuyệt đối các từ tiếng Việt bắt đầu bằng L không bị đổi thành số La Mã."""
        rec, dec, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line(negative_l)
        assert rec == negative_l
        assert dec == SymbolDecision.PRESERVE

    @pytest.mark.parametrize(
        "input_text,expected",
        [
            ("%Không có đảng viên vắng mặt", "+ Không có đảng viên vắng mặt"),
            ("4 Chỉ bộ dưới 30 đảng viên", "+ Chỉ bộ dưới 30 đảng viên"),
            ("7 Thời gian sinh hoạt chi bộ", "+ Thời gian sinh hoạt chi bộ"),
            ("1 Trừ trường hợp được miễn công tác", "* Trừ trường hợp được miễn công tác"),
            ("? Trừ trường hợp được miễn công tác", "* Trừ trường hợp được miễn công tác"),
            ("- Trừ trường hợp được miễn công tác", "* Trừ trường hợp được miễn công tác"),
        ],
    )
    def test_bullet_and_footnote_recovery(self, input_text, expected):
        """Khôi phục chính xác bullet %, 4, 7 và footnote * cho các dòng tiêu chí."""
        rec, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(input_text)
        assert rec == expected
        assert dec == SymbolDecision.AUTO_CORRECT

    @pytest.mark.parametrize(
        "input_text,expected",
        [
            (
                "Thời gian sinh hoạt chi bộ chỉ bằng % so với quy định",
                "Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định",
            ),
            (
                "Thời gian sinh hoạt chi bộ dưới % so với quy định",
                "Thời gian sinh hoạt chi bộ dưới ½ so với quy định",
            ),
            (
                "ĐẰNG CỘNG SẢN VIỆT NAM",
                "ĐẢNG CỘNG SẢN VIỆT NAM",
            ),
            (
                "ngày 28 thắng 10 năm 2024",
                "ngày 28 tháng 10 năm 2024",
            ),
            (
                "Ghi chủ",
                "Ghi chú",
            ),
            (
                "it nhất 90 phút",
                "ít nhất 90 phút",
            ),
            (
                "Trừ trường hợp thực hiện nhiệm vụ đặc",
                "Trừ trường hợp thực hiện nhiệm vụ đặc biệt",
            ),
            (
                "sinh hoạt đằng",
                "sinh hoạt đảng",
            ),
            (
                "đảng viễn",
                "đảng viên",
            ),
            (
                "CHI BỘ VĂN PHÒNG HĐND",
                "CHI BỘ VĂN PHÒNG HĐND",
            ),
        ],
    )
    def test_post_processor_text_cleanups(self, input_text, expected):
        """Chuẩn hóa phân số ½ và các thuật ngữ hành chính phổ biến."""
        cleaned = VietnameseTextPostProcessor.clean_and_restore(input_text)
        assert cleaned == expected

    def test_table_reconstruction_missing_stt_and_mutual_exclusivity(self):
        """Kiểm tra khôi phục STT 1, 2 và tính phân hoạch độc quyền của bảng."""
        boxes = [
            # Header ngoài bảng
            OCRBox("BẢNG ĐIỂM ĐÁNH GIÁ", 0.98, [[100, 30], [500, 30], [500, 55], [100, 55]]),
            # Header bảng
            OCRBox("STT", 0.95, [[30, 100], [70, 100], [70, 125], [30, 125]]),
            OCRBox("Nội dung tiêu chí", 0.98, [[100, 100], [450, 100], [450, 125], [100, 125]]),
            OCRBox("Điểm chuẩn", 0.95, [[600, 100], [700, 100], [700, 125], [600, 125]]),
            # Hàng 1 (bị mất STT 1)
            OCRBox("Chấp hành thời gian sinh hoạt", 0.98, [[100, 140], [500, 140], [500, 165], [100, 165]]),
            OCRBox("10", 0.95, [[620, 140], [660, 140], [660, 165], [620, 165]]),
            # Hàng 2 (tiêu chí con +)
            OCRBox("+ Đúng giờ", 0.98, [[100, 180], [300, 180], [300, 205], [100, 205]]),
            OCRBox("05", 0.95, [[620, 180], [660, 180], [660, 205], [620, 205]]),
            # Hàng 3 (Hàng 2 chính, STT bị dính chữ '2 ANNASTRY')
            OCRBox("2 ANNASTRY", 0.95, [[30, 220], [80, 220], [80, 245], [30, 245]]),
            OCRBox("Tỷ lệ đảng viên dự sinh hoạt", 0.98, [[100, 220], [500, 220], [500, 245], [100, 245]]),
            OCRBox("10", 0.95, [[620, 220], [660, 220], [660, 245], [620, 245]]),
        ]

        lines, tables = reconstruct_document_layout(boxes)
        assert len(tables) == 1
        tbl = tables[0]

        # Kiểm tra STT hàng 1 được khôi phục thành '1'
        assert tbl.grid[1][0] == "1"
        assert tbl.grid[1][1] == "Chấp hành thời gian sinh hoạt"

        # Kiểm tra hàng con không có STT số
        assert tbl.grid[2][0] == ""
        assert "+ Đúng giờ" in tbl.grid[2][1]

        # Kiểm tra STT hàng 3 được chuẩn hóa thành '2'
        assert tbl.grid[3][0] == "2"
        assert tbl.grid[3][1] == "Tỷ lệ đảng viên dự sinh hoạt"


class TestEmailDomainAndInternshipProgressRecovery:
    """Kiểm thử phục hồi Ký tự Email (@), Dấu gạch ngang (-) và Bảng tiến độ thực tập."""

    def test_recover_emails_and_urls_at_symbol(self):
        """Khôi phục ký tự @ trong Email bị nhận nhầm thành Q, O, (a, (D và sửa amaron.com."""
        from app.core.ocr_engine import recover_emails_and_urls

        assert recover_emails_and_urls("vpt.luan Qhutech.edu.vn") == "vpt.luan@hutech.edu.vn"
        assert recover_emails_and_urls("hunggia (Damaron.com") == "hunggia@amazon.com"
        assert recover_emails_and_urls("hunggia (Damazon.com") == "hunggia@amazon.com"
        assert recover_emails_and_urls("contact @ hutech.edu.vn") == "contact@hutech.edu.vn"
        assert recover_emails_and_urls("info (a hutech.edu.vn") == "info@hutech.edu.vn"

    def test_recover_dashes_and_hyphens(self):
        """Khôi phục dấu gạch ngang giữa ngày tháng và trong cụm họ tên – email."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        raw1 = "2. Cán bộ hướng dẫn tại đơn vị thực tập (họ tên? email): Nguyễn Gia Hưng. hunggia (Damaron.com"
        cleaned1 = VietnameseTextPostProcessor.clean_and_restore(raw1)
        assert "họ tên – email" in cleaned1
        assert "Nguyễn Gia Hưng – hunggia@amazon.com" in cleaned1 or "Nguyễn Gia Hưng - hunggia@amazon.com" in cleaned1

        raw2 = "3. Giảng viên theo dõi/ hướng dẫn thực tập tại trường (họ tên? email): Võ Phạm Thành Luân ? vpt.luan Qhutech.edu.vn"
        cleaned2 = VietnameseTextPostProcessor.clean_and_restore(raw2)
        assert "họ tên – email" in cleaned2
        assert "vpt.luan@hutech.edu.vn" in cleaned2

        raw_dates = "20/04/2026 ? 26/04/2026 và 27/04/2026 03/05/2026"
        cleaned_dates = VietnameseTextPostProcessor.clean_and_restore(raw_dates)
        assert "20/04/2026 - 26/04/2026" in cleaned_dates
        assert "27/04/2026 - 03/05/2026" in cleaned_dates

    def test_internship_progress_table_formatting(self):
        """Cố định cấu trúc bảng tiến độ thực tập 3 cột: Tuần | Ngày | Nội dung."""
        from app.core.ocr_engine import format_progress_table_row

        row1 = "1 20/04/2026 26/04/2026 | - Họp nhóm, tìm hiểu về AWS."
        formatted1 = format_progress_table_row(row1)
        assert formatted1 == "1 | 20/04/2026 - 26/04/2026 | - Họp nhóm, tìm hiểu về AWS."

        row2 = "2 | 27/04/2026 03/05/2026 - Đã tạo và cấu hình AWS Free Tier account."
        formatted2 = format_progress_table_row(row2)
        assert formatted2 == "2 | 27/04/2026 - 03/05/2026 | - Đã tạo và cấu hình AWS Free Tier account."

    def test_spell_checker_ignores_emails_and_tech_terms(self):
        """Spell checker bảo vệ 100% email, url và thuật ngữ kỹ thuật, không bôi đỏ."""
        from app.core.spell_corrector import check_text_spelling

        text = (
            "Cán bộ hướng dẫn: Nguyễn Gia Hưng – hunggia@amazon.com\n"
            "Giảng viên hướng dẫn: Võ Phạm Thành Luân – vpt.luan@hutech.edu.vn\n"
            "Cấu hình AWS Free Tier account và làm quen với AWS Management Console."
        )
        errors = check_text_spelling(text)
        error_tokens = [err.token for err in errors]

        for safe_word in ["hunggia", "amazon", "vpt", "luan", "hutech", "edu", "aws", "tier", "console", "account"]:
            assert safe_word not in error_tokens, f"Từ an toàn '{safe_word}' không được bôi đỏ trong spell checker!"


class TestSection8MasterRequirements:
    """Bộ kiểm thử bắt buộc theo Section 8 của Master Prompt."""

    def test_01_table_header_standalone(self):
        """1. Hàng tiêu đề bảng (Tuần | Ngày | Nội dung) đứng độc lập."""
        boxes = [
            # Header bảng
            OCRBox("Tuần", 0.98, [[30, 50], [80, 50], [80, 75], [30, 75]]),
            OCRBox("Ngày", 0.98, [[100, 50], [200, 50], [200, 75], [100, 75]]),
            OCRBox("Nội dung", 0.98, [[220, 50], [350, 50], [350, 75], [220, 75]]),
            # Bounding box dữ liệu đầu tiên (dù tọa độ Y gần nhau cũng không được gộp vào header)
            OCRBox("- Họp nhóm, tìm hiểu về AWS.", 0.95, [[220, 78], [600, 78], [600, 100], [220, 100]]),
            # Tuần 1
            OCRBox("1", 0.95, [[40, 105], [60, 105], [60, 130], [40, 130]]),
            OCRBox("20/04/2026 - 26/04/2026", 0.98, [[100, 105], [215, 105], [215, 130], [100, 130]]),
            OCRBox("- Đọc các nội quy.", 0.95, [[220, 105], [500, 105], [500, 130], [220, 130]]),
            # Tuần 2
            OCRBox("2", 0.95, [[40, 160], [60, 160], [60, 185], [40, 185]]),
            OCRBox("27/04/2026 - 03/05/2026", 0.98, [[100, 160], [215, 160], [215, 185], [100, 185]]),
            OCRBox("- Thực tập AWS.", 0.95, [[220, 160], [500, 160], [500, 185], [220, 185]]),
        ]
        lines, tables = reconstruct_document_layout(boxes)
        assert len(tables) == 1
        tbl = tables[0]
        header = tbl.grid[0]
        assert header[0] == "Tuần"
        assert header[1] == "Ngày"
        assert header[2] == "Nội dung"

    def test_02_multiline_internship_week_grouping(self):
        """2. Nội dung nhiều dòng của Tuần 1, Tuần 2 thuộc đúng tuần tương ứng."""
        boxes = [
            # Header
            OCRBox("Tuần", 0.98, [[30, 50], [80, 50], [80, 75], [30, 75]]),
            OCRBox("Ngày", 0.98, [[100, 50], [200, 50], [200, 75], [100, 75]]),
            OCRBox("Nội dung", 0.98, [[220, 50], [350, 50], [350, 75], [220, 75]]),
            # Tuần 1
            OCRBox("1", 0.95, [[40, 110], [60, 110], [60, 135], [40, 135]]),
            OCRBox("20/04/2026 - 26/04/2026", 0.98, [[100, 110], [215, 110], [215, 135], [100, 135]]),
            OCRBox("- Họp nhóm, tìm hiểu về AWS.", 0.95, [[220, 95], [550, 95], [550, 120], [220, 120]]),
            OCRBox("- Đọc các nội quy.", 0.95, [[220, 125], [500, 125], [500, 150], [220, 150]]),
            # Tuần 2: Bullet 1 nằm ở đầu tuần 2 (Y: 210-235), cao hơn anchor STT 2 và Ngày (Y: 250-275) do STT căn giữa dọc
            OCRBox("- Học kiến thức về AWS.", 0.95, [[220, 210], [520, 210], [520, 235], [220, 235]]),
            OCRBox("2", 0.95, [[40, 250], [60, 250], [60, 275], [40, 275]]),
            OCRBox("27/04/2026 - 03/05/2026", 0.98, [[100, 250], [215, 250], [215, 275], [100, 275]]),
            OCRBox("- Đã tạo account.", 0.95, [[220, 250], [500, 250], [500, 275], [220, 275]]),
        ]
        lines, tables = reconstruct_document_layout(boxes)
        assert len(tables) == 1
        tbl = tables[0]
        assert len(tbl.grid) == 3
        # Tuần 1
        assert tbl.grid[1][0] == "1"
        assert "20/04/2026" in tbl.grid[1][1]
        assert "Họp nhóm" in tbl.grid[1][2]
        assert "Đọc các nội quy" in tbl.grid[1][2]
        # Tuần 2
        assert tbl.grid[2][0] == "2"
        assert "27/04/2026" in tbl.grid[2][1]
        assert "Học kiến thức về AWS" in tbl.grid[2][2]
        assert "Đã tạo account" in tbl.grid[2][2]

    def test_03_and_04_sub_criterion_stt_guard_and_sequence(self):
        """3 & 4. Hàng tiêu chí con không bị gán STT và STT hàng chính không bị lệch."""
        boxes = [
            # Header
            OCRBox("STT", 0.95, [[30, 50], [70, 50], [70, 75], [30, 75]]),
            OCRBox("Tiêu chí đánh giá", 0.98, [[100, 50], [450, 50], [450, 75], [100, 75]]),
            OCRBox("Điểm tối đa", 0.95, [[600, 50], [680, 50], [680, 75], [600, 75]]),
            # Hàng chính 1
            OCRBox("1", 0.95, [[40, 90], [60, 90], [60, 115], [40, 115]]),
            OCRBox("Chấp hành thời gian, thời lượng sinh hoạt", 0.98, [[100, 90], [500, 90], [500, 115], [100, 115]]),
            OCRBox("05", 0.95, [[620, 90], [660, 90], [660, 115], [620, 115]]),
            # Hàng tiêu chí con 1
            OCRBox("+ Thời gian sinh hoạt chi bộ ít hơn quy định 15 - 30 phút", 0.98, [[100, 125], [550, 125], [550, 150], [100, 150]]),
            OCRBox("03", 0.95, [[620, 125], [660, 125], [660, 150], [620, 150]]),
            # Hàng tiêu chí con 2
            OCRBox("+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định", 0.98, [[100, 160], [550, 160], [550, 185], [100, 185]]),
            OCRBox("02", 0.95, [[620, 160], [660, 160], [660, 185], [620, 185]]),
            # Hàng chính 2 (STT 2 chuẩn)
            OCRBox("2", 0.95, [[40, 195], [60, 195], [60, 220], [40, 220]]),
            OCRBox("Tỷ lệ đảng viên dự sinh hoạt chi bộ", 0.98, [[100, 195], [500, 195], [500, 220], [100, 220]]),
            OCRBox("05", 0.95, [[620, 195], [660, 195], [660, 220], [620, 220]]),
            # Hàng tiêu chí con 3
            OCRBox("+ Từ 90% trở lên", 0.98, [[100, 230], [300, 230], [300, 255], [100, 255]]),
            OCRBox("04", 0.95, [[620, 230], [660, 230], [660, 255], [620, 255]]),
        ]
        lines, tables = reconstruct_document_layout(boxes)
        assert len(tables) == 1
        grid = tables[0].grid
        assert grid[1][0] == "1"
        assert grid[2][0] == ""
        assert "Thời gian sinh hoạt chi bộ ít hơn quy định 15 - 30 phút" in grid[2][1]
        assert grid[3][0] == ""
        assert grid[4][0] == "2"
        assert "Tỷ lệ đảng viên dự sinh hoạt chi bộ" in grid[4][1]
        assert grid[5][0] == ""
        assert "Từ 90% trở lên" in grid[5][1]

    def test_05_footnote_normalization(self):
        """5. '?Trừ trường hợp...' được chuẩn hóa thành '* Trừ trường hợp...'."""
        raw = "?Trừ trường hợp được miễn công tác, sinh hoạt đảng theo quy định."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert cleaned == "* Trừ trường hợp được miễn công tác, sinh hoạt đảng theo quy định."

        raw2 = "? Trừ trường hợp đặc biệt."
        cleaned2 = VietnameseTextPostProcessor.clean_and_restore(raw2)
        assert cleaned2 == "* Trừ trường hợp đặc biệt."

    def test_06_unmatched_closing_brackets(self):
        """6. Dấu ']' rác ở cuối câu (trong năm]) được làm sạch mà không xóa ngoặc [...] hợp lệ."""
        raw = "quá 03 lần liên tiếp trong năm]"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert cleaned == "quá 03 lần liên tiếp trong năm"

        raw_valid = "Theo quy định tại [Điều 5] của [Nghị định 01]."
        cleaned_valid = VietnameseTextPostProcessor.clean_and_restore(raw_valid)
        assert cleaned_valid == raw_valid

    def test_07_half_fraction_recovery(self):
        """7. '% so với quy định' được khôi phục thành '½ so với quy định' và phần trăm số được giữ nguyên."""
        raw = "Thời gian sinh hoạt chi bộ chỉ bằng % so với quy định"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert cleaned == "Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định"

        raw_num = "Từ 90% trở lên và đạt 100% chỉ tiêu."
        cleaned_num = VietnameseTextPostProcessor.clean_and_restore(raw_num)
        assert "90%" in cleaned_num
        assert "100%" in cleaned_num

    def test_08_casing_uppercase_preservation(self):
        """8. Casing chữ IN HOA của 'CHI BỘ VĂN PHÒNG HĐND' được bảo toàn tuyệt đối."""
        raw = "CHI BỘ VĂN PHÒNG HĐND VÀ UBND XÃ"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert cleaned == "CHI BỘ VĂN PHÒNG HĐND VÀ UBND XÃ"
        assert "chi bộ" not in cleaned

    def test_09_chi_bo_spell_check(self):
        """9. 'sinh hoạt chỉ bộ' phát hiện và bôi đỏ đúng từ 'chỉ' (gợi ý 'chi')."""
        from app.core.spell_corrector import check_text_spelling
        text = "sinh hoạt chỉ bộ định kỳ tháng 10 năm 2026"
        errs = check_text_spelling(text)
        tokens = [err.token for err in errs]
        assert "chỉ" in tokens
        err = next(e for e in errs if e.token == "chỉ")
        assert "chi" in err.suggestions

    def test_10_du_sinh_hoat_spell_check(self):
        """10. 'Tỷ lệ đảng viên dư sinh hoạt chi bộ' phát hiện và bôi đỏ đúng từ 'dư' (gợi ý 'dự')."""
        from app.core.spell_corrector import check_text_spelling
        text = "Tỷ lệ đảng viên dư sinh hoạt chi bộ"
        errs = check_text_spelling(text)
        tokens = [err.token for err in errs]
        assert "dư" in tokens
        err = next(e for e in errs if e.token == "dư")
        assert "dự" in err.suggestions

    def test_11_zero_false_positives_for_valid_words(self):
        """11. Các nghĩa đúng của 'chỉ', 'dư', 'Chi ủy' không bị báo sai (Zero False Positives)."""
        from app.core.spell_corrector import check_text_spelling
        safe_texts = [
            "Đồng chí chỉ đạo hội nghị.",
            "Tôi chỉ có một câu hỏi.",
            "Anh ấy chỉ tay về phía trước.",
            "Tài khoản có số dư 500 triệu đồng.",
            "Hàng hóa dư thừa trong kho.",
            "Dư âm của chiến thắng vẫn còn đọng lại.",
            "Chi ủy chi bộ họp định kỳ.",
        ]
        for t in safe_texts:
            errs = check_text_spelling(t)
            tokens = [err.token for err in errs]
            assert "chỉ" not in tokens, f"Từ 'chỉ' hợp lệ trong '{t}' bị báo sai!"
            assert "dư" not in tokens, f"Từ 'dư' hợp lệ trong '{t}' bị báo sai!"

    def test_12_exact_offsets_with_special_characters(self):
        """12. Offset lỗi chính tả chính xác hoàn toàn khi văn bản chứa email, URL và ký tự đặc biệt."""
        from app.core.spell_corrector import check_text_spelling
        text = "Gửi tới vpt.luan@hutech.edu.vn và tham dư buổi họp tại https://hutech.edu.vn"
        errs = check_text_spelling(text)
        assert len(errs) == 1
        err = errs[0]
        assert err.token == "dư"
        assert text[err.start_idx:err.end_idx] == "dư"
        assert err.suggestions[0] == "dự"

    def test_13_heading_split_and_aws_recovery(self):
        """13. Tách đề mục số bị dính dòng và khôi phục AW -> AWS."""
        raw_heading = "Chuyên ngành | Công nghệ phần mềm 5. Tiến độ và nội dung thực tập:"
        cleaned_heading = VietnameseTextPostProcessor.clean_and_restore(raw_heading)
        assert "Chuyên ngành | Công nghệ phần mềm\n5. Tiến độ và nội dung thực tập:" in cleaned_heading

        raw_aws = "tìm hiểu về doanh nghiệp, các dịch vụ của AW Làm quen"
        cleaned_aws = VietnameseTextPostProcessor.clean_and_restore(raw_aws)
        assert "các dịch vụ của AWS. Làm quen" in cleaned_aws


class TestGeneralPurposeOCREngineAndUniversalSpellChecker:
    """
    Bộ kiểm thử cho Hệ thống OCR & Kiểm tra chính tả Tổng quát (General-Purpose):
    1. Khôi phục số La Mã tổng quát (IIL -> III., IL/TT -> II., L -> I.)
    2. Chuẩn hóa đơn vị đo lường, số mũ (m2, m? -> m²; m3 -> m³; °C; phân số ½)
    3. Chuẩn hóa khoảng trắng dấu câu & nối dòng nhãn ngắn
    4. Tái dựng bảng: phân biệt hàng con vs hàng chính bằng hình học và cú pháp
    5. Kiểm tra chính tả 3 tầng: Zero False Positives trên từ đúng và phát hiện chính xác lỗi sai
    """

    @pytest.mark.parametrize(
        "raw_line,expected",
        [
            ("IIL NỘI DUNG ĐỀ NGHỊ", "III. NỘI DUNG ĐỀ NGHỊ"),
            ("IIl. NỘI DUNG ĐỀ NGHỊ", "III. NỘI DUNG ĐỀ NGHỊ"),
            ("11L. NỘI DUNG ĐỀ NGHỊ", "III. NỘI DUNG ĐỀ NGHỊ"),
            ("IL THÔNG TIN THỬA ĐẤT", "II. THÔNG TIN THỬA ĐẤT"),
            ("Il. THÔNG TIN THỬA ĐẤT", "II. THÔNG TIN THỬA ĐẤT"),
            ("TT. THÔNG TIN THỬA ĐẤT ĐỀ NGHỊ TÍNH/ TÍNH LẠI", "II. THÔNG TIN THỬA ĐẤT ĐỀ NGHỊ TÍNH/ TÍNH LẠI"),
            ("TI. THÔNG TIN THỬA ĐẤT", "II. THÔNG TIN THỬA ĐẤT"),
            ("L THÔNG TIN NGƯỜI SỬ DỤNG ĐẤT", "I. THÔNG TIN NGƯỜI SỬ DỤNG ĐẤT"),
            ("L. THÔNG TIN NGƯỜI SỬ DỤNG ĐẤT", "I. THÔNG TIN NGƯỜI SỬ DỤNG ĐẤT"),
            ("T. THÔNG TIN NGƯỜI SỬ DỤNG ĐẤT", "I. THÔNG TIN NGƯỜI SỬ DỤNG ĐẤT"),
            ("1. Tuyên bố lý do, giới thiệu đại biểu:", "1. Tuyên bố lý do, giới thiệu đại biểu:"),
            ("2. Chủ trì: Đồng chí Lê Văn Điền", "2. Chủ trì: Đồng chí Lê Văn Điền"),
        ],
    )
    def test_roman_numeral_general_recovery(self, raw_line, expected):
        recovered, _, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line(raw_line)
        assert recovered == expected

    @pytest.mark.parametrize(
        "raw_text,expected_substr",
        [
            ("diện tích: 293,7m2", "293,7m²"),
            ("diện tích: 547,7m?", "547,7m²"),
            ("chuyển mục đích sang đất ở 170m2", "170m²"),
            ("thể tích 25m3", "25m³"),
            ("nhiệt độ 37 oC", "37°C"),
            ("nhiệt độ 250C", "25°C"),
            ("chỉ bằng % so với quy định", "chỉ bằng ½ so với quy định"),
            ("dưới % mức quy định", "dưới ½ mức quy định"),
            ("13/04/2021.Nơi cấp: CQLHC", "13/04/2021. Nơi cấp: CQLHC"),
            ("số:473/QĐ-UBND", "số: 473/QĐ-UBND"),
        ],
    )
    def test_universal_units_and_punctuation_cleanups(self, raw_text, expected_substr):
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw_text)
        assert expected_substr in cleaned

    def test_label_and_data_line_reconnection(self):
        raw = "Sinh năm:\n1978. Số CCCD/Hộ chiếu: 077078001980"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert "Sinh năm: 1978. Số CCCD/Hộ chiếu: 077078001980" in cleaned

    def test_table_sub_row_geometry_and_syntax(self):
        """Kiểm tra hàng con không bị gán STT nhầm và không tăng expected_stt."""
        boxes = [
            # Header
            OCRBox("STT", 0.95, [[20, 10], [50, 10], [50, 30], [20, 30]]),
            OCRBox("Tiêu chí đánh giá", 0.98, [[80, 10], [350, 10], [350, 30], [80, 30]]),
            OCRBox("Điểm tối đa", 0.95, [[400, 10], [480, 10], [480, 30], [400, 30]]),
            OCRBox("Điểm tự chấm", 0.95, [[500, 10], [580, 10], [580, 30], [500, 30]]),
            # Hàng 1 (Hàng chính)
            OCRBox("1", 0.95, [[25, 45], [45, 45], [45, 65], [25, 65]]),
            OCRBox("Chấp hành thời gian, thời lượng sinh hoạt", 0.98, [[80, 45], [380, 45], [380, 65], [80, 65]]),
            OCRBox("05", 0.95, [[420, 45], [450, 45], [450, 65], [420, 65]]),
            OCRBox("05", 0.95, [[520, 45], [550, 45], [550, 65], [520, 65]]),
            # Hàng 2 (Hàng con indented + bullet +)
            OCRBox("+ Chi bộ dưới 30 đảng viên ít nhất 90 phút...", 0.98, [[105, 75], [390, 75], [390, 95], [105, 95]]),
            OCRBox("05", 0.95, [[420, 75], [450, 75], [450, 95], [420, 95]]),
            OCRBox("05", 0.95, [[520, 75], [550, 75], [550, 95], [520, 95]]),
            # Hàng 3 (Hàng con bullet -)
            OCRBox("- Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút", 0.98, [[105, 105], [390, 105], [390, 125], [105, 125]]),
            OCRBox("04", 0.95, [[420, 105], [450, 105], [450, 125], [420, 125]]),
            # Hàng 4 (Hàng chính 2)
            OCRBox("2", 0.95, [[25, 135], [45, 135], [45, 155], [25, 155]]),
            OCRBox("Tỷ lệ đảng viên dự sinh hoạt chi bộ", 0.98, [[80, 135], [380, 135], [380, 155], [80, 155]]),
            OCRBox("05", 0.95, [[420, 135], [450, 135], [450, 155], [420, 155]]),
            OCRBox("05", 0.95, [[520, 135], [550, 135], [550, 155], [520, 155]]),
        ]

        lines, tables = reconstruct_document_layout(boxes)
        assert len(tables) == 1
        tbl = tables[0]
        # Hàng 1 chính: STT = "1"
        assert tbl.grid[1][0] == "1"
        assert "Chấp hành thời gian" in tbl.grid[1][1]
        # Hàng 2 con: STT PHẢI LÀ RỖNG "" (không bị gán nhầm thành "2")
        assert tbl.grid[2][0] == ""
        # Hàng 3 con: STT PHẢI LÀ RỖNG ""
        assert tbl.grid[3][0] == ""
        # Hàng 4 chính: STT = "2"
        assert tbl.grid[4][0] == "2"
        assert "Tỷ lệ đảng viên" in tbl.grid[4][1]

    def test_zero_false_positives_on_valid_sentences(self):
        """Golden Rule #1: Các cụm từ hợp lệ không bao giờ bị bôi đỏ."""
        from app.core.spell_corrector import check_text_spelling
        clean_samples = [
            "1 | Chấp hành thời gian, thời lượng sinh hoạt | 05 | 05",
            "3. Thư ký: Đồng chí Nguyễn Văn Thắng, Chi ủy viên Chi bộ.",
            "5. Đồng chí Lê Thụy Thúy Phượng thông báo tình hình thu, nộp đảng phí;",
            "chương trình sinh hoạt chi bộ định kỳ tháng 10 năm 2026",
            "thông tin tình hình đảng viên dự sinh hoạt",
            "Thửa đất số: 593, tờ bản đồ số: 60, diện tích: 293,7m²",
            "quyền sử dụng đất và chuyển mục đích sang đất ở",
            "thanh toán tiền qua tài khoản ngân hàng",
            "kế toán trưởng giải quyết khiếu nại theo hợp đồng",
        ]
        for s in clean_samples:
            errs = check_text_spelling(s)
            assert len(errs) == 0, f"False positive detected in '{s}': {[e.token for e in errs]}"

    def test_multi_domain_accurate_error_detection(self):
        """Phát hiện chính xác lỗi chính tả trên nhiều lĩnh vực và bảo toàn casing."""
        from app.core.spell_corrector import check_text_spelling

        # 1. CHI BỐ -> BỘ
        errs = check_text_spelling("CHI BỐ VĂN PHÒNG")
        assert len(errs) == 1
        assert errs[0].token == "BỐ"
        assert errs[0].suggestions[0] == "BỘ"

        # 2. chỉ bộ -> chi
        errs = check_text_spelling("sinh hoạt chỉ bộ")
        assert len(errs) == 1
        assert errs[0].token == "chỉ"
        assert errs[0].suggestions[0] == "chi"

        # 3. dư sinh hoạt -> dự
        errs = check_text_spelling("đảng viên dư sinh hoạt")
        assert len(errs) == 1
        assert errs[0].token == "dư"
        assert errs[0].suggestions[0] == "dự"

        # 4. Vdng mặt -> Vắng
        text = "Vdng mặt 0 đồng chí"
        errs = check_text_spelling(text)
        assert len(errs) == 1
        assert errs[0].token == "Vdng"
        assert text[errs[0].start_idx:errs[0].end_idx] == "Vdng"
        assert "Vắng" in errs[0].suggestions

        # 5. NGHIÊP VỤ -> NGHIỆP
        errs = check_text_spelling("NGHIÊP VỤ")
        assert len(errs) == 1
        assert errs[0].token == "NGHIÊP"
        assert errs[0].suggestions[0] == "NGHIỆP"

        # 6. thành toán -> thanh
        errs = check_text_spelling("thành toán chuyển khoản")
        assert len(errs) == 1
        assert errs[0].token == "thành"
        assert errs[0].suggestions[0] == "thanh"

        # 7. ngần hàng -> ngân
        errs = check_text_spelling("ngần hàng nhà nước")
        assert len(errs) == 1
        assert errs[0].token == "ngần"
        assert errs[0].suggestions[0] == "ngân"

        # 8. kể toán -> kế
        errs = check_text_spelling("kể toán trưởng")
        assert len(errs) == 1
        assert errs[0].token == "kể"
        assert errs[0].suggestions[0] == "kế"

        # 9. kíh gửi -> kính
        errs = check_text_spelling("kíh gửi ban giám đốc")
        assert len(errs) == 1
        assert errs[0].token == "kíh"
        assert "kính" in errs[0].suggestions

        # 10. đc miễn -> được
        errs = check_text_spelling("đc miễn sinh hoạt")
        assert len(errs) == 1
        assert errs[0].token == "đc"
        assert "được" in errs[0].suggestions


class TestRealWorldDocumentIssues:
    """Kiểm thử toàn diện 6 vấn đề từ văn bản doanh nghiệp / thực tế mới."""

    def test_01_decorative_line_noise_under_motto_removed(self):
        """1.1. Lọc chuỗi số lặp 144444444, ******* dưới tiêu ngữ và ngắt dòng trước Địa danh, ngày..."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        # Trường hợp dính liền chung 1 dòng với số lặp 14444444
        raw1 = "Độc lập - Tự do - Hạnh phúc 14444444 Hà Nội, ngày 02 tháng 01 năm 2026"
        res1 = VietnameseTextPostProcessor.clean_and_restore(raw1)
        assert res1 == "Độc lập - Tự do - Hạnh phúc\nHà Nội, ngày 02 tháng 01 năm 2026"
        assert "14444444" not in res1

        # Trường hợp số lặp 144444444 ở dòng riêng biệt
        raw2 = "Độc lập - Tự do - Hạnh phúc\n144444444\nHà Nội, ngày 02 tháng 01 năm 2026"
        res2 = VietnameseTextPostProcessor.clean_and_restore(raw2)
        assert res2 == "Độc lập - Tự do - Hạnh phúc\nHà Nội, ngày 02 tháng 01 năm 2026"

        # Trường hợp ký tự hoa văn *******
        raw3 = "Độc lập - Tự do - Hạnh phúc\n*******\nHà Nội, ngày 02 tháng 01 năm 2026"
        res3 = VietnameseTextPostProcessor.clean_and_restore(raw3)
        assert res3 == "Độc lập - Tự do - Hạnh phúc\nHà Nội, ngày 02 tháng 01 năm 2026"

    def test_02_right_margin_logo_stamp_text_not_merged(self):
        """1.2. Box con dấu/logo ở lề phải không bị ghép dính vào dòng văn bản chính."""
        from app.core.ocr_engine import OCRBox, group_boxes_into_lines

        boxes = [
            # Dòng chính lề trái
            OCRBox("II. Nội dung cuộc họp", 0.98, [[50, 200], [250, 200], [250, 225], [50, 225]]),
            # Box logo/con dấu ở lề phải xa
            OCRBox("KHIÊN ƯNG KẾ TOÁN", 0.95, [[700, 200], [840, 200], [840, 225], [700, 225]]),
        ]
        lines = group_boxes_into_lines(boxes)
        line_texts = [" ".join(b.text for b in line) for line in lines]
        # Hai khối phải ở hai dòng riêng biệt, không gộp thành 1 dòng
        assert len(lines) == 2
        assert line_texts[0] == "II. Nội dung cuộc họp"
        assert "KHIÊN ƯNG KẾ TOÁN" in line_texts[1]

    def test_03_attached_roman_numeral_split(self):
        """2.1. Tách đề mục số La Mã bị dính vào cuối câu/dòng trước."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        raw = "Tổng số cán bộ của doanh nghiệp: 14 người. IL. Nội dung cuộc họp"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        lines = cleaned.split("\n")
        assert len(lines) == 2
        assert lines[0] == "Tổng số cán bộ của doanh nghiệp: 14 người."
        assert lines[1] == "IL. Nội dung cuộc họp"

    def test_04_title_case_roman_numeral_and_hierarchy_disambiguation(self):
        """2.2. Khôi phục số La Mã với tiêu đề Title Case và Hierarchy Disambiguation."""
        from app.core.ocr_engine import ContextSymbolRecovery

        # Title Case với IL. -> II.
        r1, dec1, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line("IL. Nội dung cuộc họp")
        assert r1 == "II. Nội dung cuộc họp"

        # Title Case với TTT. -> III.
        r2, dec2, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line("TTT. Kết luận")
        assert r2 == "III. Kết luận"

        # Hierarchy Disambiguation: Dòng 1. Thành phần... bên dưới có 1. Đại diện...
        lines = [
            "1. Thành phần tham dự",
            "1. Đại diện Công ty: ông Hoàng Trung Thật",
            "2. Đại diện người lao động: bà Nguyễn Thị Lan",
        ]
        recovered, audit = ContextSymbolRecovery.recover_document_symbols(lines)
        assert recovered[0] == "I. Thành phần tham dự"
        assert recovered[1] == "1. Đại diện Công ty: ông Hoàng Trung Thật"
        assert recovered[2] == "2. Đại diện người lao động: bà Nguyễn Thị Lan"

    def test_05_missing_opening_parenthesis_recovery(self):
        """3. Khôi phục dấu ngoặc đơn mở '(' bị thiếu ở đầu dòng kết thúc bằng ')'."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        raw1 = "Đánh dấu X vào nội dung được chọn)"
        assert VietnameseTextPostProcessor.clean_and_restore(raw1) == "(Đánh dấu X vào nội dung được chọn)"

        raw2 = "Ghi chú: trường hợp đặc biệt)"
        assert VietnameseTextPostProcessor.clean_and_restore(raw2) == "(Ghi chú: trường hợp đặc biệt)"

        # Không thêm nếu dòng đã có dấu mở ngoặc hợp lệ
        valid = "(Đánh dấu X vào nội dung được chọn)"
        assert VietnameseTextPostProcessor.clean_and_restore(valid) == valid

    def test_06_dieu_khoan_protection_and_thien_ung_correction(self):
        """4. Bảo vệ cụm 'điều khoản' và bắt lỗi 'Thiên Ung' -> 'Thiên Ưng'."""
        from app.core.spell_corrector import check_text_spelling

        # Zero false positives trên 'điều khoản'
        clean_sentence = "5. Những điều khoản không tán thành: Không có"
        errs_clean = check_text_spelling(clean_sentence)
        assert len(errs_clean) == 0, f"Không được có lỗi trên câu hợp lệ: {errs_clean}"

        # Bắt lỗi 'Thiên Ung' -> gợi ý 'Ưng' (Title Case)
        sentence_title = "Kế toán Thiên Ung"
        errs_title = check_text_spelling(sentence_title)
        assert len(errs_title) == 1
        assert errs_title[0].token == "Ung"
        assert errs_title[0].suggestions[0] == "Ưng"

        # Bắt lỗi 'THIÊN UNG' -> gợi ý 'ƯNG' (All Caps)
        sentence_upper = "CÔNG TY TNHH DỊCH VỤ KẾ TOÁN THIÊN UNG"
        errs_upper = check_text_spelling(sentence_upper)
        assert len(errs_upper) == 1
        assert errs_upper[0].token == "UNG"
        assert errs_upper[0].suggestions[0] == "ƯNG"


class TestRuntimePipelineAndDottedForm:
    """Kiểm thử chuỗi pipeline runtime, tách đề mục số La Mã và form dấu chấm điền thông tin."""

    def test_roman_numeral_recovery_on_split_lines(self):
        """Xác minh TT. và TTT. được khôi phục thành II. và III. trong văn bản biên bản cuộc họp."""
        from app.core.ocr_engine import ContextSymbolRecovery, VietnameseTextPostProcessor

        raw_input = (
            "Hôm nay, vào lúc...... giờ ....... phút, ngày ........ tháng ..... năm .....\n"
            "Tại Văn phòng công ty ........\n"
            "Chúng tôi gồm có:\n"
            "1. Thành phần tham dự:\n"
            "1. Về phía Ban Giám Đốc Công ty có:\n"
            "- Ông (bà):...\n"
            "2. Về phía tổ chức đại diện người lao động tại cơ sở:\n"
            "3. Thư ký cuộc họp:\n"
            "- Ông (bà):... 14 người. TT. Nội dung cuộc họp: Lấy ý kiến BCH Công đoàn về nội dung bản nội quy lao động của Công ty\n"
            "Mở đầu cuộc họp, đại diện Ban Giám đốc Công ty trình bày nội dung của bản nội quy lao động.\n"
            "TTT. Kết luận\n"
            "Chủ trì cuộc họp lấy ý kiến biểu quyết..."
        )

        # 1. clean_and_restore phải tách dòng dính đề mục
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw_input)
        lines = cleaned.split("\n")
        # 2. ContextSymbolRecovery khôi phục số La Mã
        recovered_lines, audit = ContextSymbolRecovery.recover_document_symbols(lines)
        full_text = "\n".join(recovered_lines)

        assert "I. Thành phần tham dự:" in full_text
        assert "II. Nội dung cuộc họp:" in full_text
        assert "III. Kết luận" in full_text
        assert "TT. Nội dung" not in full_text
        assert "TTT. Kết luận" not in full_text

    def test_spaced_concatenation_restoration(self):
        """Xác minh khôi phục khoảng trắng cho hoạtchi bộ -> hoạt chi bộ."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        text = "2 | Tỷ lệ đảng viên dự sinh hoạtchi bộ | 05 | 05"
        restored = VietnameseTextPostProcessor.clean_and_restore(text)
        assert "sinh hoạt chi bộ" in restored
        assert "hoạtchi" not in restored

    def test_dotted_fill_in_form_grouping(self):
        """Xác minh các ô form điền chấm chấm (...) trên cùng cao độ được gom vào cùng một dòng."""
        from app.core.ocr_engine import OCRBox, _group_boxes_sliding_window, _box_info_list

        boxes = [
            OCRBox("Hôm nay, vào lúc......", 0.95, [[50, 100], [220, 100], [220, 125], [50, 125]]),
            OCRBox("giờ ....... phút,", 0.95, [[260, 100], [380, 100], [380, 125], [260, 125]]),
            OCRBox("ngày ........", 0.95, [[420, 100], [530, 100], [530, 125], [420, 125]]),
            OCRBox("tháng ..... năm .....", 0.95, [[570, 100], [740, 100], [740, 125], [570, 125]]),
        ]

        box_info = _box_info_list(boxes)
        grouped = _group_boxes_sliding_window(box_info, page_min_x=50, page_w=800, gutter_thresh=60, median_h=25)
        assert len(grouped) == 1, f"Tất cả các box form điền trên cùng hàng phải được gom thành 1 dòng, thực tế: {len(grouped)}"
        line_text = " ".join(b.text for b in grouped[0])
        assert "Hôm nay, vào lúc" in line_text
        assert "tháng ..... năm ....." in line_text


class TestMasterPromptFixTransformerAndSpellCheck:
    """Kiểm thử chuyên sâu 5 nhóm tính năng trong Master Prompt."""

    def test_01_hallucinated_and_decorative_noise_filtering(self):
        """1. Kiểm tra lọc bỏ hoàn toàn chuỗi ảo giác Transformer và ký tự rác đơn lẻ."""
        from app.core.ocr_engine import is_hallucinated_or_noise_line

        # Ảo giác lặp từ VietOCR từ đường kẻ hoa văn
        hallucination_str = "đặc đặc đặc đặc đặc đặc được đặc đặc đặc đặc đặn"
        assert is_hallucinated_or_noise_line(hallucination_str) is True

        # Ký tự đơn lẻ vô nghĩa
        assert is_hallucinated_or_noise_line("%") is True
        assert is_hallucinated_or_noise_line("/") is True
        assert is_hallucinated_or_noise_line("\\") is True
        assert is_hallucinated_or_noise_line("|") is True
        assert is_hallucinated_or_noise_line("[") is True
        assert is_hallucinated_or_noise_line("]") is True

        # Đường kẻ trang trí lặp lại
        assert is_hallucinated_or_noise_line("144444444") is True
        assert is_hallucinated_or_noise_line("44444") is True
        assert is_hallucinated_or_noise_line("*****") is True
        assert is_hallucinated_or_noise_line("-----") is True

        # Dòng văn bản bình thường không bị lọc nhầm
        assert is_hallucinated_or_noise_line("Cộng hòa Xã hội Chủ nghĩa Việt Nam") is False
        assert is_hallucinated_or_noise_line("Hôm nay, vào lúc ...... giờ") is False
        assert is_hallucinated_or_noise_line("1. Về phía Ban Giám Đốc Công ty có:") is False

    def test_02_orphaned_brackets_and_footnote_cleanups(self):
        """2. Làm sạch dấu ngoặc vuông '[' mồ côi đầu dòng và khôi phục footnote '*' từ '/'."""
        from app.core.ocr_engine import VietnameseTextPostProcessor, clean_orphaned_brackets

        # Làm sạch ngoặc mồ côi
        assert clean_orphaned_brackets("[3. Thư ký cuộc họp:") == "3. Thư ký cuộc họp:"
        assert clean_orphaned_brackets("[Chủ trì cuộc họp") == "Chủ trì cuộc họp"
        assert clean_orphaned_brackets("[dung của bản nội quy") == "dung của bản nội quy"
        assert clean_orphaned_brackets("|dung của bản nội quy") == "dung của bản nội quy"

        # Bảo toàn cặp ngoặc hợp lệ
        assert clean_orphaned_brackets("[Icon] Trang chủ") == "[Icon] Trang chủ"
        assert clean_orphaned_brackets("[Chủ trì cuộc họp]") == "[Chủ trì cuộc họp]"

        # Khôi phục footnote '*' từ '/'
        sample_footnote = "/Trừ trường hợp được miễn công tác, sinh hoạt đảng theo quy định"
        cleaned_fn = VietnameseTextPostProcessor.clean_and_restore(sample_footnote)
        assert cleaned_fn.startswith("* Trừ trường hợp")

    def test_03_horizontal_form_line_merging(self):
        """3. Gom đúng hàng ngang biểu mẫu có chấm lửng (không tách nhầm thành 2 cột)."""
        from app.core.ocr_engine import OCRBox, _group_boxes_sliding_window, _box_info_list, join_ocr_line_boxes

        boxes = [
            OCRBox("+ Ông (bà): .........", 0.95, [[40, 200], [180, 200], [180, 225], [40, 225]]),
            OCRBox("(ghi tên người đại diện theo pháp luật của Công ty).", 0.95, [[260, 200], [600, 200], [600, 225], [260, 225]]),
            OCRBox("Chức danh:...........", 0.95, [[40, 235], [160, 235], [160, 260], [40, 260]]),
            OCRBox("Chủ trì cuộc họp.", 0.95, [[260, 235], [420, 235], [420, 260], [260, 260]]),
        ]

        box_info = _box_info_list(boxes)
        grouped = _group_boxes_sliding_window(box_info, page_min_x=40, page_w=800, gutter_thresh=60, median_h=25)
        assert len(grouped) == 2, f"Phải gom thành đúng 2 dòng ngang, thực tế: {len(grouped)}"

        line1_text = join_ocr_line_boxes(grouped[0])
        line2_text = join_ocr_line_boxes(grouped[1])

        assert "+ Ông (bà): ......" in line1_text
        assert "(ghi tên người đại diện" in line1_text
        assert "Chức danh: ......" in line2_text
        assert "Chủ trì cuộc họp." in line2_text

    def test_04_scorecard_text_postprocessing(self):
        """4. Xác minh '0s' -> '05', '% so với quy định' -> '½' và bullet '-' -> '+'."""
        from app.core.ocr_engine import VietnameseTextPostProcessor

        line_score = "1 | Chấp hành thời gian, thời lượng sinh hoạt | 05 | 0s"
        cleaned_score = VietnameseTextPostProcessor.clean_and_restore(line_score)
        assert "05 | 05" in cleaned_score
        assert "0s" not in cleaned_score

        line_half1 = "Thời gian sinh hoạt chi bộ chỉ bằng % so với quy định"
        cleaned_half1 = VietnameseTextPostProcessor.clean_and_restore(line_half1)
        assert "chỉ bằng ½ so với quy định" in cleaned_half1

        line_half2 = "Thời gian sinh hoạt chi bộ dưới % so với quy định"
        cleaned_half2 = VietnameseTextPostProcessor.clean_and_restore(line_half2)
        assert "dưới ½ so với quy định" in cleaned_half2

    def test_05_dot_valley_crop_splitting(self):
        """5. Phân tích hình thái học cắt sub-crops cho dòng dài có dải chấm lửng."""
        import cv2
        import numpy as np
        from app.core.ocr_engine import split_dotted_line_crop

        h, w = 30, 600
        img = np.ones((h, w, 3), dtype=np.uint8) * 255

        # Cụm chữ 1: x in [15, 95]
        cv2.rectangle(img, (15, 5), (95, 25), (0, 0, 0), -1)
        # Thung lũng chấm lửng 1: x in [96, 184]
        cv2.circle(img, (120, 27), 2, (0, 0, 0), -1)
        cv2.circle(img, (150, 27), 2, (0, 0, 0), -1)

        # Cụm chữ 2: x in [185, 235]
        cv2.rectangle(img, (185, 5), (235, 25), (0, 0, 0), -1)
        # Thung lũng chấm lửng 2: x in [236, 324]
        cv2.circle(img, (260, 27), 2, (0, 0, 0), -1)
        cv2.circle(img, (290, 27), 2, (0, 0, 0), -1)

        # Cụm chữ 3: x in [325, 415]
        cv2.rectangle(img, (325, 5), (415, 25), (0, 0, 0), -1)

        sub_crops = split_dotted_line_crop(img)
        assert len(sub_crops) >= 3, f"Dòng chấm lửng phải được tách thành ít nhất 3 lát cắt, thực tế: {len(sub_crops)}"

    def test_06_spell_checker_chi_bo_protection(self):
        """6. Không bôi đỏ chữ Chi / chi trong 'Chi bộ' / 'chi bộ' / 'chi ủy', chỉ bắt lỗi 'đồng chi'."""
        from app.core.spell_corrector import check_text_spelling

        text1 = "Chi bộ dưới 30 đảng viên; Thời gian sinh hoạt chi bộ"
        errs1 = check_text_spelling(text1)
        assert errs1 == [], f"Không được báo lỗi trên 'Chi bộ' / 'chi bộ', thực tế: {errs1}"

        text2 = "Thời gian sinh hoạt chi ủy"
        errs2 = check_text_spelling(text2)
        assert errs2 == [], f"Không được báo lỗi trên 'chi ủy', thực tế: {errs2}"

        text3 = "Các đồng chi tham dự cuộc họp"
        errs3 = check_text_spelling(text3)
        assert len(errs3) == 1
        assert errs3[0].token.lower() == "chi"
        assert any(s.lower() == "chí" for s in errs3[0].suggestions)

        text4 = "chỉ bộ tổ chức sinh hoạt"
        errs4 = check_text_spelling(text4)
        assert len(errs4) == 1
        assert errs4[0].token.lower() == "chỉ"
        assert any(s.lower() == "chi" for s in errs4[0].suggestions)

        text5 = "Các thành viên trao đổi thảo luận"
        errs5 = check_text_spelling(text5)
        assert errs5 == [], f"Cụm 'trao đổi' phải được bảo vệ, thực tế: {errs5}"


class TestMasterPromptUpgradeCVAndCropPipeline:
    """
    Test suite cho toàn bộ 5 cơ chế nâng cấp trong Master Prompt:
    1. assess_image_quality & Quality-Aware Adaptive Preprocessing
    2. clean_crop_borders (Border line stripping)
    3. pad_small_crop_for_vietocr (Aspect-ratio padding)
    4. is_decorative_pattern_crop & is_decorative_or_hallucinated
    5. final_document_cleanup
    6. spell_corrector contextual protection (chi, đổi, khoản, lượng)
    """

    def test_01_assess_image_quality_and_adaptive_preprocessing(self):
        """1. Đánh giá chất lượng ảnh văn bản sạch vs ảnh tối/mờ và tắt CLAHE trên ảnh sạch."""
        import cv2
        import numpy as np
        from app.core.ocr_engine import assess_image_quality, OCREngine

        # 1.a. Giả lập ảnh văn bản sạch (nền trắng 255, chữ đen 0)
        clean_img = np.full((100, 200, 3), 255, dtype=np.uint8)
        clean_img[40:60, 20:180] = 0  # Chữ đen
        quality_clean = assess_image_quality(clean_img)

        assert quality_clean["is_clean_document"] is True
        assert quality_clean["bg_level"] >= 195.0
        assert quality_clean["contrast_span"] >= 90.0

        # Kiểm tra tiền xử lý ảnh sạch: không bị làm mờ/haloing bởi CLAHE
        proc_clean = OCREngine.preprocess_image(clean_img, min_side=850)
        assert proc_clean is not None
        assert min(proc_clean.shape[:2]) >= 100

        # 1.b. Giả lập ảnh tối/bóng mờ (nền xám 120, chữ 80 -> tương phản thấp)
        dark_img = np.full((100, 200, 3), 120, dtype=np.uint8)
        dark_img[40:60, 20:180] = 80
        quality_dark = assess_image_quality(dark_img)

        assert quality_dark["is_clean_document"] is False

    def test_02_clean_crop_borders(self):
        """2. Tẩy đường kẻ bảng sát mép cắt (trên, dưới, trái, phải) và bảo vệ chữ cái I, 1, l."""
        import numpy as np
        from app.core.ocr_engine import clean_crop_borders

        h, w = 36, 120
        # Ảnh nền trắng 255
        crop = np.full((h, w, 3), 255, dtype=np.uint8)

        # Giả lập đường kẻ ngang đen ở hàng y=0 (kéo dài suốt chiều rộng)
        crop[0, :] = 0
        # Giả lập đường kẻ dọc đen ở cột x=0 (kéo dài suốt chiều cao)
        crop[:, 0] = 0
        # Giả lập chữ cái I ở vị trí x=20, y từ 10 đến 26 (cao 16px, < 82% h)
        crop[10:26, 20:23] = 0

        cleaned = clean_crop_borders(crop)

        # Hàng y=0 phải được tô lại màu nền trắng
        assert np.all(cleaned[0, :] > 200), "Đường kẻ ngang y=0 phải được tẩy sạch"
        # Cột x=0 phải được tô lại màu nền trắng
        assert np.all(cleaned[:, 0] > 200), "Đường kẻ dọc x=0 phải được tẩy sạch"
        # Chữ cái I ở x=20 không được bị xóa
        assert np.any(cleaned[10:26, 20:23] == 0), "Chữ cái I ở trung tâm crop phải được giữ nguyên"

    def test_03_pad_small_crop_for_vietocr(self):
        """3. Đệm canvas trắng hai bên cho ô hẹp (w < 2.5 * h) và upscale h < 36 lên 44."""
        import numpy as np
        from app.core.ocr_engine import pad_small_crop_for_vietocr

        # Ô số ngắn 05 có h=24, w=30 (w < 2.5 * h)
        h_small, w_small = 24, 30
        crop_num = np.full((h_small, w_small, 3), 255, dtype=np.uint8)
        crop_num[6:18, 10:20] = 0

        padded = pad_small_crop_for_vietocr(crop_num)

        assert padded.shape[0] == 44, f"Chiều cao phải được upscale lên 44, thực tế: {padded.shape[0]}"
        assert padded.shape[1] > w_small + 28, f"Chiều rộng phải được đệm trắng hai bên, thực tế: {padded.shape[1]}"

        # Dòng chữ dài thông thường (w >= 2.5 * h) không bị padding thêm
        h_wide, w_wide = 30, 200
        crop_wide = np.full((h_wide, w_wide, 3), 255, dtype=np.uint8)
        unpadded = pad_small_crop_for_vietocr(crop_wide)
        assert unpadded.shape == (h_wide, w_wide, 3), "Dòng chữ dài không bị biến đổi kích thước"

    def test_04_decorative_pattern_crop_and_hallucination_filter(self):
        """4. Nhận diện dải hoa văn trang trí và lọc bỏ dòng ảo giác lặp từ / chuỗi rác."""
        import numpy as np
        import cv2
        from app.core.ocr_engine import is_decorative_pattern_crop, is_decorative_or_hallucinated

        # 4.a. Tạo lát cắt chứa dải chấm nhỏ sát đáy (hoa văn trang trí)
        h, w = 24, 150
        deco_crop = np.full((h, w, 3), 255, dtype=np.uint8)
        for x in range(10, 140, 12):
            cv2.circle(deco_crop, (x, 21), 2, (0, 0, 0), -1)

        assert is_decorative_pattern_crop(deco_crop) is True, "Dải chấm nhỏ sát đáy phải được nhận diện là hoa văn"
        assert is_decorative_or_hallucinated("*******", deco_crop) is True

        # 4.b. Chuỗi text ảo giác lặp từ
        assert is_decorative_or_hallucinated("đặc đặc đặc đặc đặc") is True
        assert is_decorative_or_hallucinated("144444444") is True
        assert is_decorative_or_hallucinated("%") is True
        assert is_decorative_or_hallucinated("[") is True

        # Văn bản hợp lệ không bị lọc
        assert is_decorative_or_hallucinated("Chi bộ Văn phòng HĐND và UBND xã") is False

    def test_05_final_document_cleanup(self):
        """5. Chuẩn hóa cuối cùng trong final_document_cleanup (bảo toàn casing, sửa số, phân số, gạch đầu dòng)."""
        from app.core.ocr_engine import final_document_cleanup

        # 5.a. Chuẩn hóa chi bộ bảo toàn casing
        assert final_document_cleanup("CHI BỐ VĂN PHÒNG") == "CHI BỘ VĂN PHÒNG"
        assert final_document_cleanup("CHỈ BỘ VĂN PHÒNG") == "CHI BỘ VĂN PHÒNG"
        assert final_document_cleanup("Chi bố Văn phòng") == "Chi bộ Văn phòng"
        assert final_document_cleanup("Chỉ bộ Văn phòng") == "Chi bộ Văn phòng"
        assert final_document_cleanup("sinh hoạt chi bố") == "sinh hoạt chi bộ"
        assert final_document_cleanup("sinh hoạt chỉ bộ") == "sinh hoạt chi bộ"

        # 5.b. sinh hoatchi -> sinh hoạt chi
        assert final_document_cleanup("Tỷ lệ đảng viên dự sinh hoatchi bộ") == "Tỷ lệ đảng viên dự sinh hoạt chi bộ"

        # 5.c. (\d+) đồng ch -> \1 đồng chí
        assert final_document_cleanup("Số lượng 14 đồng ch.") == "Số lượng 14 đồng chí."
        assert final_document_cleanup("Có mặt: 14 đồng ch.") == "- Có mặt: 14 đồng chí."
        assert final_document_cleanup("Vắng mặt: 0 đồng ch") == "- Vắng mặt: 0 đồng chí"

        # 5.d. Sửa 0s -> 05
        assert final_document_cleanup("Điểm tự chấm: 0s | 0s") == "Điểm tự chấm: 05 | 05"

        # 5.e. Sửa % so với quy định -> ½
        assert final_document_cleanup("chỉ bằng % so với quy định") == "chỉ bằng ½ so với quy định"

        # 5.f. Sửa /Trừ trường hợp -> * Trừ trường hợp
        assert final_document_cleanup("/Trừ trường hợp đặc biệt") == "* Trừ trường hợp đặc biệt"

        # 5.g. Xóa [ mồ côi
        assert final_document_cleanup("[3. Thư ký cuộc họp") == "3. Thư ký cuộc họp"

        # 5.h. Khôi phục dấu '- ' đầu dòng biên bản
        raw_mins = (
            "Lý do: Tổ chức sinh hoạt chi bộ định kỳ\n"
            "Giới thiệu đại biểu: không có\n"
            "Miễn sinh hoạt: 0 đồng chí\n"
            "Có mặt: 14 đồng chí\n"
            "Vắng mặt: 0 đồng chí"
        )
        cleaned_mins = final_document_cleanup(raw_mins)
        assert "- Lý do: Tổ chức sinh hoạt chi bộ định kỳ" in cleaned_mins
        assert "- Giới thiệu đại biểu: không có" in cleaned_mins
        assert "- Miễn sinh hoạt: 0 đồng chí" in cleaned_mins
        assert "- Có mặt: 14 đồng chí" in cleaned_mins
        assert "- Vắng mặt: 0 đồng chí" in cleaned_mins

    def test_06_spell_corrector_protected_words_and_real_errors(self):
        """6. Spell corrector bảo vệ tuyệt đối chi, đổi, khoản, lượng và bắt lỗi chính xác."""
        from app.core.spell_corrector import check_text_spelling

        # 6.a. Các cụm từ ĐÚNG được bảo vệ tuyệt đối
        text_safe = (
            "Thời gian sinh hoạt chi bộ định kỳ. "
            "Chi ủy chuẩn bị nội dung và thời lượng 90 phút. "
            "Các thành viên trao đổi thảo luận theo điều khoản quy định."
        )
        errs_safe = check_text_spelling(text_safe)
        assert errs_safe == [], f"Không được gắn cờ lỗi trên văn bản đúng, thực tế: {errs_safe}"

        # 6.b. Bắt lỗi chính xác từ sai thực sự
        errs_vdng = check_text_spelling("Vdng mặt: 0 đồng chí")
        assert len(errs_vdng) == 1
        assert errs_vdng[0].token == "Vdng"

        errs_chibo = check_text_spelling("chỉ bộ sinh hoạt")
        assert len(errs_chibo) == 1
        assert errs_chibo[0].token == "chỉ"

        errs_chibo_caps = check_text_spelling("CHI BỐ HỌP ĐỊNH KỲ")
        assert len(errs_chibo_caps) == 1
        assert errs_chibo_caps[0].token == "BỐ"

        errs_du = check_text_spelling("Đảng viên dư sinh hoạt")
        assert len(errs_du) == 1
        assert errs_du[0].token == "dư"


class TestImageSpecificFixes698And6a0:
    """Kiểm thử các bản sửa lỗi cụ thể trên image_698e9c.jpg và image_6a0da4.jpg."""

    def test_01_khong_co_and_so_hieu_van_ban(self):
        """1. Khôi phục 'không c' -> 'không có.' và dấu gạch nối số hiệu 'Số 01 BB/CB' -> 'Số 01-BB/CB'."""
        from app.core.ocr_engine import final_document_cleanup

        # Khôi phục chữ bị cắt mép phải
        assert final_document_cleanup("Giới thiệu đại biểu: không c") == "- Giới thiệu đại biểu: không có."
        assert final_document_cleanup("- Giới thiệu đại biểu: không c") == "- Giới thiệu đại biểu: không có."
        assert final_document_cleanup("Giới thiệu đại biểu: không c.") == "- Giới thiệu đại biểu: không có."

        # Khôi phục dấu gạch nối trong số hiệu văn bản
        assert final_document_cleanup("Số 01 BB/CB") == "Số 01-BB/CB"
        assert final_document_cleanup("Số: 01 BB/CB") == "Số: 01-BB/CB"
        assert final_document_cleanup("Số 12 QĐ/UBND") == "Số 12-QĐ/UBND"

    def test_02_clean_crop_borders_right_edge_threshold(self):
        """1.b. clean_crop_borders không xóa mép phải nếu chiều cao nét đen < 92% h."""
        import numpy as np
        from app.core.ocr_engine import clean_crop_borders

        h, w = 50, 100
        # Nền trắng (255)
        crop = np.full((h, w, 3), 255, dtype=np.uint8)

        # Cột w-2 có 43 pixel đen (43 / 50 = 86% < 92%)
        # Với ngưỡng mới 92%, cột này KHÔNG bị tẩy thành nền trắng
        crop[:43, w - 2] = 0

        cleaned = clean_crop_borders(crop)
        # Kiểm tra cột w-2 vẫn giữ nguyên nét đen
        assert np.any(cleaned[:43, w - 2] == 0), "Cột mép phải có mật độ đen 86% (<92%) không được bị tẩy"

        # Cột w-1 có 48 pixel đen (48 / 50 = 96% >= 92% -> đường kẻ khung)
        crop_line = np.full((h, w, 3), 255, dtype=np.uint8)
        crop_line[:48, w - 1] = 0
        cleaned_line = clean_crop_borders(crop_line)
        # Cột w-1 phải được tẩy vì >= 92% h
        assert np.all(cleaned_line[:, w - 1] == 255), "Cột mép phải có mật độ đen >=92% phải được tẩy"

    def test_03_table_header_da_to_diem_toi_da(self):
        """2. Tiêu đề cột 'đa' / 'tối đa' nằm giữa 'Tiêu chí đánh giá' và 'Điểm tự chấm' -> 'Điểm tối đa'."""
        from app.core.ocr_engine import final_document_cleanup, detect_and_reconstruct_table, OCRBox

        # Chuẩn hóa qua regex trong final_document_cleanup
        row = "STT | Tiêu chí đánh giá | đa | Điểm tự chấm | Ghi chú"
        assert final_document_cleanup(row) == "STT | Tiêu chí đánh giá | Điểm tối đa | Điểm tự chấm | Ghi chú"

        row_toi_da = "STT | Tiêu chí đánh giá | tối đa | Điểm tự chấm | Ghi chú"
        assert final_document_cleanup(row_toi_da) == "STT | Tiêu chí đánh giá | Điểm tối đa | Điểm tự chấm | Ghi chú"

        # Chuẩn hóa qua detect_and_reconstruct_table
        boxes = [
            # Header
            OCRBox(text="STT", bbox=[[10, 10], [40, 10], [40, 30], [10, 30]], confidence=0.95),
            OCRBox(text="Tiêu chí đánh giá", bbox=[[50, 10], [200, 10], [200, 30], [50, 30]], confidence=0.95),
            OCRBox(text="đa", bbox=[[210, 10], [250, 10], [250, 30], [210, 30]], confidence=0.95),
            OCRBox(text="Điểm tự chấm", bbox=[[260, 10], [350, 10], [350, 30], [260, 30]], confidence=0.95),
            OCRBox(text="Ghi chú", bbox=[[360, 10], [420, 10], [420, 30], [360, 30]], confidence=0.95),
            # Row 1
            OCRBox(text="1", bbox=[[10, 40], [40, 40], [40, 60], [10, 60]], confidence=0.95),
            OCRBox(text="Tỷ lệ đảng viên dự sinh hoạt", bbox=[[50, 40], [200, 40], [200, 60], [50, 60]], confidence=0.95),
            OCRBox(text="05", bbox=[[210, 40], [250, 40], [250, 60], [210, 60]], confidence=0.95),
            OCRBox(text="05", bbox=[[260, 40], [350, 40], [350, 60], [260, 60]], confidence=0.95),
            # Row 2
            OCRBox(text="2", bbox=[[10, 70], [40, 70], [40, 90], [10, 90]], confidence=0.95),
            OCRBox(text="Thời gian sinh hoạt chi bộ", bbox=[[50, 70], [200, 70], [200, 90], [50, 90]], confidence=0.95),
            OCRBox(text="05", bbox=[[210, 70], [250, 70], [250, 90], [210, 90]], confidence=0.95),
            OCRBox(text="05", bbox=[[260, 70], [350, 70], [350, 90], [260, 90]], confidence=0.95),
        ]
        box_info = []
        for b in boxes:
            l = min(pt[0] for pt in b.bbox)
            r = max(pt[0] for pt in b.bbox)
            t = min(pt[1] for pt in b.bbox)
            bt = max(pt[1] for pt in b.bbox)
            box_info.append({
                "box": b,
                "text": b.text,
                "left": l,
                "right": r,
                "top": t,
                "bottom": bt,
                "w": r - l,
                "h": bt - t,
                "cy": (t + bt) / 2.0,
            })
        tbl, _, _, _ = detect_and_reconstruct_table(box_info, 10, 420, 1000, 20.0)
        assert tbl is not None
        assert tbl.grid[0][2] == "Điểm tối đa", f"Ô header 'đa' phải được đổi thành 'Điểm tối đa', thực tế: '{tbl.grid[0][2]}'"

    def test_04_chi_bo_phan_so_bullet_asterisk(self):
        """3. Chuẩn hóa 'chỉ bồ' -> 'chi bộ', phân số ½, bullet '+' và '* Trừ'."""
        from app.core.ocr_engine import final_document_cleanup

        # 3.1. Sửa chỉ bồ / chi bồ / chỉ bộ / chi bố -> chi bộ (bảo toàn hoa/thường)
        assert final_document_cleanup("2 | Tỷ lệ đảng viên dư sinh hoạt chỉ bồ | 05 | 05") == "2 | Tỷ lệ đảng viên dư sinh hoạt chi bộ | 05 | 05"
        assert final_document_cleanup("Tỷ lệ đảng viên dự sinh hoạt chi bồ") == "Tỷ lệ đảng viên dự sinh hoạt chi bộ"
        assert final_document_cleanup("CHỈ BỒ ĐẢNG ỦY") == "CHI BỘ ĐẢNG ỦY"
        assert final_document_cleanup("Chỉ bồ Đảng ủy") == "Chi bộ Đảng ủy"

        # 3.2. Thay thế % và ％ thành ½
        assert final_document_cleanup("chỉ bằng % so với quy định") == "chỉ bằng ½ so với quy định"
        assert final_document_cleanup("dưới % so với quy định") == "dưới ½ so với quy định"
        assert final_document_cleanup("chỉ bằng ％ so với quy định") == "chỉ bằng ½ so với quy định"
        assert final_document_cleanup("dưới ％ so với quy định") == "dưới ½ so với quy định"

        # 3.3. Bảo toàn bullet '+' cho tiêu chí con có điểm | 0\d
        sub_row = "- Chi bộ dưới 30 đảng viên ít nhất 90 phút | 05 | 05"
        assert final_document_cleanup(sub_row) == "+ Chi bộ dưới 30 đảng viên ít nhất 90 phút | 05 | 05"

        sub_row_time = "- Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút | 03"
        assert final_document_cleanup(sub_row_time) == "+ Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút | 03"

        # 3.4. Thêm khoảng trắng sau dấu * ở dòng chú thích
        assert final_document_cleanup("*Trừ trường hợp đặc biệt") == "* Trừ trường hợp đặc biệt"
        assert final_document_cleanup("*  Trừ trường hợp có lý do chính đáng") == "* Trừ trường hợp có lý do chính đáng"

    def test_05_spell_corrector_dư_hoạt_chỉ_bồ(self):
        """4. Spell corrector: KHÔNG bôi đỏ 'hoạt', chỉ bôi đỏ 'dư' (gợi ý 'dự'), bắt đúng 'chỉ bồ'."""
        from app.core.spell_corrector import check_text_spelling

        # Kiểm tra chuỗi gốc trong prompt
        raw_text = "2 | Tỷ lệ đảng viên dư sinh hoạt chỉ bồ | 05 | 05"
        errs = check_text_spelling(raw_text)

        tokens_flagged = [e.token for e in errs]
        # 'hoạt' TUYỆT ĐỐI không bao giờ bị gắn cờ
        assert "hoạt" not in tokens_flagged, f"'hoạt' không được bị bôi đỏ, danh sách bị cờ: {tokens_flagged}"

        # 'dư' bị gắn cờ với gợi ý 'dự'
        err_du = next((e for e in errs if e.token == "dư"), None)
        assert err_du is not None, "Từ 'dư' phải bị gắn cờ lỗi ngữ cảnh"
        assert err_du.suggestions[0] == "dự", f"Gợi ý đầu tiên cho 'dư' phải là 'dự', thực tế: {err_du.suggestions}"

        # 'chỉ' bị gắn cờ với gợi ý 'chi'
        err_chi = next((e for e in errs if e.token == "chỉ"), None)
        assert err_chi is not None, "Từ 'chỉ' trong 'chỉ bồ' phải bị gắn cờ lỗi"
        assert err_chi.suggestions[0] == "chi", f"Gợi ý đầu tiên cho 'chỉ' phải là 'chi', thực tế: {err_chi.suggestions}"

        # 'bồ' bị gắn cờ với gợi ý 'bộ'
        err_bo = next((e for e in errs if e.token == "bồ"), None)
        assert err_bo is not None, "Từ 'bồ' trong 'chỉ bồ' phải bị gắn cờ lỗi"
        assert err_bo.suggestions[0] == "bộ", f"Gợi ý đầu tiên cho 'bồ' phải là 'bộ', thực tế: {err_bo.suggestions}"

        # Kiểm tra cụm 'chi bồ': chỉ bắt lỗi 'bồ' (gợi ý 'bộ'), 'chi' không bị bắt lỗi
        errs_chi_bo = check_text_spelling("sinh hoạt chi bồ")
        tokens_chi_bo = [e.token for e in errs_chi_bo]
        assert "chi" not in tokens_chi_bo, "'chi' trong 'chi bồ' là đúng, không gắn cờ"
        assert "bồ" in tokens_chi_bo, "'bồ' trong 'chi bồ' phải bị gắn cờ"
        err_bồ_only = next(e for e in errs_chi_bo if e.token == "bồ")
        assert err_bồ_only.suggestions[0] == "bộ"

    def test_06_bien_ban_structure_and_roman_numeral_restoration(self):
        """Khôi phục tiêu đề BIÊN BẢN, đề mục I. Thành phần tham dự và chuẩn hóa II. Nội dung cuộc họp."""
        from app.core.ocr_engine import final_document_cleanup, ContextSymbolRecovery

        sample = (
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n"
            "Độc lập - Tự do - Hạnh phúc\n"
            "Lấy ý kiến tập thể lao động về nội dung Nội quy lao động\n"
            "1. Đại diện Công ty:\n"
            "- Ông Nguyễn Văn A - Giám đốc\n"
            "2. Đại diện người lao động:\n"
            "- Bà Trần Thị B - Chủ tịch Công đoàn\n"
            "IL. Nội dung cuộc họp\n"
            "- Ghi nhận các ý kiến...\n"
            "III. Kết luận cuộc họp"
        )
        cleaned = final_document_cleanup(sample)
        assert "BIÊN BẢN\nLấy ý kiến tập thể lao động" in cleaned, "BIÊN BẢN phải được khôi phục trên tiêu đề phụ"
        assert "I. Thành phần tham dự\n1. Đại diện Công ty:" in cleaned, "I. Thành phần tham dự phải được khôi phục"
        assert "II. Nội dung cuộc họp" in cleaned, "IL. phải đổi thành II. Nội dung cuộc họp"
        assert "III. Nội dung cuộc họp" not in cleaned, "Tuyệt đối không biến thành III. Nội dung cuộc họp"
        assert "III. Kết luận cuộc họp" in cleaned

        # Kiểm tra tính tuần tự trong ContextSymbolRecovery
        raw_lines = [
            "I. Thành phần tham dự",
            "IL. Nội dung cuộc họp",
            "III. Kết luận cuộc họp",
        ]
        rec_lines, _ = ContextSymbolRecovery.recover_document_symbols(raw_lines)
        assert rec_lines[1] == "II. Nội dung cuộc họp", f"IL. phải đổi thành II., thực tế: {rec_lines[1]}"

        # Trùng lặp 2 mục III: 'III. Nội dung cuộc họp' và 'III. Kết luận cuộc họp'
        dup_lines = [
            "I. Khai mạc",
            "III. Nội dung cuộc họp",
            "III. Kết luận cuộc họp",
        ]
        rec_dup, _ = ContextSymbolRecovery.recover_document_symbols(dup_lines)
        assert rec_dup[1] == "II. Nội dung cuộc họp", f"Mục đầu tiên phải đổi thành II., thực tế: {rec_dup[1]}"

    def test_07_clean_orphaned_brackets_and_slash_bullet(self):
        """Làm sạch ký tự rác đầu dòng [Ưng và / -Ghi nhận."""
        from app.core.ocr_engine import final_document_cleanup, clean_orphaned_brackets

        # 1. [Ưng -> Ưng
        assert clean_orphaned_brackets("[Ưng") == "Ưng"
        assert clean_orphaned_brackets("  [Ưng") == "  Ưng"
        assert clean_orphaned_brackets("CÔNG TY KẾ TOÁN THIÊN [Ưng") == "CÔNG TY KẾ TOÁN THIÊN Ưng"
        assert clean_orphaned_brackets("[Đoàn") == "Đoàn"
        assert clean_orphaned_brackets("[Ấn bản") == "Ấn bản"

        # Cặp ngoặc hợp lệ không bị xóa
        assert clean_orphaned_brackets("[Ghi chú: xem trang 2]") == "[Ghi chú: xem trang 2]"

        # 2. / -Ghi nhận -> - Ghi nhận
        cleaned_slash = final_document_cleanup("/ -Ghi nhận các ý kiến đóng góp")
        assert cleaned_slash == "- Ghi nhận các ý kiến đóng góp"
        assert final_document_cleanup("  | -Ghi nhận ý kiến") == "  - Ghi nhận ý kiến"
        assert final_document_cleanup("\\ -Nội dung") == "- Nội dung"

    def test_08_percent_to_half_and_asterisk_tru(self):
        """Khóa chốt thay thế % -> ½ và *Trừ -> * Trừ."""
        from app.core.ocr_engine import final_document_cleanup

        # % -> ½
        assert final_document_cleanup("- Thời gian sinh hoạt chi bộ chỉ bằng % so với quy định | 02") == "+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định | 02"
        assert final_document_cleanup("- Thời gian sinh hoạt chi bộ dưới % so với quy định | 01") == "+ Thời gian sinh hoạt chi bộ dưới ½ so với quy định | 01"
        assert final_document_cleanup("dưới %") == "dưới ½"

        # *Trừ -> * Trừ
        assert final_document_cleanup("*Trừ trường hợp được miễn công tác...") == "* Trừ trường hợp được miễn công tác..."
        assert final_document_cleanup("/*Trừ trường hợp...") == "* Trừ trường hợp..."
        assert final_document_cleanup("1* Trừ trường hợp...") == "* Trừ trường hợp..."

    def test_09_spell_corrector_ke_the_luat_never_flagged(self):
        """Spell Corrector: Không bao giờ bôi đỏ nhầm 'Kế', 'thể', 'luật'."""
        from app.core.spell_corrector import check_text_spelling

        sample_text = (
            "Công ty TNHH Dịch Vụ Kế Toán Thiên Ưng\n"
            "Lấy ý kiến toàn thể cán bộ\n"
            "tập thể người lao động\n"
            "các quy định của pháp \n luật hiện hành"
        )
        errors = check_text_spelling(sample_text)
        flagged_tokens = [e.token.lower() for e in errors]

        assert "kế" not in flagged_tokens, f"'kế' không được bị bôi đỏ: {flagged_tokens}"
        assert "thể" not in flagged_tokens, f"'thể' không được bị bôi đỏ: {flagged_tokens}"
        assert "luật" not in flagged_tokens, f"'luật' không được bị bôi đỏ (kể cả khi ngắt dòng): {flagged_tokens}"

    def test_10_ocr_result_properties(self):
        """OCRResult cung cấp text và lines tương thích."""
        from app.core.ocr_engine import OCRResult, OCRBox

        boxes = [
            OCRBox("Dòng 1", 0.95, [[0, 0], [10, 0], [10, 10], [0, 10]]),
            OCRBox("Dòng 2", 0.98, [[0, 20], [10, 20], [10, 30], [0, 30]]),
        ]
        res = OCRResult(boxes=boxes, full_text="Dòng 1\nDòng 2")
        assert res.text == "Dòng 1\nDòng 2"
        assert res.lines == ["Dòng 1", "Dòng 2"]

    def test_11_master_prompt_line_exact_verification(self):
        """Kiểm tra trực tiếp các trường hợp được chỉ định trong Master Prompt."""
        from app.core.ocr_engine import final_document_cleanup
        from app.core.spell_corrector import check_text_spelling

        # 1.1 final_document_cleanup: bảng điểm & phân số
        assert final_document_cleanup("- Thời gian sinh hoạt chi bộ chỉ bằng % so với quy định | 02") == "+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định | 02"

        # 1.2 final_document_cleanup: *Trừ trường hợp
        assert final_document_cleanup("*Trừ trường hợp được miễn công tác") == "* Trừ trường hợp được miễn công tác"

        # 1.3 final_document_cleanup: Thà Thành phần tham dự deduplication
        dedup_res = final_document_cleanup("Thà Thành phần tham dự\nI. Thành phần tham dự")
        assert dedup_res == "I. Thành phần tham dự", f"Thực tế: {repr(dedup_res)}"

        # 1.4 final_document_cleanup: Khôi phục phụ đề dưới BIÊN BẢN nhảy thẳng xuống Hôm nay
        sub_sample = (
            "BIÊN BẢN\n"
            "Hôm nay, vào lúc 08 giờ 30...\n"
            "Nội quy lao động"
        )
        sub_cleaned = final_document_cleanup(sub_sample)
        assert "BIÊN BẢN\nLấy ý kiến tập thể lao động về nội dung Nội quy lao động\nHôm nay" in sub_cleaned

        # 2. check_text_spelling: 0 lỗi cho 'đầy đủ' và 'Cuộc họp'
        errs_1 = check_text_spelling("14/14 đảng viên thực hiện nộp đảng phí đầy đủ theo quy định.")
        assert errs_1 == [], f"Không được có lỗi với câu 'đầy đủ': {[(e.token, e.suggestions) for e in errs_1]}"

        errs_2 = check_text_spelling("Cuộc họp kết thúc vào 11 giờ 30 cùng ngày.")
        assert errs_2 == [], f"Không được có lỗi với câu 'Cuộc họp': {[(e.token, e.suggestions) for e in errs_2]}"

    def test_12_direct_replacement_and_tail_preservation(self):
        """Kiểm tra 4 điểm sửa trực diện: cột chấm, phân số ½, bullet +, Cuộc hơp, năm 2026."""
        from app.core.ocr_engine import final_document_cleanup, split_dotted_line_crop
        from app.core.spell_corrector import check_text_spelling
        import numpy as np

        # 1. Khôi phục tiêu đề cột thứ 4: | chấm | -> | Điểm tự chấm |
        sample_hdr = "STT | Tiêu chí đánh giá | Điểm tối đa | chấm | Ghi chú"
        assert final_document_cleanup(sample_hdr) == "STT | Tiêu chí đánh giá | Điểm tối đa | Điểm tự chấm | Ghi chú"

        sample_hdr_tu = "STT | Tiêu chí đánh giá | Điểm tối đa | tự chấm | Ghi chú"
        assert final_document_cleanup(sample_hdr_tu) == "STT | Tiêu chí đánh giá | Điểm tối đa | Điểm tự chấm | Ghi chú"

        # 2. Thay thế trực tiếp % thành ½ và đổi - thành +
        sample_table = (
            "STT | Tiêu chí đánh giá | Điểm tối đa | chấm | Ghi chú\n"
            "1 | Chấp hành thời gian, thời lượng sinh hoạt | 05 | 05\n"
            "- Chi bộ dưới 30 đảng viên ít nhất 90 phút; chi bộ trên 30 đảng viên ít nhất 120 phút | 05 | 05\n"
            "- Thời gian sinh hoạt chi bộ ít hơn quy định 15 phút | 04\n"
            "- Thời gian sinh hoạt chi bộ ít hơn quy định 15 - 30 phút | 03\n"
            "- Thời gian sinh hoạt chi bộ chỉ bằng % so với quy định | 02\n"
            "- Thời gian sinh hoạt chi bộ dưới ½ so với quy định | 01\n"
            "2 | Tỷ lệ đảng viên dự sinh hoạt chi bộ | 05 | 05\n"
            "- Từ 90% trở lên | 04 | 04\n"
            "- Từ 75% đến dưới 90% | 03\n"
            "- Từ 65% đến dưới 75% | 02\n"
            "- Dưới 60% hoặc có đảng viên vắng không có lý do | 01\n"
            "- Không có đảng viên vắng mặt có lý do quá 03 lần liên tiếp trong năm | 01 | 01\n"
            "*Trừ trường hợp được miễn công tác, sinh hoạt đảng theo quy định"
        )
        cleaned_table = final_document_cleanup(sample_table)
        assert "| Điểm tối đa | Điểm tự chấm | Ghi chú" in cleaned_table
        assert "+ Chi bộ dưới 30 đảng viên" in cleaned_table
        assert "+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định | 02" in cleaned_table
        assert "+ Thời gian sinh hoạt chi bộ dưới ½ so với quy định | 01" in cleaned_table
        assert "+ Từ 90% trở lên | 04 | 04" in cleaned_table
        assert "* Trừ trường hợp được miễn công tác" in cleaned_table

        # 3. Sửa Cuộc hơp -> Cuộc họp và không bị bắt lỗi chính tả
        hop_sample = "Cuộc hơp kết thúc vào 11 giờ 30 cùng ngày."
        hop_cleaned = final_document_cleanup(hop_sample)
        assert hop_cleaned == "Cuộc họp kết thúc vào 11 giờ 30 cùng ngày."
        assert check_text_spelling(hop_cleaned) == []

        # 4. Sửa lỗi cắt mất số cuối dòng: năm 202 -> năm 2026.
        year_sample = "- Lý do: Tổ chức sinh hoạt chi bộ định kỳ tháng 10 năm 202"
        year_cleaned = final_document_cleanup(year_sample)
        assert year_cleaned == "- Lý do: Tổ chức sinh hoạt chi bộ định kỳ tháng 10 năm 2026."

        # 5. split_dotted_line_crop: Dòng chỉ có 1 khoảng trắng/thung lũng không bị cắt
        dummy_crop = np.ones((40, 200, 3), dtype=np.uint8) * 255
        # Giả lập 1 thung lũng trắng duy nhất ở giữa
        dummy_crop[:, 80:120] = 255
        assert len(split_dotted_line_crop(dummy_crop)) == 1

        # 6. Kiểm tra VietnameseTextPostProcessor.clean_and_restore trực tiếp
        from app.core.ocr_engine import VietnameseTextPostProcessor
        cr_table = VietnameseTextPostProcessor.clean_and_restore(sample_table)
        assert "| Điểm tối đa | Điểm tự chấm | Ghi chú" in cr_table
        assert "+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định | 02" in cr_table
        assert "* Trừ trường hợp được miễn công tác" in cr_table
        assert VietnameseTextPostProcessor.clean_and_restore("Cuộc hơp kết thúc") == "Cuộc họp kết thúc"
        assert VietnameseTextPostProcessor.clean_and_restore("- Lý do: sinh hoạt tháng 10 năm 202") == "- Lý do: sinh hoạt tháng 10 năm 2026."

    def test_13_fraction_1pct_footnote_star_and_admin_rules(self):
        """Kiểm tra đổi 1% -> ½, giữ * Trừ trường hợp và 3 quy tắc văn bản hành chính."""
        from app.core.text_postprocessor import final_document_cleanup, VietnameseTextPostProcessor

        s = (
            "- Thời gian sinh hoạt chi bộ chỉ bằng 1% so với quy định | 02\n"
            "- Thời gian sinh hoạt chi bộ dưới % so với quy định | 01\n"
            "*Trừ trường hợp được miễn công tác"
        )
        out = final_document_cleanup(s)
        expected = (
            "+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định | 02\n"
            "+ Thời gian sinh hoạt chi bộ dưới ½ so với quy định | 01\n"
            "* Trừ trường hợp được miễn công tác"
        )
        assert out == expected

        # 3 quy tắc văn bản hành chính
        r1 = "Yêu cầu các cá nhân, cơ quan có trách nhiệm xây dựng Nội quy lao động"
        assert final_document_cleanup(r1) == "- Yêu cầu các cá nhân, cơ quan có trách nhiệm xây dựng Nội quy lao động"

        r2 = "Tổng số cán bộ của doanh nghiệp: - 14...... người"
        assert final_document_cleanup(r2) == "Tổng số cán bộ của doanh nghiệp: ....14...... người"

        r3 = "Công ty TNHH Dịch vụ Kế Toán Thiên\nƯng thông báo"
        assert final_document_cleanup(r3) == "Công ty TNHH Dịch vụ Kế Toán Thiên Ưng thông báo"


