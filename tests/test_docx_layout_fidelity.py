from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.docx_exporter import export_to_docx
from app.core.models import DocumentBlock, OCRResult


def _export(result: OCRResult, folder: str) -> Document:
    path = Path(folder) / "layout.docx"
    assert export_to_docx(result, str(path))
    return Document(path)


def test_portrait_source_uses_a4_and_unknown_heading_alignment_is_not_guessed() -> None:
    result = OCRResult(
        image_width=1200,
        image_height=1800,
        full_text="Dòng tiêu đề",
        document_blocks=[DocumentBlock(
            kind="heading", order=0, text="Dòng tiêu đề",
            metadata={"format": {"bold_evidence": {"status": "unknown"}}},
        )],
    )
    with TemporaryDirectory() as folder:
        document = _export(result, folder)
    section = document.sections[0]
    assert round(section.page_width.mm) == 210
    assert round(section.page_height.mm) == 297
    assert document.paragraphs[0].alignment is None
    assert document.paragraphs[0].runs[0].bold is None
    assert document.paragraphs[0].text == "Dòng tiêu đề"


def test_landscape_source_uses_a4_landscape_and_keeps_explicit_alignment() -> None:
    result = OCRResult(
        image_width=1800,
        image_height=1200,
        full_text="Tiêu đề căn giữa",
        document_blocks=[DocumentBlock(
            kind="paragraph", order=0, text="Tiêu đề căn giữa",
            metadata={"format": {"alignment": "center", "bold": False}},
        )],
    )
    with TemporaryDirectory() as folder:
        document = _export(result, folder)
    section = document.sections[0]
    assert round(section.page_width.mm) == 297
    assert round(section.page_height.mm) == 210
    assert document.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[0].runs[0].bold is False


def test_unknown_alignment_and_italic_metadata_survive_export() -> None:
    result = OCRResult(
        full_text="Nội dung",
        document_blocks=[DocumentBlock(
            kind="paragraph", order=0, text="Nội dung",
            metadata={"format": {"italic": True}},
        )],
    )
    with TemporaryDirectory() as folder:
        document = _export(result, folder)
    paragraph = document.paragraphs[0]
    assert paragraph.alignment is None
    assert paragraph.runs[0].italic is True
    assert paragraph.text == result.full_text

