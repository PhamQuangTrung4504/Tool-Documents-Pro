# -*- coding: utf-8 -*-
"""
tests.test_multi_genre_documents
================================
Bộ kiểm thử toàn diện đa thể loại tài liệu (Comprehensive Multi-Genre Document Test Suite):
1. Văn bản tiếng Việt in rõ nét (Clear Printed Vietnamese).
2. Văn bản có nhiều dấu và dấu câu phức tạp (Heavy Diacritics & Punctuation).
3. Tài liệu nhiều cột (Multi-Column Documents & Spanning Headers).
4. Bảng có đường kẻ (Bordered Tables).
5. Bảng không có đường kẻ (Borderless / Light-border Tables).
6. Văn bản xen kẽ bảng (Text Interlaced with Tables & Mutual Exclusivity).
7. Tiêu đề nhiều cấp và số La Mã (Multi-Level Headings & Roman Numerals I-X).
8. Danh sách và ký hiệu đầu dòng (Lists & Bullet Normalization).
9. Hóa đơn và bảng số liệu tài chính (Invoices & Financial Tables).
10. Sách, báo và tài liệu giáo dục (Academic, Textbooks & Exam Questions).
11. Tên riêng, mã định danh và thuật ngữ (Proper Names, Technical IDs & Codes).
12. Ảnh nghiêng, mờ, nhiễu và tiền xử lý (Deskewing & Contrast Preprocessing).
13. Số liệu, phần trăm và đơn vị đo (Numbers, Percentages & Measurement Units).
14. Bộ sửa chính tả và phản ví dụ (Negative Whitelisting & Non-Destructive Inspection).
15. Đầu trang, chân trang và chú thích (Headers, Footers, Footnotes & DOCX Export).
"""

import os
import sys
import tempfile
import cv2
import numpy as np
import pytest

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from app.core.ocr_engine import (
    ContextSymbolRecovery,
    SymbolDecision,
    OCRBox,
    OCRResult,
    OCRTable,
    detect_and_reconstruct_table,
    reconstruct_document_layout,
    VietnameseTextPostProcessor,
    estimate_and_correct_skew,
    OCREngine,
    export_to_docx,
)
from app.core.spell_corrector import (
    SpellError,
    check_text_spelling,
    get_suggestions,
)


def _calc_cer(reference: str, hypothesis: str) -> float:
    """Tính Character Error Rate (CER) qua Levenshtein distance."""
    ref = list(reference)
    hyp = list(hypothesis)
    n, m = len(ref), len(hyp)
    if n == 0:
        return 0.0 if m == 0 else 1.0
    dp = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(n + 1):
        dp[i][0] = i
    for j in range(m + 1):
        dp[0][j] = j
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            cost = 0 if ref[i - 1] == hyp[j - 1] else 1
            dp[i][j] = min(dp[i - 1][j] + 1, dp[i][j - 1] + 1, dp[i - 1][j - 1] + cost)
    return dp[n][m] / float(n)


class TestGroup1ClearPrintedVietnamese:
    """1. Văn bản tiếng Việt in rõ nét: chuẩn hóa Unicode NFC, giữ nguyên dấu thanh."""

    def test_unicode_nfc_normalization(self):
        # NFD text -> NFC text
        nfd_text = "Tiê\u0301ng Viê\u0323t Nam quang vinh"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(nfd_text)
        assert cleaned == "Tiếng Việt Nam quang vinh"
        assert _calc_cer("Tiếng Việt Nam quang vinh", cleaned) == 0.0

    def test_administrative_cleanups_applied(self):
        raw = "ĐẰNG CỘNG SẢN VIỆT NAM\nngày 15 thắng 10 năm 2024"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert "ĐẢNG CỘNG SẢN VIỆT NAM" in cleaned
        assert "ngày 15 tháng 10 năm 2024" in cleaned


class TestGroup2HeavyDiacriticsAndPunctuation:
    """2. Văn bản có nhiều dấu và dấu câu phức tạp: bảo toàn khoảng trắng, dấu câu."""

    def test_complex_tones_and_diacritics(self):
        text = "Trường đại học bách khoa đào tạo kỹ sư, cử nhân chất lượng cao."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(text)
        assert cleaned == text

    def test_punctuation_spacing_preserves_file_paths(self):
        text = "Vui lòng xem file config.yml và main.py tại thư mục /app/core/ocr_engine.py ."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(text)
        # Loại bỏ khoảng trắng trước dấu chấm cuối câu, nhưng giữ nguyên file path
        assert "config.yml" in cleaned
        assert "main.py" in cleaned
        assert "/app/core/ocr_engine.py." in cleaned or "/app/core/ocr_engine.py" in cleaned


class TestGroup3MultiColumnDocuments:
    """3. Tài liệu nhiều cột: Tiêu đề trải rộng (spanning banner) trên 2 cột văn bản."""

    def test_spanning_title_and_two_columns_reading_order(self):
        """
        Bố cục bài báo 2 cột có Tiêu đề toàn trang:
        - Box 1: Tiêu đề toàn trang (w = 800, y = 50)
        - Box 2: Cột 1 dòng 1 (left=50, right=450, y=150)
        - Box 3: Cột 1 dòng 2 (left=50, right=450, y=200)
        - Box 4: Cột 2 dòng 1 (left=550, right=950, y=150)
        - Box 5: Cột 2 dòng 2 (left=550, right=950, y=200)
        Thứ tự đọc bắt buộc: Tiêu đề -> Cột 1 dòng 1 -> Cột 1 dòng 2 -> Cột 2 dòng 1 -> Cột 2 dòng 2.
        """
        boxes = [
            OCRBox("NGHIÊN CỨU PHÁT TRIỂN KHOA HỌC CÔNG NGHỆ", 0.98, [[100, 40], [900, 40], [900, 80], [100, 80]]),
            OCRBox("Đoạn văn cột bên trái dòng 1", 0.95, [[50, 140], [450, 140], [450, 170], [50, 170]]),
            OCRBox("Đoạn văn cột bên trái dòng 2", 0.95, [[50, 190], [450, 190], [450, 220], [50, 220]]),
            OCRBox("Đoạn văn cột bên phải dòng 1", 0.95, [[550, 140], [950, 140], [950, 170], [550, 170]]),
            OCRBox("Đoạn văn cột bên phải dòng 2", 0.95, [[550, 190], [950, 190], [950, 220], [550, 220]]),
        ]
        layout_res = reconstruct_document_layout(boxes)
        grouped_lines = layout_res[0]

        extracted_texts = [" ".join(b.text for b in line) for line in grouped_lines]
        assert len(extracted_texts) == 5
        assert extracted_texts[0] == "NGHIÊN CỨU PHÁT TRIỂN KHOA HỌC CÔNG NGHỆ"
        assert extracted_texts[1] == "Đoạn văn cột bên trái dòng 1"
        assert extracted_texts[2] == "Đoạn văn cột bên trái dòng 2"
        assert extracted_texts[3] == "Đoạn văn cột bên phải dòng 1"
        assert extracted_texts[4] == "Đoạn văn cột bên phải dòng 2"


class TestGroup4BorderedTables:
    """4. Bảng có đường kẻ: Nhận diện lưới bảng và phân tách cột chính xác."""

    def test_three_column_grade_table(self):
        """Bảng điểm 3 cột: STT | Môn học | Điểm số"""
        boxes = [
            # Header
            OCRBox("STT", 0.95, [[50, 100], [100, 100], [100, 130], [50, 130]]),
            OCRBox("Môn học", 0.95, [[150, 100], [400, 100], [400, 130], [150, 130]]),
            OCRBox("Điểm số", 0.95, [[450, 100], [550, 100], [550, 130], [450, 130]]),
            # Hàng 1
            OCRBox("1", 0.95, [[50, 150], [100, 150], [100, 180], [50, 180]]),
            OCRBox("Toán cao cấp", 0.95, [[150, 150], [400, 150], [400, 180], [150, 180]]),
            OCRBox("8.5", 0.95, [[450, 150], [550, 150], [550, 180], [450, 180]]),
            # Hàng 2
            OCRBox("2", 0.95, [[50, 200], [100, 200], [100, 230], [50, 230]]),
            OCRBox("Vật lý đại cương", 0.95, [[150, 200], [400, 200], [400, 230], [150, 230]]),
            OCRBox("9.0", 0.95, [[450, 200], [550, 200], [550, 230], [450, 230]]),
            # Hàng 3
            OCRBox("3", 0.95, [[50, 250], [100, 250], [100, 280], [50, 280]]),
            OCRBox("Lập trình Python", 0.95, [[150, 250], [400, 250], [400, 280], [150, 280]]),
            OCRBox("10.0", 0.95, [[450, 250], [550, 250], [550, 280], [450, 280]]),
        ]
        layout_res = reconstruct_document_layout(boxes)
        tables = layout_res[1]
        assert len(tables) == 1
        tbl = tables[0]
        assert tbl.num_rows == 4
        assert tbl.num_cols >= 3
        assert tbl.grid[0][0] == "STT"
        assert tbl.grid[1][1] == "Toán cao cấp"
        assert tbl.grid[2][2] == "9.0"
        assert tbl.grid[3][1] == "Lập trình Python"


class TestGroup5BorderlessTables:
    """5. Bảng không có đường kẻ (bảng kê / báo giá)."""

    def test_borderless_product_table_with_multiline_cells(self):
        """Bảng kê hàng hóa: Mã | Tên sản phẩm | Số lượng | Đơn giá"""
        boxes = [
            # Header
            OCRBox("Mã", 0.95, [[50, 100], [120, 100], [120, 130], [50, 130]]),
            OCRBox("Tên sản phẩm", 0.95, [[160, 100], [450, 100], [450, 130], [160, 130]]),
            OCRBox("Số lượng", 0.95, [[500, 100], [600, 100], [600, 130], [500, 130]]),
            OCRBox("Đơn giá", 0.95, [[650, 100], [780, 100], [780, 130], [650, 130]]),
            # Item 1
            OCRBox("SP01", 0.95, [[50, 150], [120, 150], [120, 180], [50, 180]]),
            OCRBox("Bàn làm việc gỗ sồi cao cấp", 0.95, [[160, 150], [450, 150], [450, 180], [160, 180]]),
            OCRBox("5", 0.95, [[500, 150], [600, 150], [600, 180], [500, 180]]),
            OCRBox("2.500.000", 0.95, [[650, 150], [780, 150], [780, 180], [650, 180]]),
            # Item 2
            OCRBox("SP02", 0.95, [[50, 200], [120, 200], [120, 230], [50, 230]]),
            OCRBox("Ghế xoay công thái học", 0.95, [[160, 200], [450, 200], [450, 230], [160, 230]]),
            OCRBox("10", 0.95, [[500, 200], [600, 200], [600, 230], [500, 230]]),
            OCRBox("1.200.000", 0.95, [[650, 200], [780, 200], [780, 230], [650, 230]]),
            # Item 3
            OCRBox("SP03", 0.95, [[50, 250], [120, 250], [120, 280], [50, 280]]),
            OCRBox("Đèn bàn LED chống cận", 0.95, [[160, 250], [450, 250], [450, 280], [160, 280]]),
            OCRBox("20", 0.95, [[500, 250], [600, 250], [600, 280], [500, 280]]),
            OCRBox("350.000", 0.95, [[650, 250], [780, 250], [780, 280], [650, 280]]),
        ]
        layout_res = reconstruct_document_layout(boxes)
        tables = layout_res[1]
        assert len(tables) == 1
        tbl = tables[0]
        assert tbl.num_rows == 4
        assert tbl.grid[1][0] == "SP01"
        assert tbl.grid[2][0] == "SP02"
        assert tbl.grid[3][0] == "SP03"
        assert "2.500.000" in tbl.grid[1][3]


class TestGroup6TextInterlacedWithTables:
    """6. Văn bản xen kẽ bảng: Nguyên tắc phân hoạch độc quyền Mutual Exclusivity."""

    def test_text_before_and_after_table(self):
        """Đoạn văn mở đầu -> Bảng biểu -> Đoạn văn kết luận."""
        box_pre1 = OCRBox("Căn cứ vào kế hoạch công tác tháng 10 năm 2024", 0.95, [[50, 20], [600, 20], [600, 50], [50, 50]])
        box_pre2 = OCRBox("Ban giám đốc thông qua bảng tổng hợp chi phí sau:", 0.95, [[50, 60], [550, 60], [550, 90], [50, 90]])

        # Bảng
        tbl_boxes = [
            OCRBox("STT", 0.95, [[50, 120], [100, 120], [100, 150], [50, 150]]),
            OCRBox("Khoản mục", 0.95, [[150, 120], [400, 120], [400, 150], [150, 150]]),
            OCRBox("Số tiền", 0.95, [[450, 120], [550, 120], [550, 150], [450, 150]]),
            OCRBox("1", 0.95, [[50, 170], [100, 170], [100, 200], [50, 200]]),
            OCRBox("Chi phí văn phòng phẩm", 0.95, [[150, 170], [400, 170], [400, 200], [150, 200]]),
            OCRBox("5.000.000", 0.95, [[450, 170], [550, 170], [550, 200], [450, 200]]),
            OCRBox("2", 0.95, [[50, 220], [100, 220], [100, 250], [50, 250]]),
            OCRBox("Tiền điện nước tháng 10", 0.95, [[150, 220], [400, 220], [400, 250], [150, 250]]),
            OCRBox("3.200.000", 0.95, [[450, 220], [550, 220], [550, 250], [450, 250]]),
            OCRBox("3", 0.95, [[50, 270], [100, 270], [100, 300], [50, 300]]),
            OCRBox("Chi phí bảo dưỡng thiết bị", 0.95, [[150, 270], [400, 270], [400, 300], [150, 300]]),
            OCRBox("2.000.000", 0.95, [[450, 270], [550, 270], [550, 300], [450, 300]]),
        ]

        box_post1 = OCRBox("Yêu cầu phòng Tài chính Kế toán thực hiện chi đúng định mức.", 0.95, [[50, 330], [650, 330], [650, 360], [50, 360]])
        box_post2 = OCRBox("Biên bản được lập thành 02 bản có giá trị như nhau.", 0.95, [[50, 370], [600, 370], [600, 400], [50, 400]])

        all_boxes = [box_pre1, box_pre2] + tbl_boxes + [box_post1, box_post2]
        layout_res = reconstruct_document_layout(all_boxes)
        grouped_lines = layout_res[0]
        tables = layout_res[1]

        assert len(tables) == 1
        lines_text = [" ".join(b.text for b in line) for line in grouped_lines]
        # Kiểm tra văn bản trước bảng xuất hiện đầu tiên
        assert "Căn cứ vào kế hoạch công tác" in lines_text[0]
        assert "Ban giám đốc thông qua" in lines_text[1]
        # Kiểm tra văn bản sau bảng xuất hiện cuối cùng
        assert "Yêu cầu phòng Tài chính Kế toán" in lines_text[-2]
        assert "Biên bản được lập thành" in lines_text[-1]


class TestGroup7MultiLevelHeadingsAndRomanNumerals:
    """7. Tiêu đề nhiều cấp và số La Mã (I đến X)."""

    @pytest.mark.parametrize(
        "input_heading,expected_heading",
        [
            ("V. KẾ HOẠCH BÀI DẠY", "V. KẾ HOẠCH BÀI DẠY"),
            ("Vl. TÀI LIỆU THAM KHẢO", "VI. TÀI LIỆU THAM KHẢO"),
            ("Vll. PHỤ LỤC VÀ BIỂU MẪU", "VII. PHỤ LỤC VÀ BIỂU MẪU"),
            ("Vlll. ĐIỀU KHOẢN THI HÀNH", "VIII. ĐIỀU KHOẢN THI HÀNH"),
            ("lX. KẾT LUẬN VÀ KIẾN NGHỊ", "IX. KẾT LUẬN VÀ KIẾN NGHỊ"),
            ("X. HIỆU LỰC HỢP ĐỒNG", "X. HIỆU LỰC HỢP ĐỒNG"),
        ],
    )
    def test_roman_numerals_recovery_i_to_x(self, input_heading, expected_heading):
        recovered, dec, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line(input_heading)
        assert recovered == expected_heading
        assert dec in {SymbolDecision.AUTO_CORRECT, SymbolDecision.PRESERVE}

    def test_negative_protection_for_non_roman_items(self):
        """Bảo vệ tuyệt đối số thứ tự thường và mã kỹ thuật."""
        negatives = [
            "1. Đặc điểm chung của đối tượng",
            "2. Mục tiêu nghiên cứu",
            "Mã số: VI-2024/KHCN",
            "Khoản 1 Điều 5",
            "ISO 9001:2015",
        ]
        for item in negatives:
            recovered, dec, _, _ = ContextSymbolRecovery.recover_roman_numeral_in_line(item)
            assert recovered == item
            assert dec == SymbolDecision.PRESERVE


class TestGroup8ListsAndBulletNormalization:
    """8. Danh sách và ký hiệu đầu dòng: bullet vs phép toán vs footnote."""

    def test_bullet_plus_recovery_from_percent(self):
        line = "% Tỷ lệ chuyên cần đạt từ 95% trở lên"
        rec, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(line)
        assert rec == "+ Tỷ lệ chuyên cần đạt từ 95% trở lên"
        assert dec == SymbolDecision.AUTO_CORRECT

    def test_math_equation_not_converted(self):
        math_line = "Biểu thức: a + b = c với x = 90%"
        rec, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(math_line)
        assert rec == math_line
        assert dec == SymbolDecision.PRESERVE

    def test_footnote_asterisk_preserved(self):
        note_line = "* Ghi chú: Số liệu đã được kiểm toán năm 2024."
        rec, dec, _ = ContextSymbolRecovery.recover_bullets_and_pluses(note_line)
        assert rec == note_line


class TestGroup9InvoicesAndFinancialTables:
    """9. Hóa đơn và bảng số liệu tài chính: tiền tệ, phần trăm, số lượng."""

    def test_vat_invoice_post_processing(self):
        raw_text = "HÓA ĐƠN GIÁ TRỊ GIA TĂNG\nTổng tiền hàng: 5O.000.000 VNĐ\nThuế suất VAT: l0%\nTiền thuế: 5.000.000 VNĐ"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw_text)
        assert "50.000.000 VNĐ" in cleaned
        assert "Thuế suất VAT: 10%" in cleaned
        assert "5.000.000 VNĐ" in cleaned

    def test_price_and_percentage_disambiguation(self):
        raw = "Chiết khấu S0% cho khách hàng mua từ 1OO đơn vị."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert "50%" in cleaned
        assert "100" in cleaned


class TestGroup10AcademicAndEducationalDocuments:
    """10. Sách, báo và tài liệu giáo dục: câu hỏi trắc nghiệm A B C D."""

    def test_exam_question_heading_recovery(self):
        raw = "Cầu 1: Thủ đô của Việt Nam là gì?\nA. Hà Nội\nB. TP. Hồ Chí Minh\nC. Đà Nẵng\nD. Cần Thơ"
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert cleaned.startswith("Câu 1:")
        assert "A. Hà Nội" in cleaned

    def test_exercise_heading_recovery(self):
        raw = "Bải 1: Tính đạo hàm của hàm số sau."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(raw)
        assert cleaned.startswith("Bài 1:")


class TestGroup11ProperNamesAndTechnicalIDs:
    """11. Tên riêng, mã định danh, thương hiệu và chữ viết tắt."""

    def test_proper_names_and_acronyms_intact(self):
        text = "Giáo sư Nguyễn Văn A chủ trì đề tài nghiên cứu cùng tập đoàn Samsung và Microsoft."
        errors = check_text_spelling(text)
        # Tên riêng và thương hiệu nổi tiếng không bị đánh dấu lỗi sai
        flagged_tokens = [err.token for err in errors]
        assert "Samsung" not in flagged_tokens
        assert "Microsoft" not in flagged_tokens
        assert "Nguyễn" not in flagged_tokens

    def test_technical_model_codes_protected(self):
        text = "Sản phẩm Model B80 tuân thủ tiêu chuẩn ISO 9001 và chuẩn Bluetooth 5.0."
        errors = check_text_spelling(text)
        flagged_tokens = [err.token for err in errors]
        assert "B80" not in flagged_tokens
        assert "Bluetooth" not in flagged_tokens


class TestGroup12SkewDenoiseAndPreprocessing:
    """12. Ảnh nghiêng, mờ, nhiễu: Thuật toán Deskewing tự động."""

    def test_deskew_straight_image_untouched(self):
        # Tạo ảnh giả lập thẳng (không nghiêng)
        img = np.full((300, 600, 3), 255, dtype=np.uint8)
        cv2.putText(img, "Cong hoa xa hoi chu nghia Viet Nam", (50, 100), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
        cv2.putText(img, "Doc lap Tu do Hanh phuc", (80, 150), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)

        deskewed, angle = estimate_and_correct_skew(img)
        assert abs(angle) < 0.5
        assert deskewed.shape == img.shape

    def test_deskew_tilted_image(self):
        # Tạo ảnh xoay nghiêng 3 độ
        img = np.full((400, 800, 3), 255, dtype=np.uint8)
        for y in range(80, 320, 50):
            cv2.putText(img, "Dong van ban tieng Viet phuc vu kiem thu nghieng", (60, y), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 0), 2)

        rot_mat = cv2.getRotationMatrix2D((400, 200), 3.0, 1.0)
        tilted = cv2.warpAffine(img, rot_mat, (800, 400), borderValue=(255, 255, 255))

        deskewed, angle = estimate_and_correct_skew(tilted)
        assert abs(angle) > 1.0  # Đã phát hiện được góc nghiêng
        assert deskewed.shape == tilted.shape


class TestGroup13NumericalAndMeasurementUnits:
    """13. Số liệu, phần trăm và đơn vị đo: chuẩn hóa OCR."""

    def test_units_and_dates_normalization(self):
        text = "Vào ngày O5 tháng O8 năm 2024, công ty đã giao 500kg gạo và 250 lít dầu ăn."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(text)
        assert "ngày 05" in cleaned
        assert "tháng 08" in cleaned
        assert "500kg" in cleaned or "500 kg" in cleaned
        assert "250 lít" in cleaned

    def test_percentages_with_letters_cleaned(self):
        text = "Hiệu suất đạt B0% trong quý 1, phấn đấu đạt B5% trong quý 2."
        cleaned = VietnameseTextPostProcessor.clean_and_restore(text)
        assert "80%" in cleaned
        assert "85%" in cleaned


class TestGroup14SpellCheckerNegativeWhitelisting:
    """14. Bộ sửa chính tả tiếng Việt: Phản ví dụ và tính bất biến nội dung."""

    def test_negative_whitelisting_multi_genre(self):
        """Các cụm từ đúng trong nhiều lĩnh vực tuyệt đối không bị sửa sai."""
        clean_texts = [
            "Tổng công ty Xây dựng Công trình Giao thông.",
            "Bước qua cây cầu là đến trường học.",
            "Cần đảm bảo an ninh và an toàn trật tự trường học.",
            "Chiến thắng lịch sử Điện Biên Phủ.",
            "Bố mẹ luôn quan tâm việc học tập của con cái.",
            "Học sinh hoàn thành bài tập về nhà đầy đủ.",
            "Bệnh nhân được bác sĩ kiểm tra huyết áp.",
        ]
        for t in clean_texts:
            errors = check_text_spelling(t)
            # Không có từ nào bị đánh dấu nhầm
            flagged = [e.token for e in errors]
            assert "công" not in flagged, f"Flagged 'công' in '{t}'"
            assert "qua" not in flagged, f"Flagged 'qua' in '{t}'"
            assert "an" not in flagged, f"Flagged 'an' in '{t}'"
            assert "thắng" not in flagged, f"Flagged 'thắng' in '{t}'"
            assert "bố" not in flagged, f"Flagged 'bố' in '{t}'"

    def test_contextual_typos_detected_with_valid_suggestions(self):
        """Phát hiện chính xác các lỗi gõ/OCR theo ngữ cảnh."""
        # 1. đáp an -> đáp án
        errs1 = check_text_spelling("Học sinh chọn đáp an đúng nhất.")
        tokens1 = [e.token for e in errs1]
        assert "an" in tokens1
        suggs1 = get_suggestions("an", context_prev="đáp")
        assert "án" in suggs1 or "Án" in suggs1

        # 2. kết qua -> kết quả
        errs2 = check_text_spelling("Nhà trường công bố kết qua thi học kỳ.")
        tokens2 = [e.token for e in errs2]
        assert "qua" in tokens2
        suggs2 = get_suggestions("qua", context_prev="kết")
        assert "quả" in suggs2 or "Quả" in suggs2

        # 3. hóa đon -> hóa đơn
        errs3 = check_text_spelling("Kế toán đã xuất hóa đon cho khách hàng.")
        tokens3 = [e.token for e in errs3]
        assert "đon" in tokens3
        suggs3 = get_suggestions("đon", context_prev="hóa")
        assert "đơn" in suggs3 or "Đơn" in suggs3

        # 4. sản phâm -> sản phẩm
        errs4 = check_text_spelling("Công ty giới thiệu sản phâm mới.")
        tokens4 = [e.token for e in errs4]
        assert "phâm" in tokens4
        suggs4 = get_suggestions("phâm", context_prev="sản")
        assert "phẩm" in suggs4 or "Phẩm" in suggs4


class TestGroup15HeadersFootersAndDocxExport:
    """15. Đầu trang, chân trang, chú thích và xuất file Word DOCX."""

    def test_export_docx_with_multi_genre_table(self):
        """Xuất file DOCX có bảng số liệu và tiêu đề cân đối."""
        from app.core.models import DocumentBlock
        tbl = OCRTable(
            grid=[
                ["STT", "Tên thiết bị", "Số lượng", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"],
                ["1", "Máy tính xách tay Dell XPS", "5", "30.000.000", "150.000.000"],
                ["2", "Màn hình Dell UltraSharp 27 inch", "10", "12.000.000", "120.000.000"],
                ["3", "Bàn phím cơ không dây", "15", "2.000.000", "30.000.000"],
            ],
            bbox=[50, 100, 750, 300],
            num_rows=4,
            num_cols=5,
        )

        full_text = (
            "HỢP ĐỒNG MUA BÁN THIẾT BỊ\n"
            "STT | Tên thiết bị | Số lượng | Đơn giá (VNĐ) | Thành tiền (VNĐ)\n"
            "1 | Máy tính xách tay Dell XPS | 5 | 30.000.000 | 150.000.000\n"
            "2 | Màn hình Dell UltraSharp 27 inch | 10 | 12.000.000 | 120.000.000\n"
            "3 | Bàn phím cơ không dây | 15 | 2.000.000 | 30.000.000\n"
            "Tổng giá trị hợp đồng: 300.000.000 VNĐ."
        )

        ocr_res = OCRResult(
            full_text=full_text,
            tables=[tbl],
            language="vi",
            document_blocks=[
                DocumentBlock(kind="heading", order=0, text="HỢP ĐỒNG MUA BÁN THIẾT BỊ"),
                DocumentBlock(kind="table", order=1, table_index=0),
                DocumentBlock(kind="paragraph", order=2, text="Tổng giá trị hợp đồng: 300.000.000 VNĐ."),
            ],
        )

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            ok = export_to_docx(ocr_res, tmp_path)
            assert ok is True
            assert os.path.exists(tmp_path)
            assert os.path.getsize(tmp_path) > 1000

            # Đọc lại bằng python-docx để thẩm định cấu trúc
            import docx
            doc = docx.Document(tmp_path)
            assert len(doc.tables) >= 1
            exported_tbl = doc.tables[0]
            assert len(exported_tbl.rows) == 4
            assert len(exported_tbl.columns) == 5
            assert exported_tbl.cell(0, 1).text.strip() == "Tên thiết bị"
            assert exported_tbl.cell(1, 1).text.strip() == "Máy tính xách tay Dell XPS"
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
