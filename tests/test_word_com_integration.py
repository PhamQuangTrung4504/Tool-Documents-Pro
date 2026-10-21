"""Opt-in Windows integration test that uses the installed Microsoft Word COM server."""
from __future__ import annotations

import os
import sys
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from PyQt5.QtCore import QEventLoop, QTimer
from PyQt5.QtWidgets import QApplication, QFileDialog, QMessageBox

from app.core.models import OCRResult, OCRTable, TableCell
from app.ui.main_window import MainWindow


@unittest.skipUnless(os.name == "nt" and os.environ.get("TDP_RUN_WORD_COM") == "1", "Set TDP_RUN_WORD_COM=1 on Windows to run real Word COM integration.")
class WordComIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication(["tool-documents-pro-word-com-test"])

    def test_word_com_pdf_preview_save_as_and_session_cleanup(self) -> None:
        import psutil
        import pymupdf
        from docx import Document

        before_word_pids = {
            process.pid for process in psutil.process_iter(["name"])
            if (process.info.get("name") or "").lower() == "winword.exe"
        }
        window = MainWindow()
        result = OCRResult(
            full_text="Biên bản kiểm tra\nTiêu đề ghép | Cột B\nNội dung trái | Nội dung phải",
            tables=[OCRTable(
                grid=[["Tiêu đề ghép", "Cột B"], ["Nội dung trái", "Nội dung phải"]],
                num_rows=2,
                num_cols=2,
                cells=[
                    TableCell(0, 0, 1, 2, "Tiêu đề ghép"),
                    TableCell(1, 0, 2, 1, "Nội dung trái"),
                    TableCell(1, 1, 2, 2, "Nội dung phải"),
                ],
            )],
        )
        loop = QEventLoop()
        preview_failures: list[str] = []
        window.word_preview.preview_ready.connect(lambda _count: loop.quit())
        window.word_preview.preview_failed.connect(lambda message: (preview_failures.append(message), loop.quit()))
        # Exercise the same OCR completion callback that automatically creates
        # the temporary DOCX and starts Word rendering in the application.
        window._on_ocr_finished(result)
        worker = window._document_worker
        self.assertIsNotNone(worker)
        worker.wait(180_000)
        self.assertFalse(worker.isRunning(), "Word conversion worker did not finish")
        QApplication.processEvents()
        session = window._document_session
        self.assertIsNotNone(session)
        self.assertTrue(session.docx_path.is_file())
        self.assertTrue(session.pdf_path.is_file())
        self.assertTrue(window.btn_export_docx.isEnabled())

        QTimer.singleShot(30_000, loop.quit)
        loop.exec_()
        QApplication.processEvents()
        self.assertFalse(preview_failures, preview_failures)
        self.assertEqual(window.word_preview.page_label.text(), "Trang 1 / 1")
        self.assertTrue(window.word_preview._render_cache)
        self.assertTrue(window.word_preview.wait_for_idle(10_000))

        document = Document(session.docx_path)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(0, 0).text, "Tiêu đề ghép")
        self.assertIs(document.tables[0].cell(0, 0)._tc, document.tables[0].cell(0, 1)._tc)
        with pymupdf.open(session.pdf_path) as pdf:
            self.assertGreater(pdf.page_count, 0)
            self.assertGreater(pdf[0].get_pixmap().width, 0)
            page_text = pdf[0].get_text()
            self.assertIn("Tiêu đề ghép", page_text)
            self.assertIn("Nội dung phải", page_text)

        with TemporaryDirectory() as target_dir:
            target = Path(target_dir) / "saved-copy.docx"
            with patch.object(QFileDialog, "getSaveFileName", return_value=(str(target), "Word Documents (*.docx)")), patch.object(QMessageBox, "information"):
                window._on_export_docx()
            saved = Document(target)
            self.assertEqual(len(saved.tables), 1)
            self.assertEqual(saved.tables[0].cell(0, 0).text, "Tiêu đề ghép")
        # This is the same owned-session invalidation used when a new image is loaded.
        window._invalidate_document()
        self.assertFalse(session.directory.exists())
        window.close()

        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            after_word_pids = {
                process.pid for process in psutil.process_iter(["name"])
                if (process.info.get("name") or "").lower() == "winword.exe"
            }
            if after_word_pids <= before_word_pids:
                break
            time.sleep(0.2)
        self.assertLessEqual(after_word_pids, before_word_pids, "A Word process created by the test remained running")


if __name__ == "__main__":
    unittest.main()
