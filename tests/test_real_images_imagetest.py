#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
tests.test_real_images_imagetest
================================
Bộ kiểm thử E2E trên 4 ảnh thực tế trong assets/imagetest/:
1. bang1.png - Bảng điểm đánh giá chất lượng sinh hoạt chi bộ (5 cột)
2. bang2.png - Mẫu số 02 (16 cột có Merge Cell 2D rowspan=2, colspan=2, colspan=7)
3. vanban1.png - Biên bản sinh hoạt chi bộ định kỳ tháng 10 năm 2026
4. vanban2.png - Biên bản lấy ý kiến tập thể lao động Công ty TNHH Dịch Vụ Kế Toán Thiên Ưng
"""

import os
from pathlib import Path

import cv2
import pytest

from app.core.docx_exporter import export_to_docx
from app.core.ocr_engine import OCREngine
from app.core.spell_corrector import VietnameseSpellChecker

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets" / "imagetest"


@pytest.fixture(scope="module")
def ocr_engine():
    """Khởi tạo OCREngine 1 lần duy nhất cho module test."""
    return OCREngine(lang="vi", use_gpu=False, det_db_box_thresh=0.25)


@pytest.fixture(scope="module")
def spell_checker():
    """Khởi tạo VietnameseSpellChecker."""
    return VietnameseSpellChecker()


class TestRealImagesImageTestE2E:
    """Bộ kiểm thử E2E trên 4 ảnh thực tế."""

    def test_01_bang1_scorecard(self, ocr_engine: OCREngine, spell_checker: VietnameseSpellChecker, tmp_path: Path):
        """Kiểm tra bang1.png: bảng điểm 5 cột, đổi 1% -> ½, dòng con bắt đầu bằng +, dòng cuối * Trừ trường hợp."""
        img_path = ASSETS_DIR / "bang1.png"
        assert img_path.exists(), f"Không tìm thấy file: {img_path}"

        img = cv2.imread(str(img_path))
        assert img is not None

        res = ocr_engine.recognize(img)
        assert len(res.boxes) > 0
        assert len(res.full_text) > 0

        # Kiểm tra bảng được phát hiện
        assert len(res.tables) >= 1
        tbl = res.tables[0]
        assert tbl.num_cols == 5

        # Kiểm tra xuất docx thành công
        docx_out = tmp_path / "bang1.docx"
        assert export_to_docx(res, docx_out) is True
        assert docx_out.exists()

        # Kiểm tra nội dung file Word (.docx)
        import docx  # type: ignore[import-untyped,import-not-found]
        doc = docx.Document(str(docx_out))
        assert len(doc.tables) == 1, f"Kỳ vọng 1 bảng, nhận được {len(doc.tables)}"
        para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for pt in para_texts:
            assert " | " not in pt, f"Đoạn văn không được chứa dấu '|': {pt}"
        assert any("Trừ trường hợp" in pt for pt in para_texts), "Phải có đoạn văn chứa 'Trừ trường hợp'"

        # Kiểm tra quy tắc văn bản
        full_text = res.full_text
        assert "½" in full_text, "Phải chứa phân số ½ thay vì 1% hoặc %"
        assert "* Trừ trường hợp" in full_text or "*Trừ trường hợp" in full_text or "Trừ trường hợp" in full_text

    def test_02_bang2_mau_so_02_complex_table(self, ocr_engine: OCREngine, spell_checker: VietnameseSpellChecker, tmp_path: Path):
        """Kiểm tra bang2.png: Mẫu số 02 có 16 cột, Merge cell 2D, Header và Danh sách đứng trước bảng, Ghi chú ở cuối."""
        img_path = ASSETS_DIR / "bang2.png"
        assert img_path.exists(), f"Không tìm thấy file: {img_path}"

        img = cv2.imread(str(img_path))
        assert img is not None

        res = ocr_engine.recognize(img)
        assert len(res.boxes) > 0
        assert len(res.full_text) > 0

        # Kiểm tra xuất docx thành công
        docx_out = tmp_path / "bang2.docx"
        assert export_to_docx(res, docx_out) is True
        assert docx_out.exists()

        import docx  # type: ignore[import-untyped,import-not-found]
        from docx.enum.section import WD_ORIENT  # type: ignore[import-untyped,import-not-found]
        doc = docx.Document(str(docx_out))
        assert doc.sections[0].orientation == WD_ORIENT.LANDSCAPE, "Bảng 16 cột phải xoay khổ ngang Landscape"
        assert len(doc.tables) >= 1
        para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for pt in para_texts:
            assert " | " not in pt, f"Đoạn văn không được chứa dấu '|': {pt}"

        full_text = res.full_text
        # Header / Title phải đứng trước tiêu đề bảng
        assert "DANH SÁCH" in full_text
        idx_ds = full_text.index("DANH SÁCH")
        if "Họ tên" in full_text:
            idx_ht = full_text.index("Họ tên")
            assert idx_ds < idx_ht, "DANH SÁCH phải đứng trước tiêu đề bảng"

        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-untyped,import-not-found]
        if len(doc.tables) >= 2:
            hdr_table = doc.tables[0]
            for p in hdr_table.cell(0, 0).paragraphs:
                if p.text.strip():
                    assert p.alignment == WD_ALIGN_PARAGRAPH.LEFT, "Khối trái header Mẫu số 02 phải căn trái"
            for p in hdr_table.cell(0, 1).paragraphs:
                if p.text.strip():
                    assert p.alignment == WD_ALIGN_PARAGRAPH.RIGHT, "Khối phải header Mẫu số 02 phải căn phải"
            main_table = doc.tables[1]
            col7_texts = [main_table.cell(r, 6).text.strip() for r in range(len(main_table.rows))]
            assert "(7)" in col7_texts, f"Cột 7 phải chứa (7), nhận được: {col7_texts}"
            assert any("(3=4+5)" in c.text for row in main_table.rows for c in row.cells), "Bảng phải chứa công thức (3=4+5)"
            assert any("người thực hiện nghi lễ truyền thống, tín ngưỡng" in c.text for row in main_table.rows for c in row.cells), "Cột (9) phải khôi phục đầy đủ cụm từ"

        # Kiểm tra công thức và cụm từ trong full_text
        assert "(3=4+5)" in full_text, "full_text phải chứa (3=4+5) thay vì (3=445)"
        assert "người thực hiện nghi lễ truyền thống, tín ngưỡng" in full_text, "Cột (9) trong full_text phải khôi phục đầy đủ"
        assert "Từ Cột (9) đến cột (15):" in full_text, "Ghi chú phải sửa thành Từ Cột (9) đến cột (15):"

        # Kiểm tra Mẫu số 02 nằm trên cùng bên phải
        if doc.paragraphs:
            assert doc.paragraphs[0].text.strip() == "Mẫu số 02"
            assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    def test_03_vanban1_bien_ban_sinh_hoat_chi_bo(self, ocr_engine: OCREngine, spell_checker: VietnameseSpellChecker, tmp_path: Path):
        """Kiểm tra vanban1.png: Biên bản sinh hoạt chi bộ định kỳ tháng 10 năm 2026."""
        import re
        img_path = ASSETS_DIR / "vanban1.png"
        assert img_path.exists(), f"Không tìm thấy file: {img_path}"

        img = cv2.imread(str(img_path))
        assert img is not None

        res = ocr_engine.recognize(img)
        assert len(res.boxes) > 0
        assert len(res.full_text) > 0

        full_text = res.full_text
        assert "BIÊN BẢN" in full_text.upper() or "SINH HOẠT" in full_text.upper()
        assert "CHI BỘ" in full_text.upper()

        docx_out = tmp_path / "vanban1.docx"
        assert export_to_docx(res, docx_out) is True
        assert docx_out.exists()

        import docx  # type: ignore[import-untyped,import-not-found]
        doc = docx.Document(str(docx_out))
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt.startswith("6. Đồng chí Lê Văn Điền"):
                assert not any(r.bold for r in p.runs), f"Mục con 6. không được in đậm: {txt}"
            if re.match(r"^(?:I|II)\.\s+", txt):
                assert any(r.bold for r in p.runs), f"Đề mục La Mã phải được in đậm: {txt}"

    def test_04_vanban2_bien_ban_thien_ung(self, ocr_engine: OCREngine, spell_checker: VietnameseSpellChecker, tmp_path: Path):
        """Kiểm tra vanban2.png: Biên bản Công ty Kế Toán Thiên Ưng."""
        img_path = ASSETS_DIR / "vanban2.png"
        assert img_path.exists(), f"Không tìm thấy file: {img_path}"

        img = cv2.imread(str(img_path))
        assert img is not None

        res = ocr_engine.recognize(img)
        assert len(res.boxes) > 0
        assert len(res.full_text) > 0

        full_text = res.full_text
        assert "THIÊN ƯNG" in full_text.upper() or "KẾ TOÁN" in full_text.upper()

        docx_out = tmp_path / "vanban2.docx"
        assert export_to_docx(res, docx_out) is True
        assert docx_out.exists()
