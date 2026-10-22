from __future__ import annotations

import pytest
from docx import Document

from app.core.docx_exporter import export_to_docx
from app.core.models import DocumentBlock, OCRBox, OCRResult, OCRTable, ParallelHeader, TableCell
from app.core.structured_editing import EditOwnershipError, StructuredDocumentEditor


def box(text: str, **metadata) -> OCRBox:
    return OCRBox(text=text, confidence=0.9, bbox=[[0, 0], [20, 0], [20, 10], [0, 10]],
                  raw_text=f"raw:{text}", metadata=metadata)


def structured_result() -> OCRResult:
    before = box("Title")
    left = box("Group")
    right = box("Name", candidates=[{"recognizer": "vietocr", "text": "Name"},
                                     {"recognizer": "paddleocr", "text": "Label",
                                      "mapping": {"status": "matched"}}])
    count = box("Count")
    a = box("A")
    b = box("B")
    five = box("5")
    after = box("End")
    table = OCRTable(
        grid=[["Group Name", "", "Count"], ["A", "B", "5"]],
        bbox=[0, 20, 100, 60], num_rows=2, num_cols=3,
        cells=[
            TableCell(0, 0, 1, 2, "Group Name", [0, 20, 60, 40], [left, right]),
            TableCell(0, 2, 1, 3, "Count", [60, 20, 100, 40], [count]),
            TableCell(1, 0, 2, 1, "A", [0, 40, 30, 60], [a]),
            TableCell(1, 1, 2, 2, "B", [30, 40, 60, 60], [b]),
            TableCell(1, 2, 2, 3, "5", [60, 40, 100, 60], [five]),
        ],
        column_boundaries=[0, 30, 60, 100], row_boundaries=[20, 40, 60],
    )
    return OCRResult(
        boxes=[before, left, right, count, a, b, five, after],
        tables=[table], full_text="Title\nGroup Name |  | Count\nA | B | 5\nEnd",
        document_blocks=[
            DocumentBlock("heading", 0, "Title", "raw:Title", [0], metadata={"format": {"bold": True}}),
            DocumentBlock("table", 1, table_index=0, box_indices=[1, 2, 3, 4, 5, 6]),
            DocumentBlock("paragraph", 2, "End", "raw:End", [7]),
        ], raw_full_text="raw source",
    )


def test_paragraph_and_heading_corrections_keep_source_and_formatting():
    source = structured_result()
    editor = StructuredDocumentEditor(source)
    editor.apply_box_correction(0, "Confirmed title")
    edited = editor.result
    assert edited.document_blocks[0].text == "Confirmed title"
    assert edited.document_blocks[0].metadata["format"]["bold"] is True
    assert edited.boxes[0].text == "Confirmed title"
    assert edited.boxes[0].raw_text == "raw:Title"
    assert edited.raw_full_text == "raw source"
    assert source.boxes[0].text == "Title"


def test_multi_box_table_cell_updates_cell_grid_and_full_text_without_topology_change():
    editor = StructuredDocumentEditor(structured_result())
    before = editor.result.tables[0]
    editor.apply_box_correction(2, "Label")
    edited = editor.result
    cell = edited.tables[0].cells[0]
    assert cell.text == "Group Label"
    assert edited.tables[0].grid[0][0] == "Group Label"
    assert edited.full_text == "Title\nGroup Label |  | Count\nA | B | 5\nEnd"
    assert (edited.tables[0].num_rows, edited.tables[0].num_cols) == (2, 3)
    assert [(c.row_start, c.col_start, c.row_end, c.col_end) for c in edited.tables[0].cells] == [
        (c.row_start, c.col_start, c.row_end, c.col_end) for c in before.cells
    ]
    assert edited.boxes[2].raw_text == "raw:Name"


def test_confirmed_candidate_is_applied_only_after_explicit_call():
    editor = StructuredDocumentEditor(structured_result())
    assert editor.result.boxes[2].text == "Name"
    editor.apply_candidate(2, 1)
    assert editor.result.boxes[2].text == "Label"
    assert editor.result.boxes[2].metadata["candidates"][0]["text"] == "Name"
    assert editor.result.boxes[2].metadata["candidates"][1]["text"] == "Label"
    assert editor.corrections[-1].source == "confirmed_candidate"
    assert editor.corrections[-1].candidate_index == 1
    assert editor.corrections[-1].candidate_recognizer == "paddleocr"
    correction = editor.result.audit_trail[-1]
    assert correction["kind"] == "user_correction"
    assert correction["candidate_recognizer"] == "paddleocr"


def test_candidate_selection_rejects_missing_metadata_or_index():
    editor = StructuredDocumentEditor(structured_result())
    with pytest.raises(IndexError):
        editor.apply_candidate(0, 0)
    with pytest.raises(IndexError):
        editor.apply_candidate(2, 4)
    assert editor.result.boxes[2].text == "Name"


def test_alternate_candidate_with_ambiguous_mapping_cannot_be_applied():
    result = structured_result()
    result.boxes[2].metadata["candidates"][1]["mapping"] = {"status": "ambiguous"}
    editor = StructuredDocumentEditor(result)
    with pytest.raises(ValueError, match="verified region mapping"):
        editor.apply_candidate(2, 1)
    assert editor.result.boxes[2].text == "Name"


def test_parallel_header_edit_updates_header_source_and_ordered_block():
    left = box("Agency")
    right = box("Republic")
    result = OCRResult(
        boxes=[left, right],
        full_text="Agency\nRepublic",
        parallel_header=ParallelHeader(left_lines=["Agency"], right_lines=["Republic"],
                                       left_boxes=[left], right_boxes=[right]),
        document_blocks=[DocumentBlock("parallel_header", 0, "Agency\nRepublic", "", [0, 1])],
    )
    editor = StructuredDocumentEditor(result)
    editor.apply_box_correction(0, "Updated agency")
    edited = editor.result
    assert edited.parallel_header.left_lines == ["Updated agency"]
    assert edited.parallel_header.right_lines == ["Republic"]
    assert edited.document_blocks[0].text == "Updated agency\nRepublic"
    assert edited.full_text == "Updated agency\nRepublic"


def test_ambiguous_ownership_does_not_mutate_any_representation():
    result = structured_result()
    result.document_blocks.append(DocumentBlock("paragraph", 3, "Title again", box_indices=[0]))
    editor = StructuredDocumentEditor(result)
    with pytest.raises(EditOwnershipError, match="ambiguous"):
        editor.apply_box_correction(0, "Changed")
    assert editor.result.boxes[0].text == "Title"
    assert editor.result.tables[0].grid[0][0] == "Group Name"
    assert not editor.dirty


def test_undo_redo_and_sequential_edits_restore_exact_document_state():
    editor = StructuredDocumentEditor(structured_result())
    editor.apply_box_correction(0, "Heading updated")
    editor.apply_box_correction(2, "Label")
    assert len(editor.corrections) == 2
    assert editor.undo().target.box_index == 2
    assert editor.result.boxes[2].text == "Name"
    assert editor.result.boxes[0].text == "Heading updated"
    assert editor.redo().target.box_index == 2
    assert editor.result.boxes[2].text == "Label"
    editor.undo()
    editor.undo()
    assert not editor.dirty
    assert editor.result.full_text == structured_result().full_text


def test_docx_export_preserves_table_merge_format_and_only_confirmed_text(tmp_path):
    editor = StructuredDocumentEditor(structured_result())
    editor.apply_box_correction(2, "Label")
    result = editor.result
    path = tmp_path / "edited.docx"
    assert export_to_docx(result, str(path))
    document = Document(path)
    assert [p.text for p in document.paragraphs if p.text.strip()] == ["Title", "End"]
    assert len(document.tables) == 1
    word_table = document.tables[0]
    assert len(word_table.rows) == 2
    assert len(word_table.columns) == 3
    assert word_table.cell(0, 0).text == "Group Label"
    assert word_table.cell(0, 0)._tc is word_table.cell(0, 1)._tc
    assert word_table.cell(1, 2).text == "5"
    assert document.paragraphs[0].runs[0].bold is True
    assert "Name" not in "\n".join([p.text for p in document.paragraphs] + [word_table.cell(0, 0).text])
