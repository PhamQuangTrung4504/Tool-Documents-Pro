#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
tests.test_c5_physical_grid_merged_cells
========================================
Kiểm thử Phase C5: Physical Grid Cell Topology Engine
1. Dựng mẫu bảng phức tạp mô phỏng Mẫu số 02 (image_441e44.jpg) với 16 cột:
   - Các cột rowspan=2: TT, Họ tên NCUT, Dân tộc, Ghi chú,...
   - Cột colspan=2: Năm sinh -> Nam | Nữ
   - Cột colspan=7: Thành phần người có uy tín -> 7 cột con
   - Ô dọc nhiều dòng: Cột 12 (9 dòng rời), Cột 13 (10 dòng rời) gom thành 1 dòng duy nhất.
2. Kiểm tra PhysicalGridTopologyEngine và extract_physical_table_topology.
3. Kiểm tra export_to_docx xuất bảng có TableCell gộp (is_merged=True).
4. Kiểm tra final_document_cleanup chuyển 1% -> ½ và giữ nguyên * Trừ trường hợp.
"""

from __future__ import annotations

import os
import zipfile
from pathlib import Path
from typing import List, Tuple

import cv2
import numpy as np
from app.core.docx_exporter import export_to_docx
from app.core.layout_analyzer import reconstruct_document_layout
from app.core.models import OCRBox, OCRResult, OCRTable
from app.core.table_engines import (
    PhysicalGridTopologyEngine,
    extract_physical_table_topology,
    split_table_crossing_boxes,
)
from app.core.text_postprocessor import final_document_cleanup


def _create_box(text: str, x1: int, y1: int, x2: int, y2: int, conf: float = 0.99) -> OCRBox:
    """Tạo OCRBox với 4 điểm góc."""
    bbox = [
        [int(x1), int(y1)],
        [int(x2), int(y1)],
        [int(x2), int(y2)],
        [int(x1), int(y2)],
    ]
    return OCRBox(text=text, confidence=conf, bbox=bbox, raw_text=text)


def build_synthetic_mau_so_02() -> Tuple[np.ndarray, List[OCRBox]]:
    """
    Tạo ảnh tổng hợp và danh sách OCRBox mô phỏng chính xác biểu mẫu Mẫu số 02 (16 cột).
    """
    img = np.ones((700, 1020, 3), dtype=np.uint8) * 255
    x_cols = [50, 90, 180, 250, 290, 330, 380, 430, 500, 560, 620, 680, 750, 820, 870, 920, 970]
    y_rows = [150, 200, 380, 420, 480, 540]

    # 1. Vẽ khung chữ nhật bao quanh bảng
    cv2.rectangle(img, (x_cols[0], y_rows[0]), (x_cols[-1], y_rows[-1]), (0, 0, 0), 2)

    # 2. Vẽ các đường kẻ ngang
    # Y1 = 200: Vạch phân cấp cho Năm sinh (cột 3..5) và Thành phần người có uy tín (cột 8..15)
    cv2.line(img, (x_cols[3], y_rows[1]), (x_cols[5], y_rows[1]), (0, 0, 0), 2)
    cv2.line(img, (x_cols[8], y_rows[1]), (x_cols[15], y_rows[1]), (0, 0, 0), 2)

    # Y2, Y3, Y4: Các đường kẻ ngang xuyên suốt toàn bộ chiều rộng bảng
    for y in y_rows[2:5]:
        cv2.line(img, (x_cols[0], y), (x_cols[-1], y), (0, 0, 0), 2)

    # 3. Vẽ các đường kẻ dọc
    for i, x in enumerate(x_cols[1:-1], 1):
        if i in [4, 9, 10, 11, 12, 13, 14]:
            # Các cột con của Năm sinh (Nam/Nữ) và Thành phần NCUT: vạch kẻ dọc bắt đầu từ Y1 = 200
            cv2.line(img, (x, y_rows[1]), (x, y_rows[-1]), (0, 0, 0), 2)
        else:
            # Các cột chính (rowspan = 2): vạch kẻ dọc xuyên suốt từ đỉnh bảng Y0 = 150
            cv2.line(img, (x, y_rows[0]), (x, y_rows[-1]), (0, 0, 0), 2)

    boxes: List[OCRBox] = []

    # Tiêu đề trên bảng (Non-table lines)
    boxes.append(_create_box("ỦY BAN NHÂN DÂN XÃ EA H'LEO", 50, 10, 300, 30))
    boxes.append(_create_box("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", 600, 10, 950, 30))
    boxes.append(_create_box("DANH SÁCH", 450, 40, 570, 65))
    boxes.append(_create_box("Công nhận người có uy tín trong vùng đồng bào dân tộc thiểu số", 220, 75, 800, 95))

    # Hàng 0 - Tiêu đề tầng 1
    boxes.append(_create_box("TT", 60, 160, 80, 185))
    boxes.append(_create_box("Họ tên NCUT", 95, 160, 175, 185))
    boxes.append(_create_box("Số lượng NCUT (3=4+5)", 185, 160, 245, 185))
    boxes.append(_create_box("Năm sinh", 260, 165, 320, 185))
    boxes.append(_create_box("Dân tộc", 335, 160, 375, 185))
    boxes.append(_create_box("Nơi cư trú (thôn)", 385, 160, 425, 185))
    boxes.append(_create_box("Trình độ học vấn/chuyên môn", 435, 160, 495, 185))
    boxes.append(_create_box("Thành phần người có uy tín", 600, 165, 800, 185))
    boxes.append(_create_box("Ghi chú", 925, 160, 965, 185))

    # Hàng 1 - Tiêu đề tầng 2 (các cột con)
    boxes.append(_create_box("Nam", 260, 220, 280, 240))
    boxes.append(_create_box("Nữ", 300, 220, 320, 240))
    boxes.append(_create_box("Già làng, trưởng dòng họ", 505, 220, 555, 240))
    boxes.append(_create_box("Chức sắc, chức việc", 565, 220, 615, 240))
    boxes.append(_create_box("Cán bộ, công chức", 625, 220, 675, 240))

    # Cột 11 (Cột 12 thực tế): 9 dòng chữ rời nhau xếp dọc trong ô
    cot_12_lines = [
        "Bí thư chi bộ,",
        "Trưởng thôn,",
        "Trưởng ban",
        "công tác mặt",
        "trận thôn, Tổ",
        "trưởng Tổ bảo",
        "vệ an ninh, trật",
        "tự, Thôn đội",
        "trưởng",
    ]
    for idx, l in enumerate(cot_12_lines):
        y_l = 205 + idx * 18
        boxes.append(_create_box(l, 685, y_l, 745, y_l + 16))

    # Cột 12 (Cột 13 thực tế): 10 dòng chữ rời nhau xếp dọc trong ô
    cot_13_lines = [
        "Nhân sỹ, trí",
        "thức, doanh",
        "nhân, nghệ",
        "nhân, nhà",
        "giáo, thầy",
        "thuốc, lương",
        "y, công nhân,",
        "nông dân,",
        "phụ nữ,",
        "thanh niên",
    ]
    for idx, l in enumerate(cot_13_lines):
        y_l = 205 + idx * 17
        boxes.append(_create_box(l, 755, y_l, 815, y_l + 15))

    boxes.append(_create_box("Đảng viên", 825, 220, 865, 240))
    boxes.append(_create_box("Thành phần khác", 875, 220, 915, 240))

    # Hàng 2 - Đánh số thứ tự cột (1) đến (16)
    for i in range(16):
        num_str = f"({i + 1})"
        boxes.append(_create_box(num_str, x_cols[i] + 5, 390, x_cols[i + 1] - 5, 410))

    # Hàng 3 - Dòng dữ liệu 1
    boxes.append(_create_box("1", 65, 440, 75, 460))
    boxes.append(_create_box("Nguyễn Văn A", 95, 440, 175, 460))
    boxes.append(_create_box("1", 210, 440, 220, 460))
    boxes.append(_create_box("1", 265, 440, 275, 460))
    boxes.append(_create_box("Kinh", 345, 440, 370, 460))
    boxes.append(_create_box("Thôn 1", 390, 440, 420, 460))
    boxes.append(_create_box("Đại học", 440, 440, 480, 460))

    # Hàng 4 - Dòng dữ liệu 2
    boxes.append(_create_box("2", 65, 500, 75, 520))
    boxes.append(_create_box("Trần Thị B", 95, 500, 165, 520))
    boxes.append(_create_box("1", 210, 500, 220, 520))
    boxes.append(_create_box("1", 305, 500, 315, 520))
    boxes.append(_create_box("Tày", 345, 500, 365, 520))
    boxes.append(_create_box("Thôn 2", 390, 500, 420, 520))
    boxes.append(_create_box("12/12", 445, 500, 475, 520))

    # Dòng ghi chú dưới bảng
    boxes.append(_create_box("Ghi chú:", 50, 560, 100, 580))
    boxes.append(_create_box("- Cột (8): Ghi cụ thể trình độ học vấn/chuyên môn.", 50, 585, 450, 605))

    return img, boxes


class TestC5PhysicalGridTopologyEngine:
    """Bộ kiểm thử chuyên sâu cho Phase C5."""

    def test_01_detect_canonical_grid_and_merged_cells(self):
        """Xác minh nhận diện chính xác 16 cột, 5 hàng và các ô gộp 2D."""
        img, boxes = build_synthetic_mau_so_02()
        tbl = extract_physical_table_topology(img, boxes)

        assert tbl is not None, "Phải phát hiện được cấu trúc bảng vật lý"
        assert tbl.num_cols == 16, f"Bảng phải có đúng 16 cột, thực tế: {tbl.num_cols}"
        assert tbl.num_rows == 5, f"Bảng phải có đúng 5 hàng, thực tế: {tbl.num_rows}"

        # 1. Kiểm tra ô Năm sinh (colspan = 2, rowspan = 1)
        nam_sinh = next((c for c in tbl.cells if "Năm sinh" in c.text), None)
        assert nam_sinh is not None, "Phải tìm thấy ô Năm sinh"
        assert nam_sinh.col_span == 2, f"Năm sinh phải có col_span=2, thực tế: {nam_sinh.col_span}"
        assert nam_sinh.row_span == 1, f"Năm sinh phải có row_span=1, thực tế: {nam_sinh.row_span}"
        assert nam_sinh.is_merged is True

        # 2. Kiểm tra ô Thành phần người có uy tín (colspan = 7, rowspan = 1)
        thanh_phan = next((c for c in tbl.cells if "Thành phần" in c.text), None)
        assert thanh_phan is not None, "Phải tìm thấy ô Thành phần người có uy tín"
        assert thanh_phan.col_span == 7, f"Thành phần NCUT phải có col_span=7, thực tế: {thanh_phan.col_span}"
        assert thanh_phan.row_span == 1, f"Thành phần NCUT phải có row_span=1, thực tế: {thanh_phan.row_span}"
        assert thanh_phan.is_merged is True

        # 3. Kiểm tra các ô gộp dọc (rowspan = 2): TT, Họ tên NCUT, Ghi chú
        tt_cell = next((c for c in tbl.cells if c.row_start == 0 and c.col_start == 0), None)
        assert tt_cell is not None
        assert "TT" in tt_cell.text
        assert tt_cell.row_span == 2, f"Cột TT phải gộp dọc rowspan=2, thực tế: {tt_cell.row_span}"
        assert tt_cell.col_span == 1

        ho_ten = next((c for c in tbl.cells if "Họ tên" in c.text), None)
        assert ho_ten is not None
        assert ho_ten.row_span == 2, f"Cột Họ tên phải gộp dọc rowspan=2, thực tế: {ho_ten.row_span}"

    def test_open_table_suffix_is_kept_inside_table_and_partitioned_by_text_geometry(self):
        """Continue a clipped ruled table through its visible, unruled suffix."""
        image = np.full((300, 400, 3), 255, dtype=np.uint8)
        x_edges = [20, 80, 200, 300, 380]
        y_edges = [50, 100, 140, 180]
        for y in y_edges:
            cv2.line(image, (x_edges[0], y), (x_edges[-1], y), (0, 0, 0), 2)
        for x in x_edges:
            cv2.line(image, (x, y_edges[0]), (x, image.shape[0] - 1), (0, 0, 0), 2)

        boxes = [
            _create_box("Header", 30, 60, 70, 80),
            _create_box("Column", 100, 60, 180, 80),
            _create_box("Existing row", 100, 110, 180, 130),
            _create_box("A-17", 38, 202, 65, 218),
            _create_box("First item", 100, 202, 170, 218),
            _create_box("B-28", 38, 242, 65, 258),
            _create_box("Second item", 100, 242, 180, 258),
            _create_box("C-39", 38, 272, 65, 288),
            _create_box("Third item", 100, 272, 175, 288),
        ]

        table = extract_physical_table_topology(image, boxes)
        assert table is not None
        assert table.bbox[3] == image.shape[0]
        assert table.num_rows == 5
        assert table.grid[-3][0] == "A-17"
        assert table.grid[-3][1] == "First item"
        assert table.grid[-2][0] == "B-28"
        assert table.grid[-2][1] == "Second item"
        assert table.grid[-1][0] == "C-39"
        assert table.grid[-1][1] == "Third item"

    def test_diagonal_cell_border_is_preserved_in_docx_ooxml(self, tmp_path: Path):
        image = np.full((200, 240, 3), 255, dtype=np.uint8)
        x_edges = [20, 120, 220]
        y_edges = [20, 100, 180]
        for y in y_edges:
            cv2.line(image, (x_edges[0], y), (x_edges[-1], y), (0, 0, 0), 2)
        for x in x_edges:
            cv2.line(image, (x, y_edges[0]), (x, y_edges[-1]), (0, 0, 0), 2)
        cv2.line(image, (x_edges[0], y_edges[0]), (x_edges[1], y_edges[1]), (0, 0, 0), 2)
        boxes = [
            _create_box("Top left", 35, 40, 85, 60),
            _create_box("Top right", 145, 40, 195, 60),
            _create_box("Bottom left", 35, 130, 85, 150),
            _create_box("Bottom right", 145, 130, 195, 150),
        ]

        table = extract_physical_table_topology(image, boxes)
        assert table is not None
        # Morphological line centers may move by a few pixels depending on
        # stroke thickness; the synthetic source edges are the ground truth.
        diagonal_cell = next(cell for cell in table.cells if cell.row_start == 0 and cell.col_start == 0)
        assert all(abs(actual - expected) <= 3 for actual, expected in zip(diagonal_cell.bbox, [20, 20, 120, 100]))
        assert diagonal_cell.diagonal_border == "tl2br"

        output = tmp_path / "diagonal.docx"
        result = OCRResult(
            boxes=boxes,
            full_text="Top left | Top right\nBottom left | Bottom right",
            tables=[table],
        )
        assert export_to_docx(result, str(output))
        with zipfile.ZipFile(output) as archive:
            document_xml = archive.read("word/document.xml")
        assert b"w:tl2br" in document_xml

    def test_legacy_export_without_ownership_preserves_text_and_reports_ambiguity(self, tmp_path: Path):
        table = OCRTable(
            grid=[["Name", "Amount"], ["Paper", "10"]],
            bbox=[20, 40, 380, 150],
            num_rows=2,
            num_cols=2,
        )
        result = OCRResult(
            full_text="Report title\nName Amount\nPaper 10\nEnd note",
            tables=[table],
        )
        output = tmp_path / "structured.docx"
        assert export_to_docx(result, str(output))
        document = __import__("docx").Document(output)
        assert len(document.tables) == 0
        assert [paragraph.text for paragraph in document.paragraphs if paragraph.text] == [
            "Report title", "Name Amount", "Paper 10", "End note"
        ]

    def test_02_multiline_intra_cell_aggregation_never_fragments_rows(self):
        """Xác minh 9 dòng rời và 10 dòng rời trong 1 ô được gom thành 1 dòng duy nhất."""
        img, boxes = build_synthetic_mau_so_02()
        tbl = extract_physical_table_topology(img, boxes)
        assert tbl is not None

        # Kiểm tra Cột 12 (Bí thư chi bộ...): gồm 9 dòng nhỏ nằm trọn trong 1 ô
        cot_12 = next((c for c in tbl.cells if "Bí thư chi bộ" in c.text), None)
        assert cot_12 is not None, "Phải tìm thấy ô Cột 12"
        expected_cot_12 = "Bí thư chi bộ, Trưởng thôn, Trưởng ban công tác mặt trận thôn, Tổ trưởng Tổ bảo vệ an ninh, trật tự, Thôn đội trưởng"
        assert cot_12.text == expected_cot_12, f"9 dòng chữ phải được nối thành 1 dòng duy nhất. Thực tế: {cot_12.text}"

        # Kiểm tra Cột 13 (Nhân sỹ, trí thức...): gồm 10 dòng nhỏ nằm trọn trong 1 ô
        cot_13 = next((c for c in tbl.cells if "Nhân sỹ" in c.text), None)
        assert cot_13 is not None, "Phải tìm thấy ô Cột 13"
        expected_cot_13 = "Nhân sỹ, trí thức, doanh nhân, nghệ nhân, nhà giáo, thầy thuốc, lương y, công nhân, nông dân, phụ nữ, thanh niên"
        assert cot_13.text == expected_cot_13, f"10 dòng chữ phải được nối thành 1 dòng duy nhất. Thực tế: {cot_13.text}"

        # Toàn bộ bảng vẫn chỉ có 5 hàng, tuyệt đối không bị phân rã thành 10+ hàng
        assert len(tbl.grid) == 5
        assert tbl.num_rows == 5

    def test_03_physical_grid_topology_engine_layout_result(self):
        """Xác minh PhysicalGridTopologyEngine trả về LayoutResult hoàn chỉnh."""
        img, boxes = build_synthetic_mau_so_02()
        engine = PhysicalGridTopologyEngine()
        grouped_lines, tables = engine(boxes, image=img)

        assert len(tables) == 1
        tbl = tables[0]
        assert tbl.num_cols == 16
        assert tbl.num_rows == 5

        # Xác minh các dòng ngoài bảng được tách riêng (DANH SÁCH, Ghi chú)
        flat_lines = [" ".join(b.text for b in line) for line in grouped_lines]
        assert any("DANH SÁCH" in line for line in flat_lines)
        assert any("Ghi chú:" in line for line in flat_lines)

    def test_04_docx_export_with_merged_cells_success(self, tmp_path: Path):
        """Xác minh export_to_docx xuất bảng có TableCell gộp (is_merged=True) thành công."""
        img, boxes = build_synthetic_mau_so_02()
        tbl = extract_physical_table_topology(img, boxes)
        assert tbl is not None

        full_text = "DANH SÁCH NGƯỜI CÓ UY TÍN\n" + "\n".join(" | ".join(row) for row in tbl.grid)
        result = OCRResult(boxes=boxes, full_text=full_text, tables=[tbl])

        out_docx = str(tmp_path / "test_mau_so_02.docx")
        ok = export_to_docx(result, out_docx)
        assert ok is True, "export_to_docx phải thành công"
        assert os.path.exists(out_docx), "File Word phải tồn tại"
        assert os.path.getsize(out_docx) > 5000, "File Word phải có nội dung hợp lệ"

    def test_05_text_postprocessor_final_cleanup_regression(self):
        """Khóa chốt: 1% -> ½, giữ * Trừ trường hợp và 3 quy tắc văn bản hành chính."""
        s = (
            "- Thời gian sinh hoạt chi bộ chỉ bằng 1% so với quy định | 02\n"
            "- Thời gian sinh hoạt chi bộ dưới % so với quy định | 01\n"
            "*Trừ trường hợp được miễn công tác"
        )
        cleaned = final_document_cleanup(s)
        expected = (
            "+ Thời gian sinh hoạt chi bộ chỉ bằng ½ so với quy định | 02\n"
            "+ Thời gian sinh hoạt chi bộ dưới ½ so với quy định | 01\n"
            "* Trừ trường hợp được miễn công tác"
        )
        assert cleaned == expected

        # 3 quy tắc văn bản hành chính
        r1 = final_document_cleanup("Yêu cầu các cá nhân, cơ quan có trách nhiệm xây dựng")
        assert r1 == "- Yêu cầu các cá nhân, cơ quan có trách nhiệm xây dựng"

        r2 = final_document_cleanup("Tổng số cán bộ của doanh nghiệp: - 14...... người")
        assert r2 == "Tổng số cán bộ của doanh nghiệp: ....14...... người"

        r3 = final_document_cleanup("Công ty TNHH Dịch vụ Kế Toán Thiên\nƯng thông báo")
        assert r3 == "Công ty TNHH Dịch vụ Kế Toán Thiên Ưng thông báo"

    def test_06_reconstruct_document_layout_reading_order_with_image(self):
        """Xác minh reconstruct_document_layout sắp xếp đúng thứ tự: Header/Title -> Bảng -> Ghi chú."""
        img, boxes = build_synthetic_mau_so_02()
        grouped_lines, tables = reconstruct_document_layout(boxes, image=img)

        assert len(tables) == 1
        tbl = tables[0]
        assert tbl.num_cols == 16
        assert tbl.num_rows == 5

        line_texts = [" ".join(b.text for b in line) for line in grouped_lines]

        # 1. pre_lines: Ủy ban nhân dân và Danh sách phải đứng đầu
        idx_ubnd = next((i for i, t in enumerate(line_texts) if "ỦY BAN NHÂN DÂN" in t), -1)
        idx_ds = next((i for i, t in enumerate(line_texts) if "DANH SÁCH" in t), -1)
        assert idx_ubnd != -1, "Phải tìm thấy dòng Ủy ban nhân dân"
        assert idx_ds != -1, "Phải tìm thấy dòng DANH SÁCH"
        assert idx_ubnd <= idx_ds, "Ủy ban nhân dân phải đứng trước hoặc cùng dòng với DANH SÁCH"

        # 2. Table rows: dòng TT / Họ tên NCUT phải xuất hiện sau DANH SÁCH
        idx_tt = next((i for i, t in enumerate(line_texts) if "TT" in t and "Họ tên" in t), -1)
        assert idx_tt != -1, "Phải tìm thấy dòng tiêu đề bảng TT / Họ tên"
        assert idx_tt > idx_ds, f"Dòng tiêu đề bảng phải đứng sau DANH SÁCH (idx_tt={idx_tt}, idx_ds={idx_ds})"

        # 3. post_lines: Ghi chú phải đứng sau toàn bộ các hàng của bảng
        idx_ghichu = next((i for i, t in enumerate(line_texts) if "Ghi chú:" in t), -1)
        assert idx_ghichu != -1, "Phải tìm thấy dòng Ghi chú"
        assert idx_ghichu > idx_tt, f"Ghi chú phải đứng sau bảng (idx_ghichu={idx_ghichu}, idx_tt={idx_tt})"

    def test_07_split_table_crossing_boxes(self):
        """Xác minh split_table_crossing_boxes tách box vắt qua vạch kẻ dọc cột."""
        img, _ = build_synthetic_mau_so_02()
        # Tạo box vắt ngang qua vạch kẻ x = 500 giữa Cột 8 (Trình độ học vấn) và Cột 9 (Già làng)
        crossing_box = _create_box("Đại học Già làng", 435, 440, 555, 460)

        class MockVietOCR:
            def predict(self, *args, **kwargs):
                return ("SubText", 0.98)

        mock_split = split_table_crossing_boxes([crossing_box], image=img, vietocr_model=MockVietOCR())
        assert len(mock_split) >= 2, f"Box vắt qua vạch x=500 phải được chia thành ít nhất 2 box, thực tế: {len(mock_split)}"
        x2_sub0 = max(pt[0] for pt in mock_split[0].bbox)
        assert x2_sub0 <= 505

    def test_08_text_postprocessor_mau_so_02_recovery_rules(self):
        """Xác minh khôi phục dòng tiêu đề bị DBNet bỏ sót và sửa 'Bí t' -> 'Bí thư Chi bộ'."""
        # 1. Khôi phục subtitle dưới DANH SÁCH
        raw_text_1 = (
            "DANH SÁCH\n"
            "(Ban hành kèm theo Quyết định số 12/QĐ-UBND)\n"
            "TT | Họ tên NCUT | Năm sinh\n"
        )
        cleaned_1 = final_document_cleanup(raw_text_1)
        assert "Công nhận người có uy tín trong vùng đồng bào dân tộc thiểu số" in cleaned_1
        assert cleaned_1.index("DANH SÁCH") < cleaned_1.index("Công nhận người có uy tín") < cleaned_1.index("(Ban hành kèm theo")

        # 2. Khôi phục 'Bí t' ở cuối dòng
        raw_text_2 = "Trưởng ban công tác mặt trận thôn, Bí t"
        cleaned_2 = final_document_cleanup(raw_text_2)
        assert cleaned_2 == "Trưởng ban công tác mặt trận thôn, Bí thư Chi bộ"

    def test_09_cadastral_land_table_normalization_rules(self):
        """Xác minh chuẩn hóa các lỗi trên Bảng danh sách thửa đất (image_612a9c.jpg)."""
        assert final_document_cleanup("Diện tích (m?)") == "Diện tích (m²)"
        assert final_document_cleanup("Diện tích (m2)") == "Diện tích (m²)"
        assert final_document_cleanup("Số tờ 2 Chiner") == "Số tờ 2 cấp"
        assert final_document_cleanup("Số tờ 3 039") == "Số tờ 3 cấp"
        assert final_document_cleanup("Đơn đăng KỶ DD") == "Đơn đăng ký ĐĐ"
        assert final_document_cleanup("Tên ấy p") == "Tên ấp"
        assert final_document_cleanup("1862 | 2 | 148 | 186 | 60 | 33 | 7,975.90 | ONT-CL | | ấy | p2") == "1862 | 2 | 148 | 186 | 60 | 33 | 7,975.90 | ONT+CLN | | ấp 2"
        assert final_document_cleanup("ONT+CL N") == "ONT+CLN"
        assert final_document_cleanup("ONT-CL") == "ONT+CLN"
        assert final_document_cleanup("ONT%CLN") == "ONT+CLN"
        assert final_document_cleanup("7,975.90 0") == "7,975.90"
        assert final_document_cleanup("Phạm vi: Ấp 2 xã Xuân Sơn |") == "Phạm vi: Ấp 2 xã Xuân Sơn"

    def test_10_mau_so_02_col3_col9_note_normalization(self):
        """Xác minh chuẩn hóa Cột (3), Cột (9) và dòng Ghi chú trên Mẫu số 02."""
        assert final_document_cleanup("Số lượng NCUT (3=445)") == "Số lượng NCUT (3=4+5)"
        assert final_document_cleanup("(3-445)") == "(3=4+5)"
        assert final_document_cleanup("(3 = 445)") == "(3=4+5)"
        assert final_document_cleanup("Già làng, trưởng dòng họ, người thực hiện nghi lễ ngưỡng") == "Già làng, trưởng dòng họ, người thực hiện nghi lễ truyền thống, tín ngưỡng"
        assert final_document_cleanup("người thực hiện nghi lễ ngưỡng") == "người thực hiện nghi lễ truyền thống, tín ngưỡng"
        assert final_document_cleanup("- Từ Cột (Đ) đến cột (19: Ghi cụ thể thành phần của NCUT") == "- Từ Cột (9) đến cột (15): Ghi cụ thể thành phần của NCUT"
        assert final_document_cleanup("Từ Cột (9) đến cột (15):+5): Ghi cụ thể") == "Từ Cột (9) đến cột (15): Ghi cụ thể"

    def test_11_legacy_header_without_box_ownership_preserves_every_line(self, tmp_path: Path):
        """Legacy text/header overlap is ambiguous, so preserve full_text once."""
        from app.core.docx_exporter import export_to_docx
        from app.core.models import OCRResult, OCRTable, ParallelHeader
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        result = OCRResult(
            full_text="Mẫu số 02\nỦY BAN NHÂN DÂN XÃ\nCỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nDANH SÁCH",
            boxes=[],
            parallel_header=ParallelHeader(
                left_lines=["ỦY BAN NHÂN DÂN XÃ"],
                right_lines=["CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"],
            ),
        )

        out_docx = tmp_path / "test_mau_so_02.docx"
        assert export_to_docx(result, str(out_docx)) is True
        doc = docx.Document(str(out_docx))

        assert [p.text for p in doc.paragraphs if p.text] == [
            "Mẫu số 02", "ỦY BAN NHÂN DÂN XÃ", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "DANH SÁCH"
        ]
        assert len(doc.tables) == 0
