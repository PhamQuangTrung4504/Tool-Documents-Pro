#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""A/B đo OCR QThread và worker process; cần chạy trên Windows với model thật."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

# Đồng bộ một số biến môi trường khởi động của run.py trước khi nạp Paddle/Torch.
os.environ.setdefault("KMP_DUPLICATE_LIB_OK", "TRUE")
os.environ.setdefault("PADDLE_PDX_EAGER_INIT", "False")
os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.native_runtime import bootstrap_torch_before_paddle

bootstrap_torch_before_paddle()

import cv2
import numpy as np
from PyQt5.QtCore import QEventLoop, QObject, QTimer
from PyQt5.QtWidgets import QApplication, QLabel

from app.core.docx_exporter import export_to_docx
from app.core.ocr_engine import OCREngine, OCRResult
from app.ui.main_window import OCRWorker
from app.ui.ocr_process_adapter import OCRProcessClient, OCRProcessWorker
from app.core.ocr_process_worker import ocr_result_fingerprint


def _private_bytes(pid: int) -> Optional[int]:
    if os.name != "nt":
        return None
    import ctypes
    from ctypes import wintypes

    class Counters(ctypes.Structure):
        _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD),
                    ("PeakWorkingSetSize", ctypes.c_size_t), ("WorkingSetSize", ctypes.c_size_t),
                    ("QuotaPeakPagedPoolUsage", ctypes.c_size_t), ("QuotaPagedPoolUsage", ctypes.c_size_t),
                    ("QuotaPeakNonPagedPoolUsage", ctypes.c_size_t), ("QuotaNonPagedPoolUsage", ctypes.c_size_t),
                    ("PagefileUsage", ctypes.c_size_t), ("PeakPagefileUsage", ctypes.c_size_t),
                    ("PrivateUsage", ctypes.c_size_t)]
    handle = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)
    if not handle:
        return None
    try:
        counters = Counters()
        counters.cb = ctypes.sizeof(counters)
        if ctypes.windll.psapi.GetProcessMemoryInfo(handle, ctypes.byref(counters), counters.cb):
            return int(counters.PrivateUsage)
        return None
    finally:
        ctypes.windll.kernel32.CloseHandle(handle)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class StageProbe(QObject):
    """Lấy mẫu QTimer thưa, không ghi log từng tick."""
    INTERVAL_MS = 50

    def __init__(self, app: QApplication) -> None:
        super().__init__(app)
        self.timer = QTimer(self)
        self.timer.setInterval(self.INTERVAL_MS)
        self.timer.timeout.connect(self._tick)
        self.samples: List[tuple[int, float]] = []
        self.events: List[dict] = []
        self.deadline_ns = 0

    def start(self) -> None:
        self.samples = []
        self.events = []
        self.deadline_ns = time.perf_counter_ns() + self.INTERVAL_MS * 1_000_000
        self.timer.start()

    def _tick(self) -> None:
        now = time.perf_counter_ns()
        delay = max(0.0, (now - self.deadline_ns) / 1_000_000)
        self.samples.append((now, delay))
        self.deadline_ns = now + self.INTERVAL_MS * 1_000_000

    def on_stage(self, generation: int, stage: str, transition: str, timestamp_ns: int) -> None:
        self.events.append({"generation": generation, "stage": stage, "transition": transition,
                            "timestamp_ns": int(timestamp_ns), "received_ns": time.perf_counter_ns()})

    @staticmethod
    def _stats(values: List[float]) -> dict:
        if not values:
            return {"n": 0, "max_ms": 0.0, "p95_ms": 0.0,
                    "over_100ms": 0, "over_250ms": 0, "over_500ms": 0}
        ordered = sorted(values)
        index = min(len(ordered) - 1, int(len(ordered) * 0.95 + 0.999) - 1)
        return {"n": len(ordered), "max_ms": round(ordered[-1], 1),
                "p95_ms": round(ordered[index], 1),
                "over_100ms": sum(x > 100 for x in ordered),
                "over_250ms": sum(x > 250 for x in ordered),
                "over_500ms": sum(x > 500 for x in ordered)}

    def finish(self, start_ns: int, end_ns: int) -> dict:
        self.timer.stop()
        intervals, active = [], None
        for item in sorted(self.events, key=lambda e: e["timestamp_ns"]):
            if item["transition"] == "start":
                if active:
                    intervals.append(active)
                active = {"stage": item["stage"], "start": item["timestamp_ns"], "end": None}
            elif active and item["stage"] == active["stage"]:
                active["end"] = item["timestamp_ns"]
                intervals.append(active)
                active = None
        if active:
            active["end"] = end_ns
            intervals.append(active)
        per_stage: Dict[str, List[float]] = {}
        for sample_ns, delay in self.samples:
            owner = next((x["stage"] for x in reversed(intervals)
                          if x["start"] <= sample_ns <= (x["end"] or end_ns)), "between_stages")
            per_stage.setdefault(owner, []).append(delay)
        return {
            "interpretation": "Độ trễ timer là tín hiệu đáp ứng event loop, không phải thời gian chính xác của tác vụ.",
            "overall": self._stats([x[1] for x in self.samples]),
            "stages": {key: self._stats(value) for key, value in per_stage.items()},
            "stage_durations_ms": [{"stage": x["stage"], "duration_ms": round((x["end"] - x["start"]) / 1_000_000, 1)}
                                   for x in intervals],
            "worker_signal_delivery": self._stats([
                max(0.0, (x["received_ns"] - x["timestamp_ns"]) / 1_000_000) for x in self.events
            ]),
            "wall_seconds": round((end_ns - start_ns) / 1_000_000_000, 3),
        }


def _engine_config(engine: OCREngine) -> dict:
    return {key: getattr(engine, key) for key in (
        "lang", "use_gpu", "use_angle_cls", "det_db_thresh", "det_db_box_thresh",
        "det_db_unclip_ratio", "drop_score", "rec_batch_num", "ocr_version", "vietocr_model_name",
    )}


def _run_worker(app: QApplication, probe: StageProbe, worker, generation: int) -> tuple[OCRResult, dict]:
    loop = QEventLoop()
    outcome: List[Any] = []
    worker.stage_event.connect(probe.on_stage)
    worker.result_ready.connect(lambda value: outcome.append(value))
    worker.finished.connect(loop.quit)
    start_ns = time.perf_counter_ns()
    probe.start()
    worker.start()
    loop.exec_()
    end_ns = time.perf_counter_ns()
    if not outcome:
        raise RuntimeError("Worker kết thúc mà không gửi kết quả.")
    if isinstance(outcome[0], BaseException):
        raise outcome[0]
    return outcome[0], probe.finish(start_ns, end_ns)


def _verify_docx(result: OCRResult, folder: Path, label: str) -> dict:
    from docx import Document
    path = folder / f"{label}.docx"
    exported = export_to_docx(result, str(path))
    opened = Document(str(path))
    return {"export_ok": bool(exported), "open_ok": True,
            "paragraphs": len(opened.paragraphs), "tables": len(opened.tables),
            "bytes": path.stat().st_size}


def main() -> int:
    parser = argparse.ArgumentParser(description="So sánh OCR trong QThread và tiến trình spawn riêng.")
    parser.add_argument("--image", required=True, type=Path, help="Đường dẫn ảnh đầu vào")
    parser.add_argument("--output", type=Path, default=None, help="JSON báo cáo; mặc định nằm trong TEMP")
    parser.add_argument("--mode", choices=("both", "qthread", "subprocess"), default="both")
    args = parser.parse_args()
    image_path = args.image.expanduser().resolve()
    if not image_path.is_file():
        parser.error(f"Không tìm thấy ảnh: {image_path}")
    if args.mode == "both":
        with tempfile.TemporaryDirectory(prefix="tdp-ocr-ab-runs-") as scratch:
            scratch_path = Path(scratch)
            child_reports = {}
            for mode in ("qthread", "subprocess"):
                child_report = scratch_path / f"{mode}.json"
                command = [sys.executable, str(Path(__file__).resolve()), "--image", str(image_path),
                           "--mode", mode, "--output", str(child_report)]
                completed = subprocess.run(command, cwd=str(ROOT), capture_output=True, text=True, check=False)
                if completed.returncode != 0 or not child_report.is_file():
                    sys.stderr.write(completed.stdout[-4000:])
                    sys.stderr.write(completed.stderr[-4000:])
                    raise RuntimeError(f"Chạy A/B mode={mode} thất bại, exit={completed.returncode}.")
                child_reports[mode] = json.loads(child_report.read_text(encoding="utf-8"))
            runs = {mode: report["run"] for mode, report in child_reports.items()}
            hashes_match = runs["qthread"]["result_sha256"] == runs["subprocess"]["result_sha256"]
            report = {
                "schema_version": 1, "image": image_path.name, "image_sha256": _file_sha256(image_path),
                "image_size": child_reports["qthread"]["image_size"],
                "opencv_threads": child_reports["qthread"]["opencv_threads"], "runs": runs,
                "result_comparison": {
                    "same_full_text": runs["qthread"]["full_text_sha256"] == runs["subprocess"]["full_text_sha256"],
                    "same_result_fingerprint": hashes_match,
                    "same_box_count": runs["qthread"]["boxes"] == runs["subprocess"]["boxes"],
                    "same_table_count": runs["qthread"]["tables"] == runs["subprocess"]["tables"],
                },
                "docx": {mode: item["docx"] for mode, item in child_reports.items()},
                "note": "Hai nhánh được chạy trong tiến trình Python mới để cô lập số đo RAM.",
            }
        output = args.output or (Path(tempfile.gettempdir()) / f"tdp_ocr_process_ab_{time.strftime('%Y%m%d_%H%M%S')}.json")
        output = output.expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        print(f"Báo cáo: {output}")
        return 0
    image = cv2.imread(str(image_path), cv2.IMREAD_UNCHANGED)
    if image is None:
        parser.error("Không đọc được ảnh đầu vào.")

    app = QApplication.instance() or QApplication(sys.argv[:1])
    # Giữ một widget thật để event loop/GUI subsystem khởi tạo tương tự ứng dụng desktop.
    source_view = QLabel("A/B OCR đang chạy")
    source_view.resize(360, 80)
    source_view.show()
    app.processEvents()
    probe = StageProbe(app)
    engine = OCREngine.get_engine(lang="vi")
    client = OCRProcessClient()
    report: Dict[str, Any] = {"schema_version": 1, "image": image_path.name,
                              "image_sha256": _file_sha256(image_path),
                              "image_size": [int(image.shape[1]), int(image.shape[0])],
                              "opencv_threads": int(cv2.getNumThreads()), "mode": args.mode}
    try:
        if args.mode == "qthread":
            worker = OCRWorker(engine, image, generation=1)
            result, metrics = _run_worker(app, probe, worker, 1)
            memory = _private_bytes(os.getpid())
            run = {**metrics, "boxes": len(result.boxes), "tables": len(result.tables),
                   "result_sha256": ocr_result_fingerprint(result),
                   "full_text_sha256": hashlib.sha256(result.full_text.encode("utf-8")).hexdigest(),
                   "process_private_memory_mb": round(memory / (1024 * 1024), 1) if memory is not None else None}
            with tempfile.TemporaryDirectory(prefix="tdp-ocr-ab-docx-") as temp_dir:
                docx = _verify_docx(result, Path(temp_dir), "qthread")
        else:
            worker = OCRProcessWorker(client, image, _engine_config(engine), generation=2)
            result, metrics = _run_worker(app, probe, worker, 2)
            child_memory = client._process_memory_bytes(client._process)
            memory = _private_bytes(os.getpid())
            run = {**metrics, "boxes": len(result.boxes), "tables": len(result.tables),
                   "result_sha256": ocr_result_fingerprint(result),
                   "full_text_sha256": hashlib.sha256(result.full_text.encode("utf-8")).hexdigest(),
                   "parent_private_memory_mb": round(memory / (1024 * 1024), 1) if memory is not None else None,
                   "child_private_memory_mb": round(child_memory / (1024 * 1024), 1) if child_memory is not None else None,
                   "worker_ready_ms": client.worker_ready_ms}
            with tempfile.TemporaryDirectory(prefix="tdp-ocr-ab-docx-") as temp_dir:
                docx = _verify_docx(result, Path(temp_dir), "subprocess")
        report["run"] = run
        report["docx"] = docx
    finally:
        shutdown_thread = threading.Thread(target=client.shutdown, kwargs={"timeout": 5.0}, daemon=True)
        shutdown_thread.start()
        shutdown_loop = QEventLoop()
        shutdown_timer = QTimer()
        shutdown_timer.setInterval(50)
        shutdown_timer.timeout.connect(lambda: shutdown_loop.quit() if not client.is_running else None)
        shutdown_timer.start()
        shutdown_loop.exec_()
        shutdown_timer.stop()
        source_view.close()
    output = args.output or (Path(tempfile.gettempdir()) / f"tdp_ocr_process_ab_{args.mode}_{time.strftime('%Y%m%d_%H%M%S')}.json")
    output = output.expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(f"Báo cáo: {output}")
    return 0


if __name__ == "__main__":
    import multiprocessing
    multiprocessing.freeze_support()
    raise SystemExit(main())

