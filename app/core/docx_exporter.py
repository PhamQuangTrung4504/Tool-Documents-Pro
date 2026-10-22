#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
app.core.docx_exporter
======================
Module xuất kết quả nhận dạng OCR ra định dạng Microsoft Word (.docx):
- Bảo toàn cấu trúc bảng biểu thành Word Table chuẩn.
- Bảo toàn tiêu đề 2 khối song song (Quốc hiệu / Cơ quan ban hành).
- Bảo toàn định dạng căn lề, in đậm và cỡ chữ chuẩn văn bản hành chính.
"""

from __future__ import annotations

import re
from typing import Any, Optional
from loguru import logger

from app.core.models import DocumentBlock, OCRResult, OCRTable, ParallelHeader

__all__ = [
    "export_to_docx",
]


def _estimate_column_widths(table: OCRTable, usable_width_inches: float, font_size_pt: float) -> list[float]:
    """Estimate Word column widths from source geometry, constrained by text."""
    count = max(1, table.num_cols)
    weights: list[float] = []
    boundaries = table.column_boundaries
    if len(boundaries) == count + 1 and all(b > a for a, b in zip(boundaries, boundaries[1:])):
        origin = boundaries[0]
        weights = [max(1.0, boundaries[i + 1] - boundaries[i]) for i in range(count)]
    else:
        by_column: list[list[float]] = [[] for _ in range(count)]
        for cell in table.cells:
            if cell.bbox and cell.col_span == 1 and 0 <= cell.col_start < count:
                width = max(0, cell.bbox[2] - cell.bbox[0])
                if width:
                    by_column[cell.col_start].append(float(width))
        if any(by_column):
            known = [sorted(values)[len(values) // 2] if values else None for values in by_column]
            fallback = float(table.bbox[2] - table.bbox[0]) / count if len(table.bbox) >= 4 else 1.0
            weights = [float(value if value is not None else fallback) for value in known]
        elif len(table.bbox) >= 4 and table.bbox[2] > table.bbox[0]:
            weights = [float(table.bbox[2] - table.bbox[0]) / count] * count
        else:
            weights = [1.0] * count

    total_weight = max(1.0, sum(weights))
    geometry_widths = [usable_width_inches * weight / total_weight for weight in weights]

    # Keep long unbreakable OCR tokens from forcing a narrow cell into vertical
    # character wrapping. This is a lower-bound constraint only; geometry sets
    # the preferred proportions whenever the page can accommodate them.
    minima = []
    for col in range(count):
        longest_token = max(
            (len(token) for row in table.grid if col < len(row) for token in re.findall(r"\S+", row[col])),
            default=1,
        )
        minima.append(min(0.9, max(0.32, longest_token * font_size_pt * 0.048 / 72.0)))

    min_sum = sum(minima)
    if min_sum >= usable_width_inches:
        scale = usable_width_inches / max(min_sum, 1e-9)
        return [minimum * scale for minimum in minima]

    remainder = usable_width_inches - min_sum
    residuals = [max(0.0, geometry - minimum) for geometry, minimum in zip(geometry_widths, minima)]
    residual_sum = sum(residuals)
    if residual_sum <= 1e-9:
        residuals = weights
        residual_sum = sum(residuals)
    return [minimum + remainder * residual / residual_sum for minimum, residual in zip(minima, residuals)]


def _document_blocks_for_export(result: OCRResult) -> list[DocumentBlock]:
    """Use explicit blocks, or build a geometry-linked compatibility view.

    The legacy no-box case cannot establish where a table belongs among text.
    It preserves legacy text verbatim and reports the missing structural link
    instead of suppressing paragraphs by comparing their words to table cells.
    """
    if result.document_blocks:
        return sorted(result.document_blocks, key=lambda block: block.order)

    blocks: list[DocumentBlock] = []
    box_indices = {id(box): index for index, box in enumerate(result.boxes)}
    if result.boxes:
        from app.core.layout_analyzer import group_boxes_into_lines

        table_ownership: dict[int, set[int]] = {}
        for table_index, table in enumerate(result.tables):
            owned = {id(box) for row in table.row_boxes for box in row}
            owned.update(id(box) for cell in table.cells for box in cell.boxes)
            table_ownership[table_index] = owned
        header_owned = set()
        if result.parallel_header is not None:
            header_owned = {id(box) for box in result.parallel_header.left_boxes + result.parallel_header.right_boxes}
            header_boxes = [box for box in result.boxes if id(box) in header_owned]
            blocks.append(DocumentBlock(
                kind="parallel_header", order=0,
                text="\n".join(result.parallel_header.left_lines + result.parallel_header.right_lines),
                raw_text="\n".join(box.raw_text for box in header_boxes),
                box_indices=[box_indices[id(box)] for box in header_boxes],
                bbox=_boxes_bbox(header_boxes),
            ))
        emitted: set[int] = set()
        for line in group_boxes_into_lines(result.boxes):
            if not line or all(id(box) in header_owned for box in line):
                continue
            owners = [
                (sum(id(box) in table_ownership[index] for box in line), index)
                for index in range(len(result.tables))
            ]
            owner_count, table_index = max(owners, default=(0, -1))
            owned_ids = table_ownership.get(table_index, set()) if owner_count else set()
            line_table_boxes = [box for box in line if id(box) in owned_ids]
            paragraph_boxes = [box for box in line if id(box) not in owned_ids]
            if line_table_boxes and table_index not in emitted:
                table = result.tables[table_index]
                blocks.append(DocumentBlock(
                    kind="table", order=len(blocks), table_index=table_index,
                    bbox=list(table.bbox) if table.bbox else None,
                ))
                emitted.add(table_index)
            if paragraph_boxes:
                blocks.append(DocumentBlock(
                    kind="paragraph", order=len(blocks),
                    text=" ".join(box.text.strip() for box in paragraph_boxes if box.text.strip()),
                    raw_text=" ".join(box.raw_text.strip() for box in paragraph_boxes if box.raw_text.strip()),
                    box_indices=[box_indices[id(box)] for box in paragraph_boxes],
                    bbox=_boxes_bbox(paragraph_boxes),
                ))

        # Insert empty/boxless tables at their physical vertical position.
        for table_index, table in enumerate(result.tables):
            if table_index in emitted:
                continue
            top = table.bbox[1] if len(table.bbox) >= 4 else float("inf")
            insert_at = next((i for i, block in enumerate(blocks)
                              if block.kind != "parallel_header" and block.bbox
                              and block.bbox[1] >= top), len(blocks))
            blocks.insert(insert_at, DocumentBlock(
                kind="table", order=insert_at, table_index=table_index,
                bbox=list(table.bbox) if table.bbox else None,
            ))
    else:
        blocks.extend(
            DocumentBlock(kind="paragraph", order=index, text=line, raw_text=line)
            for index, line in enumerate(result.full_text.splitlines())
        )
        if result.tables and result.full_text.strip():
            logger.warning(
                "Legacy OCRResult has tables and text but no OCRBox ownership; "
                "table placement is ambiguous, so table structure is omitted to avoid duplicate text"
            )
        else:
            for table_index, table in enumerate(result.tables):
                blocks.append(DocumentBlock(
                    kind="table", order=len(blocks), table_index=table_index,
                    bbox=list(table.bbox) if table.bbox else None,
                ))
    for order, block in enumerate(blocks):
        block.order = order
    return blocks


def _boxes_bbox(boxes: list[Any]) -> Optional[list[float]]:
    if not boxes:
        return None
    return [
        min(point[0] for box in boxes for point in box.bbox),
        min(point[1] for box in boxes for point in box.bbox),
        max(point[0] for box in boxes for point in box.bbox),
        max(point[1] for box in boxes for point in box.bbox),
    ]

def export_to_docx(result: OCRResult, output_path: str) -> bool:
    """
    Xuất kết quả OCR sang tài liệu Microsoft Word (.docx).
    - Bảo toàn cấu trúc bảng biểu thành Word Table chuẩn khi có bảng.
    - Tuyệt đối KHÔNG in lặp lại nội dung bảng thành văn bản thường.
    - Tự động xoay khổ giấy Ngang (Landscape) cho bảng rộng (>= 8 cột).
    - Bảo toàn tiêu đề 2 khối song song và căn chỉnh chuẩn văn bản hành chính.
    """
    try:
        import docx  # type: ignore[import-untyped,import-not-found]
        from docx.shared import Inches, Pt, RGBColor, Mm  # type: ignore[import-untyped,import-not-found]
        from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-untyped,import-not-found]
        from docx.enum.table import WD_TABLE_ALIGNMENT  # type: ignore[import-untyped,import-not-found]
        from docx.enum.section import WD_ORIENT  # type: ignore[import-untyped,import-not-found]
        from docx.oxml import OxmlElement, parse_xml  # type: ignore[import-untyped,import-not-found]
        from docx.oxml.ns import nsdecls, qn  # type: ignore[import-untyped,import-not-found]
    except ImportError:
        logger.error("python-docx is not installed. Cannot export to DOCX.")
        return False

    try:
        import re
        doc: Any = docx.Document()

        # Kiểm tra bảng rộng (num_cols >= 8) để tự động xoay khổ giấy Ngang (Landscape)
        has_wide_table = bool(result.tables and any(tbl.num_cols >= 8 for tbl in result.tables))
        section = doc.sections[0]
        if has_wide_table:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width = Mm(297)
            section.page_height = Mm(210)
            section.left_margin = Mm(12)
            section.right_margin = Mm(12)
            section.top_margin = Mm(15)
            section.bottom_margin = Mm(15)
        else:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(0.8)

        # Thiết lập style Normal
        style: Any = doc.styles["Normal"]
        font = style.font
        font.name = "Times New Roman"
        font.size = Pt(12)
        font.color.rgb = RGBColor(0, 0, 0)
        document_blocks = _document_blocks_for_export(result)
        has_parallel_header_block = any(block.kind == "parallel_header" for block in document_blocks)

        # Parallel headers are represented separately and retain their source text.
        if (
            result.parallel_header
            and has_parallel_header_block
            and (result.parallel_header.left_lines or result.parallel_header.right_lines)
        ):
            hdr = result.parallel_header

            hdr_table: Any = doc.add_table(rows=1, cols=2)
            hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER

            try:
                tblPr: Any = hdr_table._tbl.tblPr
                borders = parse_xml(
                    r'<w:tblBorders %s><w:top w:val="none"/><w:left w:val="none"/>'
                    r'<w:bottom w:val="none"/><w:right w:val="none"/>'
                    r'<w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>' % nsdecls("w")
                )
                tblPr.append(borders)
            except Exception as e_border:
                logger.debug(f"Could not remove table borders: {e_border}")

            valid_left_lines = [line for line in hdr.left_lines if line.strip()]
            align_left = WD_ALIGN_PARAGRAPH.LEFT
            cell_left = hdr_table.cell(0, 0)
            for idx, line_text in enumerate(valid_left_lines):
                p = cell_left.paragraphs[0] if idx == 0 else cell_left.add_paragraph()
                p.alignment = align_left
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(line_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

            valid_right_lines = [line for line in hdr.right_lines if line.strip()]
            align_right = WD_ALIGN_PARAGRAPH.RIGHT
            cell_right = hdr_table.cell(0, 1)
            for idx, line_text in enumerate(valid_right_lines):
                p = cell_right.paragraphs[0] if idx == 0 else cell_right.add_paragraph()
                p.alignment = align_right
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(line_text)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_before = Pt(4)
            p_spacer.paragraph_format.space_after = Pt(4)

        def _render_word_table(tbl: OCRTable) -> Any:
            """Vẽ đúng 1 bảng Word từ OCRTable."""
            is_wide = tbl.num_cols >= 8
            font_size_pt = Pt(9.0) if is_wide else Pt(11.0)
            doc_table: Any = doc.add_table(rows=len(tbl.grid), cols=tbl.num_cols)
            doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            doc_table.style = "Table Grid"
            doc_table.autofit = False

            topology_errors = tbl.topology_issues()
            if topology_errors:
                raise ValueError(f"Invalid table topology: {', '.join(sorted(set(topology_errors)))}")

            total_width_inches = (
                section.page_width - section.left_margin - section.right_margin
            ) / Inches(1)
            col_widths_inches = _estimate_column_widths(tbl, total_width_inches, float(font_size_pt.pt))

            # Keep tblGrid and fixed table layout in agreement. python-docx
            # autofit can otherwise override the source-derived grid in Word.
            tbl_pr = doc_table._tbl.tblPr
            layout = tbl_pr.find(qn("w:tblLayout"))
            if layout is None:
                layout = OxmlElement("w:tblLayout")
                tbl_pr.append(layout)
            layout.set(qn("w:type"), "fixed")
            tbl_width = tbl_pr.find(qn("w:tblW"))
            if tbl_width is None:
                tbl_width = OxmlElement("w:tblW")
                tbl_pr.append(tbl_width)
            tbl_width.set(qn("w:w"), str(int(round(total_width_inches * 1440))))
            tbl_width.set(qn("w:type"), "dxa")
            for column, width in zip(doc_table.columns, col_widths_inches):
                column.width = Inches(width)

            # Repeat only rows explicitly identified as headers by upstream
            # layout metadata; row zero is not assumed to be a header.
            for header_row_index in sorted(set(tbl.header_rows)):
                if not 0 <= header_row_index < len(doc_table.rows):
                    continue
                try:
                    header_tr: Any = doc_table.rows[header_row_index]._tr.get_or_add_trPr()
                    header_tr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls("w")))
                except Exception as header_error:
                    raise RuntimeError(f"Could not mark declared header row {header_row_index}") from header_error

            def _clean_cell_content(cell_str: str, col_idx: int) -> str:
                # Keep recognized content unchanged; semantic correction belongs
                # in an auditable OCR correction workflow with visual evidence.
                return cell_str

            has_merged_cells = bool(tbl.cells and any(c.is_merged for c in tbl.cells))
            if has_merged_cells:
                # Đổ dữ liệu các ô không gộp
                for c in tbl.cells:
                    if not c.is_merged:
                        try:
                            w_cell = doc_table.cell(c.row_start, c.col_start)
                            w_cell.text = _clean_cell_content(c.text, c.col_start)
                        except Exception as e_cell:
                            raise RuntimeError(
                                f"Could not populate table cell at ({c.row_start}, {c.col_start})"
                            ) from e_cell
                # Đổ dữ liệu các ô gộp
                for c in tbl.cells:
                    if c.is_merged:
                        try:
                            c_tl = doc_table.cell(c.row_start, c.col_start)
                            c_br = doc_table.cell(c.row_end - 1, c.col_end - 1)
                            merged_w_cell = c_tl.merge(c_br)
                            merged_w_cell.text = _clean_cell_content(c.text, c.col_start)
                        except Exception as e_merge:
                            raise RuntimeError(
                                f"Could not merge table cell span ({c.row_start}, {c.col_start})-({c.row_end}, {c.col_end})"
                            ) from e_merge
            else:
                for r_i, row_data in enumerate(tbl.grid):
                    for c_i in range(min(len(row_data), tbl.num_cols)):
                        cell = doc_table.cell(r_i, c_i)
                        cell.text = _clean_cell_content(row_data[c_i], c_i)

            # Apply source edge metadata and explicit fixed widths to Word XML.
            cells_by_origin = {
                (cell.row_start, cell.col_start): cell
                for cell in tbl.cells
            }
            for row_index, row in enumerate(doc_table.rows):
                grid_col = 0
                for tc in row._tr.tc_lst:
                    tc_pr = tc.get_or_add_tcPr()
                    span_el = tc_pr.find(qn("w:gridSpan"))
                    col_span = int(span_el.get(qn("w:val"))) if span_el is not None else 1
                    col_span = max(1, min(col_span, tbl.num_cols - grid_col))
                    cell_width = sum(col_widths_inches[grid_col:grid_col + col_span])
                    tc_width = tc_pr.find(qn("w:tcW"))
                    if tc_width is None:
                        tc_width = OxmlElement("w:tcW")
                        tc_pr.append(tc_width)
                    tc_width.set(qn("w:w"), str(int(round(cell_width * 1440))))
                    tc_width.set(qn("w:type"), "dxa")

                    source_cell = cells_by_origin.get((row_index, grid_col))
                    if source_cell is not None and source_cell.border_edges is not None:
                        borders = tc_pr.find(qn("w:tcBorders"))
                        if borders is None:
                            borders = OxmlElement("w:tcBorders")
                            tc_pr.append(borders)
                        for edge_name in ("top", "bottom", "left", "right"):
                            edge = borders.find(qn(f"w:{edge_name}"))
                            if edge is None:
                                edge = OxmlElement(f"w:{edge_name}")
                                borders.append(edge)
                            edge.set(qn("w:val"), "single" if source_cell.border_edges.get(edge_name) else "nil")
                    grid_col += col_span

            # Diagonal cell borders are stored as explicit topology metadata and
            # emitted as editable Word table-cell borders, not rasterized content.
            for source_cell in tbl.cells:
                if source_cell.diagonal_border not in {"tl2br", "tr2bl"}:
                    continue
                word_cell = doc_table.cell(source_cell.row_start, source_cell.col_start)
                if source_cell.bbox and source_cell.boxes:
                    x1, y1, x2, y2 = source_cell.bbox
                    cell_w = max(1, x2 - x1)
                    cell_h = max(1, y2 - y1)
                    box_info = []
                    for box in source_cell.boxes:
                        left = min(point[0] for point in box.bbox)
                        right = max(point[0] for point in box.bbox)
                        top = min(point[1] for point in box.bbox)
                        bottom = max(point[1] for point in box.bbox)
                        box_info.append({
                            "box": box,
                            "cx": (left + right) / 2.0,
                            "cy": (top + bottom) / 2.0,
                            "height": max(1, bottom - top),
                        })
                    median_height = sorted(item["height"] for item in box_info)[len(box_info) // 2]
                    text_lines = []
                    for item in sorted(box_info, key=lambda value: (value["cy"], value["cx"])):
                        existing = next((
                            line for line in text_lines
                            if abs(item["cy"] - line["cy"]) <= 0.55 * median_height
                        ), None)
                        if existing is None:
                            text_lines.append({"cy": item["cy"], "items": [item]})
                        else:
                            existing["items"].append(item)

                    first_paragraph = True
                    for line in sorted(text_lines, key=lambda value: value["cy"]):
                        line_items = sorted(line["items"], key=lambda value: value["cx"])
                        line_text = " ".join(item["box"].text.strip() for item in line_items if item["box"].text.strip())
                        if not line_text:
                            continue
                        center_x = sum(item["cx"] for item in line_items) / len(line_items)
                        center_y = sum(item["cy"] for item in line_items) / len(line_items)
                        relative_x = min(1.0, max(0.0, (center_x - x1) / cell_w))
                        if source_cell.diagonal_border == "tl2br":
                            diagonal_y = y1 + relative_x * cell_h
                            alignment = WD_ALIGN_PARAGRAPH.RIGHT if center_y < diagonal_y else WD_ALIGN_PARAGRAPH.LEFT
                        else:
                            diagonal_y = y1 + (1.0 - relative_x) * cell_h
                            alignment = WD_ALIGN_PARAGRAPH.LEFT if center_y < diagonal_y else WD_ALIGN_PARAGRAPH.RIGHT
                        paragraph = word_cell.paragraphs[0] if first_paragraph else word_cell.add_paragraph()
                        first_paragraph = False
                        paragraph.clear()
                        paragraph.alignment = alignment
                        paragraph.add_run(line_text)

                tc_pr = word_cell._tc.get_or_add_tcPr()
                borders = tc_pr.find(qn("w:tcBorders"))
                if borders is None:
                    borders = OxmlElement("w:tcBorders")
                    tc_pr.append(borders)
                diagonal = OxmlElement(f"w:{source_cell.diagonal_border}")
                diagonal.set(qn("w:val"), "single")
                diagonal.set(qn("w:sz"), "8")
                borders.append(diagonal)

            # Định dạng font, cỡ chữ, căn lề và độ rộng cột
            for r_i, row in enumerate(doc_table.rows):
                for c_i, cell in enumerate(row.cells):
                    if c_i < len(col_widths_inches):
                        try:
                            cell.width = Inches(col_widths_inches[c_i])
                        except Exception:
                            pass
                    cell_text = cell.text.strip()
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(1) if is_wide else Pt(2)
                        p.paragraph_format.space_after = Pt(1) if is_wide else Pt(2)
                        p.paragraph_format.line_spacing = 1.0 if is_wide else 1.15
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = font_size_pt
                            if r_i in tbl.header_rows:
                                run.bold = True
                        if r_i == 0 or (c_i != 1 and len(cell_text) <= 6) or re.match(r"^\(\d+\)$", cell_text):
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            return doc_table

        # 3. Export the explicit document order; legacy results are converted
        # to geometry-linked blocks above without matching their words.
        rendered_table_indexes: set[int] = set()
        last_was_major_title = False
        hit_major_title = False

        for block in document_blocks:
            if block.kind == "parallel_header":
                # The structured header has already been emitted in its own
                # two-column Word layout above.
                continue
            if block.kind == "table":
                if block.table_index is None or not 0 <= block.table_index < len(result.tables):
                    raise ValueError("Document table block refers to an invalid table index")
                if block.table_index in rendered_table_indexes:
                    raise ValueError(f"Table {block.table_index} appears more than once in document order")
                _render_word_table(result.tables[block.table_index])
                rendered_table_indexes.add(block.table_index)
                continue
            if block.kind not in {"paragraph", "heading", "caption"}:
                raise ValueError(f"Unsupported document block kind: {block.kind}")

            line_str = block.text.strip()
            if not line_str:
                continue

            # Dòng văn bản thông thường ngoài bảng (Pre-table hoặc Post-table)
            in_table_range = False
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15

            visual_format = block.metadata.get("format", {}) if isinstance(block.metadata, dict) else {}
            alignment = visual_format.get("alignment")
            alignment_map = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            p.alignment = alignment_map.get(
                alignment,
                WD_ALIGN_PARAGRAPH.CENTER if block.kind == "heading" else WD_ALIGN_PARAGRAPH.LEFT,
            )
            if block.kind == "heading":
                p.paragraph_format.space_before = Pt(6)
            run = p.add_run(line_str)
            run.font.name = "Times New Roman"
            run.font.size = Pt(visual_format.get("font_size_pt", 13 if block.kind == "heading" else 12))
            # Preserve three-state style semantics: explicit True/False is
            # authoritative, and evidence marked unknown must not be promoted
            # to bold merely because the structural block is a heading.
            if "bold" in visual_format:
                run.bold = visual_format["bold"] if isinstance(visual_format["bold"], bool) else None
            elif "bold_evidence" not in visual_format and block.kind == "heading":
                # Compatibility for legacy heading blocks without visual
                # style metadata.
                run.bold = True
            if "italic" in visual_format and isinstance(visual_format["italic"], bool):
                run.italic = visual_format["italic"]
        # If geometry-linked legacy fallback could not include tables because
        # there were no OCRBox ownership links, do not silently misplace them.
        legacy_ambiguous = bool(result.tables and result.full_text.strip() and not result.boxes and not result.document_blocks)
        expected_table_indexes = set() if legacy_ambiguous else set(range(len(result.tables)))
        if rendered_table_indexes != expected_table_indexes:
            missing = sorted(expected_table_indexes - rendered_table_indexes)
            if missing:
                raise ValueError(f"Document order omits table blocks: {missing}")

        doc.save(output_path)
        logger.info(f"DOCX exported successfully: {output_path}")
        return True
    except Exception as exc:
        logger.error(f"Error exporting to DOCX: {exc}")
        return False

